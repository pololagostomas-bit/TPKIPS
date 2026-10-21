"""Private alert history on disposable SQLite and synthetic Excel only."""
import csv
import io
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from backend.services import purchase_alerts as alerts
from tests import test_purchase_alerts as fixtures


def source(rows):
    book = Workbook()
    book.active.title = 'BUYER'
    book.active.append(['ORDEN DE COMPRA', 'PEDIDO TRITON', 'OV'])
    for row in rows:
        book.active.append(row)
        for cell in book.active[book.active.max_row]:
            if cell.data_type == 'f':
                cell.data_type = 's'
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class PurchaseAgingTests(unittest.TestCase):
    setUp = fixtures.PurchaseTests.setUp
    tearDown = fixtures.PurchaseTests.tearDown
    set_rows = fixtures.PurchaseTests.set_rows

    def set_source(self, rows):
        self.snapshot = alerts.read_cases(source(rows), [])
        self.snapshot['checked_at'] = alerts._stamp(self.now)

    def history(self, **kwargs):
        return alerts.case_history(now=self.now, **kwargs)

    def test_x_is_ignored_for_both_rules_but_not_prefix_x(self):
        cases = alerts.read_cases(source([
            ['PRE1', ' x ', 'COT-1'], ['PRE2', 'X', 'OV1'],
            ['PRE3', 'X123', 'OV2']]), [])['cases']
        self.assertEqual([c['ip'] for c in cases], ['X123'])

    def test_cot_is_grave_with_any_oc_or_ip_even_approved_oc(self):
        cases = alerts.read_cases(source([
            ['12345', 'IP1', ' cot-1 '], ['PRE1', '', 'COT2'],
            ['', 'IP3', 'COT3'], ['', '', 'COT4'], ['PRE9', 'IP9', 'OV9']]), [])['cases']
        grave = [c for c in cases if c.get('issue_type') == 'COT']
        self.assertEqual({c['ov'] for c in grave}, {'cot-1', 'COT2', 'COT3'})
        self.assertTrue(all(c['severity'] == 'GRAVE' for c in grave))

    def test_pre_edits_and_reordering_keep_clock_then_resolution_freezes_time(self):
        alerts.run(force=True, now=self.now)
        first = self.history()['items'][0]
        self.set_rows([['12345', 'UNRELATED'], ['PRE-CHANGED', 'IP-1']])
        alerts.run(force=True, now=self.now + 7200)
        item = alerts.case_history(now=self.now + 10800)['items'][0]
        self.assertEqual(item['id'], first['id'])
        self.assertEqual(item['elapsed_hours'], 3)
        self.set_rows([['12345', 'IP-1']])
        alerts.run(force=True, now=self.now + 14400)
        resolved = alerts.case_history(state='all', now=self.now + 90000)['items'][0]
        self.assertEqual(resolved['elapsed_hours'], 4)
        self.assertEqual(resolved['status'], 'RESUELTO')
        self.set_rows([['PRE-AGAIN', 'IP-1']])
        alerts.run(force=True, now=self.now + 18000)
        self.assertEqual(alerts.case_history(state='all')['total'], 2)

    def test_failed_or_missing_sheet_snapshot_never_resolves_or_advances_observation(self):
        alerts.run(force=True, now=self.now)
        first = self.history()['items'][0]
        self.snapshot['cases'] = []
        self.snapshot['sheets'] = ['OTHER']
        with self.assertRaises(ValueError):
            alerts.run(force=True, now=self.now + 3600)
        self.assertEqual(self.history()['items'][0]['last_seen_at'], first['last_seen_at'])
        self.assertEqual(self.history()['total'], 1)

    def test_cot_and_pre_clocks_are_independent_and_x_closes_current_episodes(self):
        self.set_source([['PRE1', 'IP1', 'OV1']])
        alerts.run(force=True, now=self.now)
        self.set_source([['PRE2', 'IP1', 'COT1']])
        alerts.run(force=True, now=self.now + 3600)
        items = alerts.case_history(now=self.now + 7200)['items']
        self.assertEqual({c['issue_type']: c['elapsed_hours'] for c in items}, {'PRELIMINAR': 2, 'COT': 1})
        self.set_source([['PRE2', ' x ', 'COT1']])
        alerts.run(force=True, now=self.now + 10800)
        self.assertEqual(self.history()['total'], 0)
        self.assertEqual(alerts.case_history(state='all')['total'], 2)
        self.post_mock.assert_not_called()

    def test_history_pagination_filters_and_formula_safe_export(self):
        self.set_source([['PRE1', '=DANGER', 'COT1'], ['PRE2', 'IP2', 'OV2']])
        alerts.run(force=True, now=self.now)
        self.assertEqual(self.history(severity='GRAVE')['total'], 1)
        self.assertEqual(len(self.history(limit=1)['items']), 1)
        text = alerts.history_csv(state='all', now=self.now).decode('utf-8-sig')
        rows = list(csv.reader(io.StringIO(text)))
        self.assertTrue(any("'=DANGER" in row for row in rows))
        with patch.object(alerts, 'fetch_snapshot', side_effect=AssertionError('History cannot download Excel')):
            self.history()
        self.post_mock.assert_not_called()

    def test_excluded_sheet_does_not_fake_resolution_and_ov_formula_fails_closed(self):
        alerts.run(force=True, now=self.now)
        alerts.save({'excluded_sheets': ['BUYER']}, 'admin.qa')
        self.snapshot = {'cases': [], 'sheets': ['OTHER'], 'excluded_sheets': ['BUYER']}
        alerts.run(force=True, now=self.now + 3600)
        self.assertEqual(self.history()['total'], 1)
        book = Workbook()
        book.active.append(['ORDEN DE COMPRA', 'PEDIDO TRITON', 'OV'])
        book.active.append(['PRE1', 'IP1', '=1+1'])
        output = io.BytesIO(); book.save(output); book.close()
        with self.assertRaises(ValueError):
            alerts.read_cases(output.getvalue(), [])

    def test_cot_priority_and_aging_are_in_mail_without_changing_saved_speech(self):
        self.set_source([['PRE1', 'IP1', 'COT1']])
        alerts.run(force=True, now=self.now)
        result = alerts.run(force=True, now=self.now + 7200)
        self.assertIn('GRAVE: cotizacion COT', result['preview']['html'])
        self.assertIn('2.0 h', result['preview']['html'])
        self.assertEqual(alerts.settings()['body'], alerts.DEFAULTS['body'])
        self.post_mock.assert_not_called()

    def test_removed_ov_column_cannot_resolve_a_cot_alert(self):
        self.set_source([['12345', 'IP1', 'COT1']])
        alerts.run(force=True, now=self.now)
        self.set_rows([['12345', 'IP1']])
        with self.assertRaises(ValueError):
            alerts.run(force=True, now=self.now + 3600)
        self.assertEqual(self.history(severity='GRAVE')['total'], 1)

    def test_excluded_duplicate_source_remains_unverified_until_rechecked(self):
        self.snapshot['sheets'].append('BUYER-B')
        self.snapshot['cases'][0]['sources'].append({'sheet': 'BUYER-B', 'row': 2})
        alerts.run(force=True, now=self.now)
        alerts.save({'excluded_sheets': ['BUYER-B']}, 'admin.qa')
        self.set_rows([['PRE-100001', 'IP-1']])
        alerts.run(force=True, now=self.now + 3600)
        self.set_rows([['12345', 'IP-1']])
        alerts.run(force=True, now=self.now + 7200)
        item = self.history()['items'][0]
        self.assertEqual(item['unverified_sheets'], ['BUYER-B'])
        alerts.save({'excluded_sheets': []}, 'admin.qa')
        self.snapshot['sheets'].append('BUYER-B')
        alerts.run(force=True, now=self.now + 10800)
        self.assertEqual(self.history()['total'], 0)


if __name__ == '__main__':
    unittest.main()
