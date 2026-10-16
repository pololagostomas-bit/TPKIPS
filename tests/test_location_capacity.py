import base64
import sqlite3
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from backend.services import location_capacity as capacity


class LocationCapacityTests(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(':memory:')
        self.db.row_factory = sqlite3.Row
        capacity.init_schema(self.db)
        self.db.execute('''CREATE TABLE inventory_stock(item_key TEXT,warehouse_key TEXT,snapshot_qty REAL,
            default_location TEXT,is_current INTEGER,snapshot_at TEXT)''')

    def tearDown(self):
        self.db.close()

    def save(self, payload):
        checked = capacity.preview(self.db, payload, 'ADMINISTRADOR')
        return capacity.save_master(self.db, {**payload,'revision':checked['revision'],
            'digest':checked['digest'],'reason':'Maestro de prueba'}, 'qa.admin', 'ADMINISTRADOR')

    def csv(self, text):
        return {'filename':'qa.csv','file_base64':base64.b64encode(text.encode('utf-8')).decode()}

    def add_stock(self, code, sku='QA-NP', qty=1, warehouse='1', current=1):
        self.db.execute('INSERT INTO inventory_stock VALUES(?,?,?,?,?,?)',
                        (sku,warehouse,qty,code,current,'2026-10-14T09:00:00'))

    def test_no_master_or_cut_does_not_invent_capacity_or_occupancy(self):
        data = capacity.catalog(self.db, 'PICKER')
        self.assertEqual(data['summary']['registered'], 0)
        self.assertIsNone(data['summary']['ratio'])
        self.add_stock('01-01-A-01-01')
        data = capacity.catalog(self.db, 'PICKER')
        self.assertEqual(data['summary']['installed'], 0)
        self.assertEqual(data['summary']['unmapped'], 1)
        self.assertIsNone(data['summary']['ratio'])
        self.assertEqual(data['items'], [])

    def test_master_is_separate_from_stock_defaults_and_metrics_remain_global(self):
        self.save(self.csv('Ubicacion,Zona,Rack,Estado\nQA-001,A,R1,HABILITADA\nQA-002,A,R1,HABILITADA\nQA-003,B,R2,BLOQUEADA\nQA-004,B,R2,BAJA'))
        self.add_stock(' qa-001 ', 'NP1');self.add_stock('QA-001', 'NP2')
        self.add_stock('QA-003', 'NP3');self.add_stock('QA-004', 'NP4')
        self.add_stock('', 'MISSING');self.add_stock('QA-002', 'ZERO', 0)
        self.add_stock('QA-002', 'NEGATIVE', -1);self.add_stock('QA-002', 'OTHER', 1, '2')
        self.add_stock('QA-002', 'OLD', 9, current=0)
        data = capacity.catalog(self.db, 'PICKER', search='QA-001')
        self.assertEqual(data['total'], 1)
        self.assertEqual(data['items'][0]['sku_count'], 2)
        self.assertEqual(data['summary'], {'registered':4,'installed':3,'enabled':2,'blocked':1,
            'with_stock':1,'blocked_with_stock':1,'without_stock':1,'unmapped':1,'ratio':50.0})
        self.assertEqual(data['stock']['missing_location'], 1)
        self.assertEqual(capacity.catalog(self.db,'PICKER',evidence='UNMAPPED')['items'][0]['code'], 'QA-004')

    def test_import_adds_updates_and_audits_without_deleting_or_reactivating_omitted_fields(self):
        saved = self.save(self.csv('Ubicacion,Zona,Rack,Estado\nQA-001,A,R1,BLOQUEADA\nQA-002,B,R2,HABILITADA'))
        self.assertEqual(saved['added'], 2)
        saved = self.save({'text':'qa-001'})
        self.assertEqual(saved['unchanged'], 1)
        self.assertEqual(saved['revision'], 1)
        self.assertEqual(self.db.execute("SELECT state FROM warehouse_locations WHERE code='QA-001'").fetchone()[0], 'BLOQUEADA')
        self.assertEqual(self.db.execute('SELECT count(*) FROM warehouse_locations').fetchone()[0], 2)
        self.save(self.csv('Ubicacion,Zona,Rack,Estado\nQA-001,A,R1,HABILITADA'))
        history = capacity.history(self.db, 'QA-001', 'ADMINISTRADOR')
        self.assertEqual(len(history), 2)
        self.assertEqual(history[0]['username'], 'qa.admin')
        self.assertEqual(history[0]['reason'], 'Maestro de prueba')
        self.assertIn('BLOQUEADA', history[0]['before_json'])

    def test_preview_is_read_only_and_digest_revision_reason_guard_commits(self):
        before = self.db.total_changes
        payload = {'text':'QA-001'}
        checked = capacity.preview(self.db, payload, 'ADMINISTRADOR')
        self.assertEqual(self.db.total_changes, before)
        for patch_values in [{'digest':'tampered'}, {'revision':True}, {'reason':''}]:
            with self.assertRaises(ValueError):
                capacity.save_master(self.db, {**payload,**checked,'reason':'Carga de prueba',**patch_values}, 'qa', 'ADMINISTRADOR')
        self.save({'text':'QA-OTHER'})
        with self.assertRaisesRegex(ValueError, 'cambió'):
            capacity.save_master(self.db, {**payload,**checked,'reason':'Carga antigua'}, 'qa', 'ADMINISTRADOR')
        self.assertIsNone(self.db.execute("SELECT code FROM warehouse_locations WHERE code='QA-001'").fetchone())

    def test_bad_batch_never_partially_updates_the_master(self):
        for payload in [{'text':'QA-001\nqa-001'}, {'text':'QA-001\n<script>'},
                        self.csv('Ubicacion,Estado\nQA-001,HABILITADA\nQA-002,INVALID')]:
            with self.assertRaises(ValueError):
                self.save(payload)
        self.assertEqual(self.db.execute('SELECT count(*) FROM warehouse_locations').fetchone()[0], 0)

    def test_files_csv_single_column_utf8_and_excel_sheet_selection(self):
        self.assertEqual(capacity.parse_master(self.csv('Ubicación\n01-01-A-01-01'))[0][0]['code'],'01-01-A-01-01')
        workbook = Workbook();workbook.active.title = 'Otra'
        sheet = workbook.create_sheet('Ubicaciones');sheet.append(['Ubicación','Zona','Rack'])
        sheet.append(['01-01-A-01-01','ZONA A','R1'])
        stream = BytesIO();workbook.save(stream);workbook.close()
        payload = {'filename':'maestro.xlsx','sheet':'Ubicaciones','file_base64':base64.b64encode(stream.getvalue()).decode()}
        rows, _, _ = capacity.parse_master(payload)
        self.assertEqual(rows, [{'code':'01-01-A-01-01','zone':'ZONA A','rack':'R1'}])
        self.assertEqual(self.save(payload)['added'], 1)

    def test_upload_bounds_headers_and_file_types(self):
        bad = [{'text':''},{'text':'QA-1','file_base64':'AAAA'}, {'filename':'bad.xls','file_base64':'AAAA'},
               {'filename':'bad.xlsx','file_base64':'????'}, self.csv('Nombre,Estado\nQA-1,HABILITADA'),
               {'text':'\n'.join('QA-'+str(i) for i in range(capacity.MAX_ROWS + 1))}]
        for payload in bad:
            with self.assertRaises(ValueError):
                capacity.parse_master(payload)

    def excel(self, workbook, **options):
        stream = BytesIO();workbook.save(stream);workbook.close()
        return {'filename':'positions-qa.xlsx','file_base64':base64.b64encode(stream.getvalue()).decode(),**options}

    def test_detects_master_not_active_label_sheet_and_preserves_dimensions(self):
        workbook = Workbook();label = workbook.active;label.title = 'Etiqueta'
        label.append(['FORMATO AUXILIAR']);label.append(['=1+1'])
        sheet = workbook.create_sheet('Posiciones');sheet.append([]);sheet.append([])
        sheet.append(['Posición de ubicación','ALMACEN','PASILLO','RACK','NIVEL','CAJA'])
        sheet.append(['01-03-B-02-07',1,3,'B',2,7]);sheet.append(['01-03-B-03',1,3,'B',3,None])
        payload = self.excel(workbook)
        rows, source, _ = capacity.parse_master(payload)
        self.assertTrue(source.endswith(' / Posiciones'))
        self.assertEqual(rows[0], {'code':'01-03-B-02-07','warehouse':'01','aisle':'03','rack':'B','level':'02','bin':'07'})
        self.assertEqual(rows[1]['bin'], '')
        self.assertEqual(self.save(payload)['added'], 2)
        data = capacity.catalog(self.db, 'PICKER', search='01-03-B-02-07')['items'][0]
        self.assertEqual([data[key] for key in capacity.DIMENSIONS], ['01','03','02','07'])
        self.save({'text':'01-03-B-02-07'})
        self.assertEqual(capacity.catalog(self.db,'PICKER',search='01-03-B-02-07')['items'][0]['aisle'],'03')

    def test_explicit_sheet_and_ambiguous_sheets_never_choose_arbitrarily(self):
        workbook = Workbook();workbook.active.append(['Ubicacion']);workbook.active.append(['QA-001'])
        other = workbook.create_sheet('Otra');other.append(['Ubicacion']);other.append(['QA-002'])
        payload = self.excel(workbook)
        with self.assertRaisesRegex(ValueError, 'varias hojas'):
            capacity.parse_master(payload)
        self.assertEqual(capacity.parse_master({**payload,'sheet':'Otra'})[0],[{'code':'QA-002'}])

    def test_formulas_in_master_are_rejected_but_auxiliary_sheets_are_not_executed(self):
        workbook = Workbook();workbook.active.append(['Ubicacion','Rack']);workbook.active.append(['QA-001','=1+1'])
        with self.assertRaisesRegex(ValueError, 'fórmulas'):
            self.save(self.excel(workbook))
        self.assertEqual(capacity.catalog(self.db,'PICKER')['total'],0)

    def test_foreign_warehouse_rejected_and_numeric_zero_dimensions_preserved(self):
        with self.assertRaisesRegex(ValueError, 'almacén 01'):
            self.save(self.csv('Ubicacion,Almacen\nQA-001,02'))
        self.save(self.csv('Ubicacion,Almacen,Pasillo,Nivel,Caja\nQA-001,01,0,001,00'))
        row = capacity.catalog(self.db,'PICKER')['items'][0]
        self.assertEqual((row['aisle'],row['level'],row['bin']),('00','001','00'))

    def test_large_master_is_paginated_and_reimport_is_idempotent(self):
        payload = {'text':'\n'.join(f'QA-LARGE-{i:06}' for i in range(6001))}
        first = self.save(payload)
        self.assertEqual(first['added'],6001)
        self.assertEqual(self.save(payload)['unchanged'],6001)
        self.assertEqual(capacity.catalog(self.db,'PICKER',offset=6000)['total'],6001)
        self.assertEqual(len(capacity.catalog(self.db,'PICKER',offset=6000)['items']),1)
        self.assertEqual(self.db.execute('SELECT count(*) FROM warehouse_location_history').fetchone()[0],6001)

    def test_legacy_schema_migration_preserves_rows_and_is_repeatable(self):
        self.db.execute('DROP TABLE warehouse_locations')
        self.db.execute('CREATE TABLE warehouse_locations(code TEXT PRIMARY KEY,zone TEXT,rack TEXT,state TEXT,updated_at TEXT,updated_by TEXT)')
        self.db.execute("INSERT INTO warehouse_locations VALUES('QA-LEGACY','Z','R','BLOQUEADA','2026-01-01','qa')")
        capacity.init_schema(self.db);capacity.init_schema(self.db)
        row = capacity.catalog(self.db,'PICKER')['items'][0]
        self.assertEqual((row['warehouse'],row['aisle'],row['state']),('01','','BLOQUEADA'))

    def test_pagination_filtering_and_aim_are_bounded_and_literal(self):
        self.save({'text':'\n'.join('QA-'+str(i).zfill(3) for i in range(115))})
        for offset, expected in [(0,50),(50,50),(100,15)]:
            data = capacity.catalog(self.db, 'PICKER', offset=offset)
            self.assertEqual(len(data['items']), expected)
            self.assertEqual(data['total'], 115)
        self.assertEqual(capacity.catalog(self.db,'PICKER',search='%')['total'], 0)
        self.assertEqual(capacity.catalog(self.db,'PICKER',search=']d1QA-001')['items'][0]['code'], 'QA-001')
        self.assertEqual(capacity.catalog(self.db,'PICKER',search=']D1QA-001')['items'][0]['code'], 'QA-001')
        for args in [{'offset':'bad'},{'offset':-1},{'state':'INVALID'},{'evidence':'FREE'}]:
            with self.assertRaises(ValueError):
                capacity.catalog(self.db,'PICKER',**args)

    def test_read_and_write_roles_are_enforced(self):
        for role in ('PICKER','GUIADOR','ASISTENTE_RECEPCION'):
            self.assertEqual(capacity.catalog(self.db, role)['items'], [])
            with self.assertRaises(PermissionError):
                capacity.preview(self.db, {'text':'QA-1'}, role)
            with self.assertRaises(PermissionError):
                capacity.history(self.db, 'QA-1', role)
        with self.assertRaises(PermissionError):
            capacity.catalog(self.db, 'INVALID')

    def test_labels_only_registered_active_positions_copies_and_private_logo(self):
        self.save(self.csv('Ubicacion,Estado\nQA-001,HABILITADA\nQA-002,BLOQUEADA\nQA-003,BAJA'))
        with patch.object(capacity, 'label_brand', return_value={'brand_logo_url':'/api/reception-brand-logo?brand=TRITON&v='+'a'*64}):
            before = self.db.total_changes
            result = capacity.labels(self.db, {'codes':['QA-001','QA-002'],'copies':2}, 'PICKER', Path('qa.sqlite'))
            self.assertEqual(result['codes'], ['QA-001','QA-001','QA-002','QA-002'])
            self.assertTrue(result['logo'].startswith('/api/inventory/brand-logo?'))
            self.assertEqual(self.db.total_changes, before)
            for payload in [{'codes':[]},{'codes':['QA-001']*201},{'codes':['QA-003']},
                            {'codes':['MISSING']},{'codes':['QA-001','QA-001']},
                            {'codes':['QA-001'],'copies':True},{'codes':['QA-001'],'copies':11}]:
                with self.assertRaises(ValueError):
                    capacity.labels(self.db, payload, 'PICKER', Path('qa.sqlite'))
        with patch.object(capacity, 'label_brand', return_value={'brand_logo_url':''}):
            with self.assertRaisesRegex(ValueError, 'logo'):
                capacity.labels(self.db, {'codes':['QA-001']}, 'PICKER', Path('qa.sqlite'))


if __name__ == '__main__':
    unittest.main()
