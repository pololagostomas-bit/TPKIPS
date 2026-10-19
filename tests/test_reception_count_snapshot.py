"""Receipt work stays stable while synthetic documentary cuts are refreshed."""
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

from backend import app
from backend.services import reception
from backend.services.cloud_import_policy import documentary_only
from tests import test_reception_importation_customer as fixtures


class CountSnapshotTests(unittest.TestCase):
    load = fixtures.ImportationCustomerTests.load

    def setUp(self):
        fixtures.ImportationCustomerTests.setUp(self)
        setting = patch.dict("os.environ", {"TRITON_IMPORTACIONES_EMAILS": "qa@example.invalid"})
        setting.start()
        self.addCleanup(setting.stop)

    def workbook(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.title = "COMPRADOR QA"
        sheet.append(["GUIA DE IMPORTACION", "CODIGO", "PEDIDO TRITON", "OV",
                      "ORDEN DE COMPRA", "CANTIDAD SOLICITADA", "DESCRIPCION"])
        for guide, np_code, ip, ov, oc, quantity in rows:
            sheet.append([guide, np_code, ip, ov, oc, quantity, "Pieza QA"])
        return book

    def start(self, guide="QA-BL", *, verified=0, legacy=False):
        with app.db() as db:
            db.execute("""UPDATE reception_shipments SET app_status='REVISION SISTEMA',
                current_assistant='qa.worker',current_auxiliary='qa.worker',
                system_quantities_initialized=1 WHERE bl_awb=?""", (guide,))
            shipment = db.execute("SELECT * FROM reception_shipments WHERE bl_awb=?", (guide,)).fetchone()
            if legacy:
                db.execute("""UPDATE reception_lines SET received_qty=?,np_scanned_at='2026-10-19 08:00:00',
                    np_scanned_by='qa.worker',final_location='BIN-QA',count_location_saved=1
                    WHERE shipment_id=?""", (verified, shipment["id"]))
                return shipment["id"], None
            attention = reception._create_reception_attention(db, shipment, None, "qa.admin", "REVISION SISTEMA")
            db.execute("UPDATE reception_attentions SET system_quantities_initialized=1 WHERE id=?", (attention["id"],))
            db.execute("""UPDATE reception_attention_lines SET verified_qty=?,
                np_scanned_at='2026-10-19 08:00:00',np_scanned_by='qa.worker',
                final_location='BIN-QA',count_location_saved=1 WHERE attention_id=?""", (verified, attention["id"]))
            reception._sync_reception_line_totals(db, shipment["id"])
            return shipment["id"], attention["id"]

    def add_fr(self, shipment_id, ip="IP-1"):
        with app.db() as db:
            db.execute("""INSERT INTO reception_accounting_refs
                (shipment_id,source_key,ip_reference,fr_number,created_at,updated_at)
                VALUES (?, ?, ?, 'FR-QA', '2026-10-19', '2026-10-19')""",
                (shipment_id, "qa-fr-" + str(shipment_id), ip))

    def detail(self, shipment_id, attention_id=None, role="ADMINISTRADOR"):
        with app.db() as db:
            return reception.reception_detail(db, shipment_id, role, "qa.worker", attention_id)

    def test_pre_to_oc_keeps_plan_for_every_bl_and_role_while_source_updates(self):
        initial = [("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE-1", 5),
                   ("QA-BL-2", "QA-NP", "IP-2", "OV-2", "PRE-2", 7)]
        self.load(self.workbook(initial))
        contexts = [self.start(row[0], verified=3) for row in initial]
        with app.db() as db:
            before = [tuple(row) for row in db.execute("SELECT * FROM reception_attention_lines ORDER BY id")]
            ids = [row[0] for row in db.execute("SELECT id FROM reception_lines ORDER BY id")]
        updated = [(guide, np, ip, ov, "54321", qty + 4) for guide, np, ip, ov, _, qty in initial]
        result = self.load(self.workbook(list(reversed(updated))))
        self.assertEqual((result["lines_created"], result["lines_updated"]), (0, 2))
        with app.db() as db:
            self.assertEqual([row[0] for row in db.execute("SELECT id FROM reception_lines ORDER BY id")], ids)
            self.assertEqual([row[0] for row in db.execute("SELECT expected_qty FROM reception_lines ORDER BY id")], [9, 11])
            self.assertEqual([tuple(row) for row in db.execute("SELECT * FROM reception_attention_lines ORDER BY id")], before)
        for (sid, aid), original in zip(contexts, initial):
            for role in ("ADMINISTRADOR", "ASISTENTE_RECEPCION", "AUXILIAR_RECEPCION"):
                with self.subTest(guide=original[0], role=role):
                    line = self.detail(sid, aid, role)["lines"][0]
                    self.assertEqual((line["np_code"], line["expected_qty"], line["attention_planned_qty"]),
                                     ("QA-NP", original[-1], original[-1]))
                    self.assertEqual((line["oc_number"], line["attention_verified_qty"]), ("54321", 3))

    def source_only_line(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, aid = self.start(verified=3)
        self.add_fr(sid)
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9),
                                 ("QA-BL", "QA-NP", "IP-1", "OV-2", "65432", 7)]))
        with app.db() as db:
            extra = db.execute("SELECT id FROM reception_lines WHERE ov_number='OV-2'").fetchone()[0]
        return sid, aid, extra

    def test_new_documentary_line_is_not_advertised_as_current_fr_work(self):
        sid, aid, extra = self.source_only_line()
        for role in ("ADMINISTRADOR", "ASISTENTE_RECEPCION", "AUXILIAR_RECEPCION"):
            with self.subTest(role=role):
                detail = self.detail(sid, aid, role)
                line = next(line for line in detail["lines"] if line["id"] == extra)
                self.assertFalse(line["in_selected_attention"])
                self.assertEqual((line["attention_planned_qty"], line["attention_verified_qty"],
                                  line.get("attention_remaining_qty")), (0, 0, 0))
                linked = detail["accounting_refs"][0]["linked_lines"]
                self.assertEqual([(row["ov_number"], row["expected_qty"]) for row in linked], [("OV-1", 5)])

    def test_foreign_attention_line_is_rejected_even_without_scan_flag(self):
        sid, aid, extra = self.source_only_line()
        for role in ("ADMINISTRADOR", "ASISTENTE_RECEPCION", "AUXILIAR_RECEPCION"):
            for scan in (False, True):
                with self.subTest(role=role, scan=scan), app.db() as db:
                    before = list(db.iterdump())
                    with self.assertRaisesRegex(ValueError, "no pertenece a esta atenci"):
                        reception.update_reception_line_quantity(db, sid, extra,
                            {"attention_id": aid, "received_qty": 7, "scan_np": scan}, "qa.worker", role)
                    self.assertEqual(list(db.iterdump()), before)

    def test_updated_source_quantity_does_not_send_false_discrepancy(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, aid = self.start()
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        with app.db() as db:
            line_id = db.execute("SELECT id FROM reception_lines").fetchone()[0]
            before = db.execute("SELECT COUNT(*) FROM reception_notifications").fetchone()[0]
            reception.update_reception_line_quantity(db, sid, line_id,
                {"attention_id": aid, "received_qty": 5}, "qa.worker", "ASISTENTE_RECEPCION")
            self.assertEqual(db.execute("SELECT COUNT(*) FROM reception_notifications").fetchone()[0], before)

    def test_real_discrepancy_uses_attention_plan_not_latest_source(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, aid = self.start()
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        with app.db() as db:
            line_id = db.execute("SELECT id FROM reception_lines").fetchone()[0]
            reception.update_reception_line_quantity(db, sid, line_id,
                {"attention_id": aid, "received_qty": 4}, "qa.worker", "ASISTENTE_RECEPCION")
            notification = db.execute("SELECT * FROM reception_notifications ORDER BY id DESC LIMIT 1").fetchone()
            self.assertIsNotNone(notification)
            self.assertIn("Esperado: 5", notification["body"])
            self.assertNotIn("Esperado: 9", notification["body"])

    def test_later_arrival_uses_latest_source_balance_without_changing_closed_plan(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, first = self.start(verified=5)
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 8)]))
        with app.db() as db:
            db.execute("UPDATE reception_attentions SET app_status='CERRADO' WHERE id=?", (first,))
            shipment = db.execute("SELECT * FROM reception_shipments WHERE id=?", (sid,)).fetchone()
            later = reception._create_reception_attention(db, shipment, None, "qa.admin", "REVISION SISTEMA")["id"]
        self.assertEqual(self.detail(sid, first)["lines"][0]["expected_qty"], 5)
        self.assertEqual(self.detail(sid, later)["lines"][0]["expected_qty"], 3)

    def test_legacy_count_is_preserved_before_refreshing_its_source(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, _ = self.start(verified=3, legacy=True)
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        line = self.detail(sid)["lines"][0]
        self.assertEqual(line["expected_qty"], 5)
        self.assertEqual(line["attention_verified_qty"], 3)
        self.assertEqual((line["np_scanned_by"], line["final_location"]), ("qa.worker", "BIN-QA"))
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        with app.db() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM reception_attentions").fetchone()[0], 1)

    def test_same_np_with_separate_ovs_or_no_ov_keeps_each_plan(self):
        rows = [("QA-BL", "QA-NP", "IP-1", ov, "PRE-" + str(index), index + 2)
                for index, ov in enumerate(("OV-1", "OV-2", ""))]
        self.load(self.workbook(rows))
        sid, aid = self.start()
        before = {line["ov_number"]: (line["id"], line["expected_qty"])
                  for line in self.detail(sid, aid)["lines"]}
        self.load(self.workbook([(guide, np, ip, ov, "54321", qty + 4)
                                 for guide, np, ip, ov, _, qty in reversed(rows)]))
        after = self.detail(sid, aid)["lines"]
        self.assertEqual({line["ov_number"]: (line["id"], line["expected_qty"]) for line in after}, before)
        self.assertTrue(all(line["oc_number"] == "54321" for line in after))

    def test_bytes_and_local_file_imports_share_the_snapshot_rule(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, aid = self.start()
        for quantity, use_path in ((9, False), (11, True)):
            book = self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", quantity)])
            self.addCleanup(book.close)
            with app.db() as db:
                if use_path:
                    path = Path(self.scratch.name) / "qa.xlsx"
                    book.save(path)
                    reception.import_reception_excel_path(db, path)
                else:
                    content = io.BytesIO()
                    book.save(content)
                    reception.import_reception_excel_bytes(db, content.getvalue(), "qa.xlsx", "qa.admin", "ADMINISTRADOR")
                self.assertEqual(db.execute("SELECT expected_qty FROM reception_lines").fetchone()[0], quantity)
            self.assertEqual(self.detail(sid, aid)["lines"][0]["expected_qty"], 5)

    def test_legacy_quantities_without_new_scan_markers_render_as_recorded(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, _ = self.start(verified=3, legacy=True)
        with app.db() as db:
            db.execute("""UPDATE reception_lines SET np_scanned_at=NULL,np_scanned_by=NULL,
                count_location_saved=0,final_location=NULL""")
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        line = self.detail(sid)["lines"][0]
        node = shutil.which("node")
        if not node:
            self.skipTest("Node is required for the real count renderer")
        source = (Path(__file__).resolve().parents[1] / "frontend/templates/reception.html").read_text(encoding="utf-8")
        renderer = re.search(r"function physicalLineCells\(line,status\)\{[^\n]+\}", source).group(0)
        for verified in (3, 0):
            with self.subTest(verified=verified):
                value = dict(line, attention_verified_qty=verified)
                result = subprocess.run([node, "-e", "const esc=String;" + renderer +
                    ";console.log(physicalLineCells(" + json.dumps(value) + ",'REVISION SISTEMA'));"],
                    capture_output=True, text=True, check=True)
                self.assertIn(f'value="{verified}"', result.stdout)

    def test_legacy_transfer_confirmation_is_not_reset_by_snapshot_capture(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, _ = self.start(verified=5, legacy=True)
        with app.db() as db:
            db.execute("""UPDATE reception_shipments SET app_status='SOLICITUD TRANSFERENCIA',
                transfer_assistant_checked=1 WHERE id=?""", (sid,))
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        self.assertEqual(self.detail(sid)["transfer_assistant_checked"], 1)

    def report_row(self, content):
        book = load_workbook(io.BytesIO(content), data_only=True)
        self.addCleanup(book.close)
        rows = list(book["Packing List SAP"].iter_rows(values_only=True))
        return dict(zip(rows[0], rows[1]))

    def test_review_export_uses_count_plan_instead_of_later_documentary_quantity(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, _ = self.start(verified=5)
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 9)]))
        with app.db() as db:
            row = self.report_row(reception.reception_operational_report_bytes(db, sid, "ADMINISTRADOR", "qa.admin"))
        self.assertEqual((row["Esperado"], row["Verificado"], row["Diferencia"], row["Resultado"]), (5, 5, 0, "CONFORME"))

    def report_request(self, sid, attention_id):
        from backend import wsgi
        captured = {}
        env = {"PATH_INFO": f"/api/receptions/{sid}/system-report", "QUERY_STRING": f"attention_id={attention_id}",
               "REQUEST_METHOD": "GET", "wsgi.input": io.BytesIO(), "CONTENT_LENGTH": "0",
               "HTTP_X_ROLE": "ADMINISTRADOR", "HTTP_X_USER": "qa.admin"}
        with patch.dict("os.environ", {"TRITON_AUTH_MODE": "demo"}):
            content = b"".join(wsgi.application(env, lambda status, headers: captured.update(status=status)))
        return int(captured["status"][:3]), content

    def test_review_export_http_keeps_the_explicitly_selected_arrival(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, first = self.start(verified=5)
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "54321", 8)]))
        with app.db() as db:
            db.execute("UPDATE reception_attentions SET app_status='CERRADO' WHERE id=?", (first,))
            shipment = db.execute("SELECT * FROM reception_shipments WHERE id=?", (sid,)).fetchone()
            later = reception._create_reception_attention(db, shipment, None, "qa.admin", "REVISION SISTEMA")["id"]
            db.execute("UPDATE reception_attention_lines SET verified_qty=2 WHERE attention_id=?", (later,))
            reception._sync_reception_line_totals(db, sid)
        for aid, expected, verified in ((first, 5, 5), (later, 3, 2)):
            with self.subTest(attention_id=aid):
                status, content = self.report_request(sid, aid)
                self.assertEqual(status, 200)
                row = self.report_row(content)
                self.assertEqual((row["Esperado"], row["Verificado"]), (expected, verified))

    def test_review_export_http_rejects_a_foreign_attention(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, _ = self.start()
        status, _ = self.report_request(sid, 99999)
        self.assertEqual(status, 400)

    def test_download_button_passes_the_selected_attention_in_its_get_url(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node is required for the real download handler")
        source = (Path(__file__).resolve().parents[1] / "frontend/templates/reception.html").read_text(encoding="utf-8")
        handler = re.search(r"async function downloadSystemReport\(\)\{[^\n]+\}", source).group(0)
        program = """const selected={bl_awb:'QA-BL'},selectedAttentionId=7;
            let requested='';const notify=()=>{};const setTimeout=()=>{};
            const shipmentUrl=suffix=>'/api/receptions/11'+suffix;
            const document={body:{append(){}},createElement:()=>({click(){},remove(){}})};
            URL.createObjectURL=()=> 'blob:qa';
            const api=async url=>{requested=url;return {ok:true,blob:async()=>new Blob(['QA'])}};
        """ + handler + ";downloadSystemReport().then(()=>console.log(requested));"
        result = subprocess.run([node, "-e", program], capture_output=True, text=True, check=True)
        self.assertIn("/api/receptions/11/system-report?attention_id=7", result.stdout)

    def test_fr_em_import_remains_documentary_and_does_not_replace_plan(self):
        self.load(self.workbook([("QA-BL", "QA-NP", "IP-1", "OV-1", "PRE", 5)]))
        sid, aid = self.start(verified=3)
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = reception.ACCOUNTING_SHEET_NAME
        sheet.append(["FACTURA RESERVA", "ENTRADA DE MERCANCIA", "FECHA DE RECEPCION", "FECHA DE EM", "IP", "AWB"])
        sheet.append(["FR-QA", "EM-QA", "2026-10-19", "2026-10-19", "IP-1", "QA-BL"])
        with app.db() as db:
            context = documentary_only.set(True)
            try:
                reception.import_reception_accounting_workbook(db, book, "qa.xlsx", "qa.admin", "ADMINISTRADOR")
            finally:
                documentary_only.reset(context)
        detail = self.detail(sid, aid)
        self.assertEqual((detail["app_status"], detail["lines"][0]["expected_qty"]), ("REVISION SISTEMA", 5))
        self.assertEqual((detail["accounting_refs"][0]["fr_number"], detail["accounting_refs"][0]["em_numbers"]), ("FR-QA", ["EM-QA"]))


if __name__ == "__main__":
    unittest.main()
