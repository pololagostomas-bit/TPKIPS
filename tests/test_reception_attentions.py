import sqlite3
import unittest
from io import BytesIO
from uuid import uuid4
from pathlib import Path

import app
import backend.services.reception as reception
from openpyxl import load_workbook


class ReceptionAttentionTests(unittest.TestCase):
    def setUp(self):
        self.db_path = Path(__file__).parent / f".wms-reception-attentions-{uuid4().hex}.db"
        self.original_db_path = app.DB_PATH
        app.DB_PATH = self.db_path
        app.init_db()
        self.connection = sqlite3.connect(self.db_path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute(
            """INSERT INTO reception_shipments
               (bl_awb, transport_type, supplier, expected_packages, app_status,
                current_assistant, current_auxiliary, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            ("BL-ATT-001", "AEREO", "Proveedor prueba", 20, "PROGRAMADO",
             "asistente.test", "auxiliar.test", reception.reception_now(), reception.reception_now()),
        )
        self.shipment_id = self.connection.execute(
            "SELECT last_insert_rowid()"
        ).fetchone()[0]
        self.connection.execute(
            """INSERT INTO reception_lines
               (shipment_id, np_code, description, expected_qty, received_qty)
               VALUES (?, ?, ?, ?, ?)""",
            (self.shipment_id, "NP-001", "Repuesto de prueba", 20, 0),
        )
        self.connection.commit()

    def tearDown(self):
        self.connection.close()
        self.db_path.unlink(missing_ok=True)
        app.DB_PATH = self.original_db_path

    def test_new_arrival_after_closed_attention_uses_remaining_quantity(self):
        first = reception.add_physical_receipt(
            self.connection, self.shipment_id,
            {"received_packages": 2, "location": "ZONA-A-01"},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.assertEqual(first["active_attention_sequence"], 1)
        line_id = first["lines"][0]["id"]

        reception.update_reception_references(
            self.connection, self.shipment_id,
            {"fr_number": "FR-001"}, "admin.test", "ADMINISTRADOR",
        )
        reception.change_reception_status(
            self.connection, self.shipment_id,
            {"status": "REVISION SISTEMA", "require_location": True},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        reception.update_reception_line_quantity(
            self.connection, self.shipment_id, line_id,
            {"received_qty": 6}, "asistente.test", "ASISTENTE_RECEPCION",
        )

        self.connection.execute(
            """UPDATE reception_shipments
               SET app_status = 'SOLICITUD TRANSFERENCIA', transfer_assistant_checked = 1
               WHERE id = ?""", (self.shipment_id,)
        )
        self.connection.execute(
            """UPDATE reception_attentions
               SET app_status = 'SOLICITUD TRANSFERENCIA', transfer_assistant_checked = 1
               WHERE shipment_id = ? AND sequence_no = 1""", (self.shipment_id,)
        )
        self.connection.commit()
        reception.change_reception_status(
            self.connection, self.shipment_id,
            {"status": "CERRADO"}, "asistente.test", "ASISTENTE_RECEPCION",
        )

        second = reception.add_physical_receipt(
            self.connection, self.shipment_id,
            {"received_packages": 1, "location": "ZONA-B-02"},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.assertEqual(second["attention_count"], 2)
        self.assertEqual(second["active_attention_sequence"], 2)
        self.assertEqual([a["label"] for a in second["attentions"]], ["Atención 1/2", "Atención 2/2"])
        self.assertEqual(second["attentions"][1]["lines"][0]["planned_qty"], 14)
        self.assertEqual(second["attentions"][0]["lines"][0]["verified_qty"], 6)
        self.assertEqual(second["receipts"][1]["location_text"], "ZONA-B-02")

        queue = reception.list_receptions(self.connection, "BL-ATT-001", "ADMINISTRADOR")
        self.assertEqual([row["attention_label"] for row in queue], ["Atención 1/2", "Atención 2/2"])
        self.assertEqual([row["attention_sequence"] for row in queue], [1, 2])
        self.assertEqual([row["app_status"] for row in queue], ["CERRADO", "ARRIBADO"])

        # La nueva atención debe inicializarse con su propio esperado aunque
        # la BL ya haya inicializado las cantidades de la atención anterior.
        second_review = reception.change_reception_status(
            self.connection, self.shipment_id,
            {"status": "REVISION SISTEMA", "require_location": True},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.assertEqual(second_review["lines"][0]["received_qty"], 20)
        self.assertEqual(second_review["attentions"][1]["lines"][0]["verified_qty"], 14)

        # La corrección administrativa de una ubicación histórica no debe
        # depender de que la BL siga en ARRIBADO ni cambiar otra atención.
        self.connection.execute(
            "UPDATE reception_shipments SET app_status = 'REVISION SISTEMA' WHERE id = ?",
            (self.shipment_id,),
        )
        self.connection.execute(
            "UPDATE reception_attentions SET app_status = 'REVISION SISTEMA' "
            "WHERE shipment_id = ? AND sequence_no = 2",
            (self.shipment_id,),
        )
        self.connection.commit()
        corrected = reception.update_reception_location(
            self.connection,
            self.shipment_id,
            {"location": "ZONA-A-CORREGIDA", "receipt_id": 1},
            "admin.test",
            "ADMINISTRADOR",
        )
        self.assertEqual(corrected["receipts"][0]["location_text"], "ZONA-A-CORREGIDA")
        self.assertEqual(corrected["receipts"][1]["location_text"], "ZONA-B-02")
        with self.assertRaises(PermissionError):
            reception.update_reception_location(
                self.connection,
                self.shipment_id,
                {"location": "NO-AUTORIZADO", "receipt_id": 1},
                "asistente.test",
                "ASISTENTE_RECEPCION",
            )

    def test_location_is_required_only_when_ui_requests_system_review(self):
        reception.add_physical_receipt(
            self.connection, self.shipment_id,
            {"received_packages": 1}, "asistente.test", "ASISTENTE_RECEPCION",
        )
        with self.assertRaises(PermissionError):
            reception.change_reception_status(
                self.connection, self.shipment_id,
                {"status": "REVISION SISTEMA", "require_location": True},
                "asistente.test", "ASISTENTE_RECEPCION",
            )
        reception.update_reception_location(
            self.connection, self.shipment_id,
            {"location": "RACK-01", "receipt_id": 1},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        detail = reception.change_reception_status(
            self.connection, self.shipment_id,
            {"status": "REVISION SISTEMA", "require_location": True},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.assertEqual(detail["app_status"], "REVISION SISTEMA")

    def test_review_accepts_excess_and_report_labels_difference(self):
        reception.add_physical_receipt(
            self.connection, self.shipment_id,
            {"received_packages": 1, "location": "ZONA-A-01"},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        reception.change_reception_status(
            self.connection, self.shipment_id,
            {"status": "REVISION SISTEMA", "require_location": True},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        line_id = self.connection.execute(
            "SELECT id FROM reception_lines WHERE shipment_id = ?", (self.shipment_id,)
        ).fetchone()[0]

        detail = reception.update_reception_line_quantity(
            self.connection, self.shipment_id, line_id,
            {"received_qty": 25, "reason": "Llegaron 5 unidades adicionales"},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.assertEqual(detail["lines"][0]["received_qty"], 25)

        report = reception.reception_operational_report_bytes(
            self.connection, self.shipment_id, "ADMINISTRADOR", "admin.test"
        )
        workbook = load_workbook(BytesIO(report), data_only=True)
        sheet = workbook["Packing List SAP"]
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0]
        values = rows[1]
        self.assertIn("Resultado", header)
        self.assertEqual(values[header.index("Esperado")], 20)
        self.assertEqual(values[header.index("Verificado")], 25)
        self.assertEqual(values[header.index("Diferencia")], 5)
        self.assertEqual(values[header.index("Resultado")], "EXCEDENTE")

    def test_validation_has_one_default_location_check_and_persists_comment(self):
        reception.add_physical_receipt(
            self.connection, self.shipment_id,
            {"received_packages": 1, "location": "ZONA-A-01"},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.connection.execute(
            "UPDATE reception_shipments SET app_status = 'VALIDACION' WHERE id = ?",
            (self.shipment_id,),
        )
        self.connection.commit()

        detail = reception.reception_detail(
            self.connection, self.shipment_id,
            "ASISTENTE_RECEPCION", "asistente.test",
        )
        line = detail["lines"][0]
        self.assertEqual(line["validation_location_checked"], 1)
        self.assertEqual(line["validation_sap_checked"], 0)

        saved = reception.update_reception_validation(
            self.connection, self.shipment_id,
            {"checks": [{
                "line_id": line["id"],
                "location_checked": True,
                "comment": "Ubicación confirmada en rack A-01",
            }]},
            "asistente.test", "ASISTENTE_RECEPCION",
        )
        self.assertEqual(saved["lines"][0]["validation_comment"], "Ubicación confirmada en rack A-01")


if __name__ == "__main__":
    unittest.main()
