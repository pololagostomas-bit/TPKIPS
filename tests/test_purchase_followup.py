"""Date scopes and actual workbook/mail bytes on synthetic private history."""
import base64
import io
import json
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from backend.services import purchase_alerts as alerts
from tests import test_purchase_report_dates as date_fixtures


class PurchaseFollowupTests(unittest.TestCase):
    setUp = date_fixtures.PurchaseReportDateTests.setUp
    tearDown = date_fixtures.PurchaseReportDateTests.tearDown
    set_rows = date_fixtures.PurchaseReportDateTests.set_rows
    make_snapshot = date_fixtures.PurchaseReportDateTests.make_snapshot

    def observe(self, rows, offset=0):
        self.snapshot = self.make_snapshot(rows)
        return alerts.run(force=True, now=self.now + offset)

    def test_mail_cutoff_is_inclusive_and_newest_request_first_within_level(self):
        result = self.observe([
            ['PRE1', 'IP-OLD', 'OV1', '30/09/2026'],
            ['PRE2', 'IP-FIRST', 'OV2', '01/10/2026'],
            ['PRE3', 'IP-NEWEST', 'OV3', '24/10/2026']])
        body = result['preview']['html']
        self.assertNotIn('IP-OLD', body)
        self.assertIn('IP-FIRST', body)
        self.assertLess(body.index('IP-NEWEST'), body.index('IP-FIRST'))
        self.assertEqual(len(result['preview']['case_keys']), 2)

    def test_date_scope_does_not_close_or_delete_older_pending_episode(self):
        self.observe([['PRE1', 'IP-OLD', 'OV1', '30/09/2026']])
        self.assertEqual(alerts.case_history(period='recent')['total'], 0)
        old = alerts.case_history(period='historical')['items'][0]
        self.assertEqual(old['status'], 'PENDIENTE')
        self.assertIsNone(old['resolved_at'])

    def test_scope_and_request_sort_apply_before_pagination(self):
        self.observe([['PRE1', 'IP-OLD', 'OV1', '30/09/2026'],
                      ['PRE2', 'IP-FIRST', 'OV2', '01/10/2026'],
                      ['PRE3', 'IP-NEWEST', 'OV3', '24/10/2026']])
        result = alerts.case_history(period='recent', limit=1)
        self.assertEqual(result['total'], 2)
        self.assertEqual(result['items'][0]['ip'], 'IP-NEWEST')
        self.assertTrue(result['has_more'])
        self.assertEqual(alerts.case_history(period='recent', limit=1, offset=1)['items'][0]['ip'], 'IP-FIRST')
        with self.assertRaises(ValueError):
            alerts.case_history(period='invented')

    def test_period_and_severity_predicates_agree_for_invalid_and_legacy_dates(self):
        self.snapshot = self.make_snapshot([
            ['12345', 'IP-GRAVE', 'COT1', '01/10/2026'],
            ['PRE1', 'IP-OLD', 'OV1', '30/09/2026'],
            ['PRE2', 'IP-MISSING', 'OV2', None]])
        for value in (None, '', '0000-01-01', '1899-12-31', '2026-02-30', '20261001', 'invalid'):
            with self.subTest(value=value):
                next(case for case in self.snapshot['cases'] if case['ip'] == 'IP-MISSING')['sources'][0]['request_date'] = value
                alerts.run(force=True, now=self.now)
                self.assertEqual(alerts.case_history(period='recent', severity='GRAVE')['total'], 1)
                self.assertEqual(alerts.case_history(period='historical', severity='REGULARIZAR')['total'], 1)
                self.assertEqual(alerts.case_history(period='undated', severity='REGULARIZAR')['total'], 1)

    def test_workbook_keeps_resolved_recent_and_old_episodes_with_frozen_clock(self):
        self.observe([['PRE1', 'IP-OLD', 'OV1', '30/09/2026'],
                      ['PRE2', 'IP-RECENT', 'OV2', '01/10/2026']])
        self.observe([['12345', 'IP-OLD', 'OV1', '30/09/2026'],
                      ['12345', 'IP-RECENT', 'OV2', '01/10/2026']], 3600)
        book = load_workbook(io.BytesIO(alerts.history_xlsx(now=self.now + 7200)), data_only=True)
        try:
            self.assertEqual(book.sheetnames, ['Recientes', 'Historico'])
            for sheet, ip in [('Recientes', 'IP-RECENT'), ('Historico', 'IP-OLD')]:
                rows = list(book[sheet].values)
                row = dict(zip(rows[0], rows[1]))
                self.assertEqual(row['IP / Pedido Triton'], ip)
                self.assertEqual(row['Estado'], 'RESUELTO')
                self.assertEqual(row['Horas calendario'], 1)
                self.assertTrue(row['Fin observado UTC'])
                self.assertEqual(book[sheet].freeze_panes, 'A2')
                self.assertTrue(book[sheet].auto_filter.ref)
        finally:
            book.close()

    def test_mixed_sources_are_partitioned_without_losing_episode_or_missing_date(self):
        result = self.observe([['PRE1', 'IP-MIXED', 'OV1', '30/09/2026'],
                               ['PRE1', 'IP-MIXED', 'OV1', '01/10/2026'],
                               ['PRE1', 'IP-MIXED', 'OV1', None]])
        self.assertNotIn('30/09/2026', result['preview']['html'])
        self.assertIn('Sin fecha', result['preview']['html'])
        ids = {alerts.case_history(period=scope)['items'][0]['id']
               for scope in ('recent', 'historical', 'undated')}
        self.assertEqual(len(ids), 1)
        book = load_workbook(io.BytesIO(alerts.history_xlsx()), data_only=True)
        self.assertEqual(book.sheetnames, ['Recientes', 'Historico', 'Sin fecha'])
        self.assertTrue(all(book[sheet].max_row == 2 for sheet in book.sheetnames))
        book.close()

    def test_sent_payload_has_real_excel_attachment_and_transport_keeps_it(self):
        self.observe([['PRE1', 'IP-RECENT', 'OV1', '01/10/2026']])
        result = alerts.run(force=True, send_now=True, now=self.now + 3600)
        self.assertTrue(result['sent'])
        payload = self.post_mock.call_args.args[1]
        attachment = payload['attachments'][0]
        self.assertEqual(attachment['@odata.type'], '#microsoft.graph.fileAttachment')
        book = load_workbook(io.BytesIO(base64.b64decode(attachment['contentBytes'])))
        self.assertEqual(book.sheetnames, ['Recientes', 'Historico'])
        book.close()
        with alerts.db() as connection:
            recorded = json.loads(connection.execute('SELECT payload FROM outbox').fetchone()[0])
        self.assertNotIn('contentBytes', recorded['attachments'][0])
        self.post.stop()
        with patch.object(alerts.requests, 'post', return_value=Mock(status_code=202)) as post:
            alerts._post_mail('synthetic-token', payload)
        self.assertEqual(post.call_args.kwargs['json']['message']['attachments'], payload['attachments'])

    def test_large_attachment_fails_before_mail_or_outbox_without_discarding_history(self):
        self.observe([['PRE1', 'IP-RECENT', 'OV1', '01/10/2026']])
        with patch.object(alerts, 'history_xlsx', return_value=b'x' * 2500001):
            with self.assertRaisesRegex(ValueError, 'limite del adjunto'):
                alerts.run(force=True, send_now=True, now=self.now + 3600)
        self.post_mock.assert_not_called()
        self.assertEqual(alerts.case_history()['total'], 1)
        with alerts.db() as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM outbox').fetchone()[0], 0)

    def test_daily_report_keeps_resolved_history_when_no_cases_remain_pending(self):
        self.observe([['PRE1', 'IP-RECENT', 'OV1', '01/10/2026']])
        alerts.run(now=self.now + 3600)
        self.snapshot = self.make_snapshot([['12345', 'IP-RECENT', 'OV1', '01/10/2026']])
        result = alerts.run(now=self.now + 7200)  # Scheduled 10:00 Lima.
        self.assertTrue(result['sent'])
        payload = self.post_mock.call_args.args[1]
        self.assertEqual(payload['case_keys'], [])
        book = load_workbook(io.BytesIO(base64.b64decode(payload['attachments'][0]['contentBytes'])))
        self.assertEqual(book['Recientes'].cell(2, 4).value, 'RESUELTO')
        book.close()

    def test_date_edit_moves_scope_without_changing_episode_clock(self):
        self.observe([['PRE1', 'IP1', 'OV1', '30/09/2026']])
        old = alerts.case_history()['items'][0]
        self.observe([['PRE1', 'IP1', 'OV1', '01/10/2026']], 3600)
        item = alerts.case_history(period='recent', now=self.now + 3600)['items'][0]
        self.assertEqual((item['id'], item['first_seen_at']), (old['id'], old['first_seen_at']))
        self.assertEqual(item['elapsed_hours'], 1)
        self.assertEqual(alerts.case_history(period='historical')['total'], 0)

    def test_historical_new_case_is_recorded_without_triggering_incremental_email(self):
        alerts.save({'daily_time': '11:00'}, 'admin.qa')
        self.observe([['PRE1', 'IP-RECENT', 'OV1', '01/10/2026']])
        alerts.run(now=self.now + 3600)
        calls = self.post_mock.call_count
        self.snapshot = self.make_snapshot([['PRE1', 'IP-RECENT', 'OV1', '01/10/2026'],
                                            ['PRE2', 'IP-OLD', 'OV2', '30/09/2026']])
        result = alerts.run(now=self.now + 7200)
        self.assertFalse(result['sent'])
        self.assertEqual(self.post_mock.call_count, calls)
        self.assertEqual(alerts.case_history(period='historical')['total'], 1)

    def test_excel_cells_are_text_not_formulas(self):
        self.snapshot = self.make_snapshot([['PRE1', 'IP1', '@OV', '01/10/2026']])
        self.snapshot['cases'][0]['ip'] = '=IP-UNTRUSTED'
        alerts.run(force=True, now=self.now)
        book = load_workbook(io.BytesIO(alerts.history_xlsx()))
        values = [cell for row in book['Recientes'] for cell in row]
        self.assertFalse(any(cell.data_type == 'f' for cell in values))
        self.assertTrue(any(cell.value == '=IP-UNTRUSTED' for cell in values))
        book.close()
