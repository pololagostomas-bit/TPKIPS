"""Real automatic import path, synthetic workbook/network, disposable database."""
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from backend import app
from backend.services import cloud_sync, reception
from tests import test_cloud_connection as fixtures


class CloudQuotationSyncTests(unittest.TestCase):
    setUp = fixtures.CloudIntegrationTests.setUp
    tearDown = fixtures.CloudIntegrationTests.tearDown

    def source(self, ov, minute):
        book = Workbook()
        sheet = book.active
        sheet.title = 'COMPRADOR QA'
        sheet.append(['GUIA DE IMPORTACION','CODIGO','PEDIDO TRITON','OV','ORDEN DE COMPRA',
                      'CANTIDAD SOLICITADA','MODO DE TRANSPORTE','Cliente'])
        sheet.append(['QA-BL','QA-NP','QA-IP',ov,'54321',5,'AEREO','Cliente '+ov])
        content = BytesIO()
        book.save(content)
        book.close()
        return {'source_type':'importation','filename':'qa.xlsx','content':content.getvalue(),
                'modified_at':f'2026-10-23T10:{minute:02d}:00Z'}

    def test_automatic_cycles_refresh_quote_then_real_ov_without_resetting_attention(self):
        source = self.source('COT 001234',0)
        with patch.dict('os.environ', {'WMS_CLOUD_AUTO_SYNC':'1','WMS_CLOUD_INTERVAL_SECONDS':'300',
                                       'TRITON_MAIL_ENABLED':'0'}), \
             patch.object(cloud_sync.cloud_auth,'status',return_value={'authorized':True}), \
             patch.object(cloud_sync.cloud_excel,'download_cloud_excels',side_effect=lambda **kwargs:[source]):
            cloud_sync._worker_cycle(app)
            with app.db() as c:
                old = dict(c.execute('SELECT * FROM reception_lines').fetchone())
                c.execute("UPDATE reception_shipments SET app_status='REVISION SISTEMA',received_packages=1 WHERE id=?",(old['shipment_id'],))
                shipment = c.execute('SELECT * FROM reception_shipments WHERE id=?',(old['shipment_id'],)).fetchone()
                attention = reception._create_reception_attention(c,shipment,None,'admin.qa','REVISION SISTEMA')
                c.execute("""UPDATE reception_attention_lines SET verified_qty=3,np_quantity_saved=1,
                    np_scanned_at='2026-10-23T09:00:00',np_scanned_by='admin.qa',final_location='QA-BIN',
                    count_location_saved=1 WHERE attention_id=?""",(attention['id'],))
                physical = tuple(c.execute('SELECT * FROM reception_attention_lines WHERE attention_id=?',(attention['id'],)).fetchone())
            for ov, minute in [('COT 005678',5),('159001',10),('159001',10)]:
                source = self.source(ov,minute)
                cloud_sync._worker_cycle(app)
                with app.db() as c:
                    self.assertEqual([(r['id'],r['ov_number'],r['customer_name']) for r in c.execute('SELECT * FROM reception_lines')],
                                     [(old['id'],ov,'Cliente '+ov)])
                    self.assertEqual(tuple(c.execute('SELECT * FROM reception_attention_lines WHERE attention_id=?',(attention['id'],)).fetchone()),physical)
                    detail = reception.reception_detail(c,old['shipment_id'],attention_id=attention['id'])
                    self.assertEqual(detail['lines'][0]['ov_number'],ov)
                    self.assertEqual(detail['lines'][0]['attention_verified_qty'],3)
                    status = c.execute("SELECT state FROM cloud_sync_status WHERE source_type='importation'").fetchone()
                    self.assertIn(status['state'],('updated','unchanged'))
                    runtime = json.loads(c.execute('SELECT value FROM cloud_sync_runtime WHERE id=1').fetchone()[0])
                    self.assertEqual(runtime['worker_state'],'waiting')
                    data = reception.reception_operational_report_bytes(c,old['shipment_id'],attention_id=attention['id'])
                exported = load_workbook(BytesIO(data),data_only=True)
                try:
                    self.assertEqual([(r[4],r[5],r[6]) for r in exported['Packing List SAP'].iter_rows(min_row=2,values_only=True)],[(ov,5,3)])
                finally:
                    exported.close()
