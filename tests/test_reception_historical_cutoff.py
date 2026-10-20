"""Owner-directed historical assumptions never manufacture physical evidence."""
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from backend.services import reception
from backend.services.reception_historical_cutoff import (
    apply_historical_cutoff,
    preview_historical_cutoff,
)
from backend.services.reception_work import reception_work_report
from backend.services.traceability import init_traceability_schema


STAMP = '2026-10-19 12:00:00'
ADMIN = 'ADMINISTRADOR'
REASON = 'Propietario solicita asumir historicos por Excel desordenado'


class HistoricalCutoffTests(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(':memory:')
        self.db.row_factory = sqlite3.Row
        self.addCleanup(self.db.close)
        reception.init_reception_schema(self.db)
        init_traceability_schema(self.db)
        self.db.execute('''CREATE TABLE order_importation_refs
            (bl_awb TEXT,sap_ov TEXT,transport_type TEXT,ip_reference TEXT,oc_number TEXT)''')

    def shipment(self, day='2026-07-31', bl=None, status='PROGRAMADO'):
        bl = bl or 'QA-CUTOFF-' + str(self.db.execute(
            'SELECT COUNT(*) FROM reception_shipments').fetchone()[0])
        sid = self.db.execute('''INSERT INTO reception_shipments
            (bl_awb,scheduled_date,app_status,ip_reference,expected_packages,created_at,updated_at)
            VALUES (?,?,?,'IP-QA',5,?,?)''', (bl, day, status, STAMP, STAMP)).lastrowid
        self.db.execute('''INSERT INTO reception_lines
            (shipment_id,np_code,ip_reference,ov_number,expected_qty)
            VALUES (?,'NP-QA','IP-QA','OV-QA',9)''', (sid,))
        return sid

    def preview(self, cutoff='2026-07-31', role=ADMIN):
        return preview_historical_cutoff(self.db, cutoff=cutoff, role=role)

    def apply(self, preview=None, **changes):
        preview = preview if preview is not None else self.preview()
        data = {'cutoff': preview['cutoff'], 'token': preview['token'],
                'confirmed': True, 'reason': REASON, **changes}
        return apply_historical_cutoff(self.db, data, 'qa.admin', ADMIN)

    def rows(self, table):
        return [dict(row) for row in self.db.execute(f'SELECT * FROM {table} ORDER BY rowid')]

    def workbook(self, bl='QA-EXCEL', day='2026-07-31', confirmed=None):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = 'COMPRADOR QA'
        sheet.append(['GUIA DE IMPORTACION', 'CODIGO', 'PEDIDO TRITON', 'OV',
                      'CANTIDAD SOLICITADA', 'FECHA TRITON ESTIMADA', 'FECHA CONFIRMADA TRITON'])
        sheet.append([bl, 'NP-QA', 'IP-QA', 'OV-QA', 9, day, confirmed])
        return book

    def test_preview_includes_july_boundary_and_older_years_using_source_date_parser(self):
        eligible = [self.shipment(day) for day in
                    ('2026-07-31', '2026-07-31 23:59:59', '31/07/2026', '2025-12-31')]
        closed = self.shipment(status='CERRADO')
        missing = self.shipment(None)
        invalid = self.shipment('2026-02-30')
        future = self.shipment('2026-08-01')
        self.db.execute('PRAGMA query_only=ON')
        before = self.db.total_changes
        result = self.preview()
        self.assertEqual(result['eligible_ids'], eligible)
        self.assertEqual(result['excluded_ids'], [closed, missing, invalid, future])
        self.assertEqual(result['counts'], {
            'total': 8, 'eligible': 4, 'excluded': 4,
            'by_reason': {'ALREADY_CLOSED': 1, 'UNKNOWN_DATE': 1,
                          'INVALID_DATE': 1, 'AFTER_CUTOFF': 1}})
        self.assertEqual(result['date_field'], 'scheduled_date')
        self.assertEqual(result['eligible'][2]['scheduled_day'], '2026-07-31')
        self.assertEqual(self.preview()['token'], result['token'])
        self.assertEqual(self.db.total_changes, before)

    def test_cutoff_can_only_be_strict_iso_on_or_before_july_31(self):
        old = self.shipment('2026-06-30')
        self.shipment('2026-07-01')
        self.assertEqual(self.preview('2026-06-30')['eligible_ids'], [old])
        for cutoff in ('2026-08-01', '20260731', '2026-7-31', '31/07/2026',
                       '2026-07-31T00:00:00', ' 2026-07-31', '2026-02-30', '', None, 20260731):
            with self.subTest(cutoff=cutoff), self.assertRaises(ValueError):
                self.preview(cutoff)

    def test_apply_preserves_pending_accounting_and_every_quantity_and_physical_table(self):
        sid = self.shipment()
        self.db.execute('''INSERT INTO reception_accounting_refs
            (shipment_id,source_key,ip_reference,fr_number,created_at,updated_at)
            VALUES (?,'QA-REF','IP-QA','FR-QA',?,?)''', (sid, STAMP, STAMP))
        self.db.execute("UPDATE reception_shipments SET accounting_status='PENDIENTE EM',fr_number='FR-QA'")
        tables = ('reception_lines', 'reception_receipts', 'reception_attentions',
                  'reception_attention_lines', 'reception_truck_package_scans',
                  'reception_truck_bl_arrivals', 'inventory_lots',
                  'reception_accounting_refs', 'reception_accounting_ems',
                  'reception_transfer_requests', 'reception_notifications')
        preserved = {table: self.rows(table) for table in tables}
        original = self.rows('reception_shipments')[0]
        preview = self.preview()
        result = self.apply(preview)
        self.assertEqual(result['applied_ids'], [sid])
        self.assertEqual(result['applied_count'], 1)
        current = self.rows('reception_shipments')[0]
        self.assertEqual(current['app_status'], 'CERRADO')
        self.assertEqual(current['condition_status'], 'HISTORICO POR CORTE')
        self.assertEqual(current['accounting_status'], 'PENDIENTE EM')
        allowed = {'app_status', 'condition_status', 'updated_at',
                   'physical_initialized', 'system_quantities_initialized'}
        self.assertEqual({k: v for k, v in current.items() if k not in allowed},
                         {k: v for k, v in original.items() if k not in allowed})
        for table, rows in preserved.items():
            self.assertEqual(self.rows(table), rows, table)
        history = self.rows('reception_history')
        status_event = next(row for row in history if row['field_name'] == 'app_status')
        self.assertEqual((status_event['event_type'], status_event['old_value'],
                          status_event['new_value'], status_event['username']),
                         ('CIERRE HISTORICO POR CORTE', 'PROGRAMADO', 'CERRADO', 'qa.admin'))
        self.assertIn(REASON, status_event['reason'])
        audit = json.loads(next(row['new_value'] for row in history
                                if row['field_name'] == 'historical_cutoff'))
        self.assertEqual(audit['basis'], 'OWNER_INSTRUCTION_ASSUMPTION')
        self.assertFalse(audit['fr_em_sap_verified'])
        self.assertEqual(audit['cutoff'], '2026-07-31')
        self.assertEqual(audit['scheduled_date'], '2026-07-31')
        self.assertEqual(audit['preview_token'], preview['token'])

    def test_admin_and_explicit_confirmation_reason_and_identity_are_required(self):
        self.shipment()
        preview = self.preview()
        before = self.rows('reception_shipments')
        data = {'confirmed': True, 'reason': REASON, 'token': preview['token']}
        for role in ('ASISTENTE_RECEPCION', 'AUXILIAR_RECEPCION', 'PICKER', '', None):
            with self.subTest(role=role), self.assertRaises(PermissionError):
                self.preview(role=role)
            with self.subTest(role=role), self.assertRaises(PermissionError):
                apply_historical_cutoff(self.db, data, 'qa', role)
        for change in ({'confirmed': False}, {'confirmed': 'true'}, {'confirmed': 1},
                       {'reason': 'corto'}, {'reason': ''}, {'token': ''}, {'token': None}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                self.apply(preview, **change)
        with self.assertRaises(ValueError):
            apply_historical_cutoff(self.db, data, '  ', ADMIN)
        self.assertEqual(self.rows('reception_shipments'), before)
        self.assertEqual(self.rows('reception_history'), [])

    def test_active_or_unproven_physical_work_is_excluded_even_if_source_status_is_programmed(self):
        sid = self.shipment()
        cases = [
            "UPDATE reception_shipments SET received_packages=1",
            "UPDATE reception_shipments SET app_status='REVISION SISTEMA'",
            "UPDATE reception_shipments SET first_arrival_at='invalid physical date'",
            "UPDATE reception_shipments SET completed_arrival_at='2026-07-01'",
            "UPDATE reception_shipments SET bultos_closed=1",
            "UPDATE reception_shipments SET location_text='QA-ZONE'",
            "UPDATE reception_shipments SET transfer_assistant_checked=1",
            "UPDATE reception_lines SET received_qty=1",
            "UPDATE reception_lines SET np_scanned_at='2026-07-01'",
            "UPDATE reception_lines SET np_scanned_by='qa.worker'",
            "UPDATE reception_lines SET count_location_saved=1",
            "UPDATE reception_lines SET final_location='QA-01'",
            "UPDATE reception_lines SET validation_sap_checked=1",
            """INSERT INTO reception_attentions
                (shipment_id,sequence_no,app_status,created_at,updated_at)
                VALUES (1,1,'ARRIBADO','qa','qa')""",
            """INSERT INTO reception_attentions
                (shipment_id,sequence_no,app_status,condition_status,created_at,updated_at)
                VALUES (1,1,'CERRADO','ANULADA POR REVERSION','qa','qa')""",
            """INSERT INTO reception_receipts
                (shipment_id,sequence_no,received_packages,username,received_at)
                VALUES (1,1,0,'qa','qa')""",
            """INSERT INTO reception_truck_bl_arrivals
                (truck_guide,shipment_id,received_packages,arrival_key,username,arrived_at)
                VALUES ('QA-TRUCK',1,1,'QA-ARRIVAL','qa','qa')""",
            """INSERT INTO reception_truck_package_scans
                (truck_guide,shipment_id,package_code,normalized_code,scanned_by,scanned_at)
                VALUES ('QA-TRUCK',1,'QA-PACK','QA-PACK','qa','qa')""",
            """INSERT INTO reception_truck_bl_manifest
                (truck_guide,shipment_id,planned_packages,created_by,created_at,updated_at)
                VALUES ('QA-TRUCK',1,0,'qa','qa','qa')""",
            """INSERT INTO reception_truck_manifest
                (truck_guide,shipment_id,expected_packages,created_at)
                VALUES ('QA-TRUCK',1,1,'qa')""",
            "UPDATE reception_shipments SET truck_guide='QA-MISSING'",
            """INSERT INTO reception_transfer_requests
                (shipment_id,status,created_at) VALUES (1,'PENDING',1)""",
            """INSERT INTO inventory_lots
                (lot_code,reception_shipment_id,item_key,item_code,created_by,created_at,updated_at)
                VALUES ('QA-LOT',1,'NP-QA','NP-QA','qa','qa','qa')""",
        ]
        for sql in cases:
            with self.subTest(sql=sql):
                self.db.execute('SAVEPOINT qa_case')
                self.db.execute(sql)
                preview = self.preview()
                self.assertEqual(preview['eligible_ids'], [])
                self.assertEqual(preview['excluded_ids'], [sid])
                self.assertIn(preview['excluded'][0]['reason'],
                              ('PHYSICAL_WORK', 'TRUCK_WORK', 'INVENTORY', 'TRANSFER_WORK'))
                self.db.execute('ROLLBACK TO qa_case')
                self.db.execute('RELEASE qa_case')

    def test_legacy_excel_arrival_is_documentary_and_does_not_gain_a_received_date(self):
        sid = self.shipment(status='ARRIBADO')
        self.db.execute('UPDATE reception_shipments SET first_arrival_at=?', (STAMP,))
        reception._write_history(self.db, sid, 'IMPORTACION EXCEL', 'app_status',
                                 'PROGRAMADO', 'ARRIBADO', 'qa.sync', 'Legacy source availability')
        preview = self.preview()
        self.assertEqual(preview['eligible_ids'], [sid])
        self.apply(preview)
        self.assertEqual(self.rows('reception_shipments')[0]['first_arrival_at'], STAMP)
        item = reception_work_report(self.db, view='history', now=STAMP)['items'][0]
        self.assertEqual(item['app_status'], 'CERRADO')
        self.assertIsNone(item['work_started_at'])
        self.assertEqual(self.rows('reception_receipts'), [])

    def test_raw_arrived_master_without_physical_evidence_is_eligible_without_import_history(self):
        sid = self.shipment(status='ARRIBADO')
        preview = self.preview()
        self.assertEqual(preview['eligible_ids'], [sid])
        self.assertEqual(self.apply(preview)['applied_ids'], [sid])
        current = self.rows('reception_shipments')[0]
        self.assertEqual(current['app_status'], 'CERRADO')
        self.assertIsNone(current['first_arrival_at'])
        self.assertEqual(current['received_packages'], 0)
        self.assertEqual(current['accounting_status'], 'PENDIENTE CONTABILIDAD')
        self.assertEqual(self.rows('reception_receipts'), [])

    def test_automatic_sample_location_defaults_do_not_exclude_documentary_arrival(self):
        sid = self.shipment(status='ARRIBADO')
        reception._validation_sample_ids(self.db, sid)
        self.assertEqual(self.rows('reception_lines')[0]['validation_location_checked'], 1)
        self.assertEqual(self.rows('reception_shipments')[0]['validation_sample_json'], '[1]')
        self.assertEqual(self.preview()['eligible_ids'], [sid])

    def test_explicit_sample_validation_still_excludes_even_after_operator_unchecks_it(self):
        sid = self.shipment(status='ARRIBADO')
        reception._validation_sample_ids(self.db, sid)
        self.db.execute('UPDATE reception_lines SET validation_location_checked=0')
        reception._write_history(self.db, sid, 'VALIDACION', 'linea:1:muestra', '',
            '{"validation_location_checked": 0}', 'qa.worker', 'Control registrado por operador')
        self.assertEqual(self.preview()['eligible_ids'], [])
        self.assertEqual(self.preview()['excluded'][0]['reason'], 'PHYSICAL_WORK')

    def test_reason_contract_accepts_500_characters_and_rejects_501_without_mutation(self):
        self.shipment()
        preview = self.preview()
        with self.assertRaises(ValueError):
            self.apply(preview, reason='x' * 501)
        self.assertEqual(self.rows('reception_shipments')[0]['app_status'], 'PROGRAMADO')
        self.assertEqual(self.rows('reception_history'), [])
        self.assertEqual(self.apply(preview, reason='x' * 500)['applied_count'], 1)

    def test_stale_token_rejects_same_timestamp_document_and_physical_changes(self):
        sid = self.shipment()
        preview = self.preview()
        changes = [
            "UPDATE reception_shipments SET scheduled_date='2026-07-30'",
            "UPDATE reception_shipments SET supplier='changed'",
            "UPDATE reception_lines SET expected_qty=10",
            "UPDATE reception_lines SET received_qty=1",
            """INSERT INTO reception_accounting_refs
                (shipment_id,source_key,fr_number,created_at,updated_at)
                VALUES (1,'QA-NEW','FR-NEW','qa','qa')""",
            """INSERT INTO reception_history
                (shipment_id,event_type,field_name,username,created_at)
                VALUES (1,'MODIFICACION','qa-field','qa','qa')""",
        ]
        for sql in changes:
            with self.subTest(sql=sql):
                self.db.execute('SAVEPOINT qa_change')
                self.db.execute(sql)
                with self.assertRaisesRegex(ValueError, 'cambio|vigente|previa'):
                    self.apply(preview)
                self.assertNotEqual(self.rows('reception_shipments')[0]['app_status'], 'CERRADO')
                self.db.execute('ROLLBACK TO qa_change')
                self.db.execute('RELEASE qa_change')
        self.assertEqual(self.preview()['eligible_ids'], [sid])

    def test_token_binds_cutoff_and_entire_current_eligible_set(self):
        first = self.shipment('2026-06-01')
        preview = self.preview()
        with self.assertRaises(ValueError):
            self.apply(preview, cutoff='2026-06-30')
        second = self.shipment('2026-05-01')
        with self.assertRaises(ValueError):
            self.apply(preview)
        self.assertEqual(self.apply()['applied_ids'], [first, second])

    def test_changes_to_still_excluded_records_do_not_change_eligible_snapshot_token(self):
        self.shipment()
        excluded = self.shipment('2026-08-01')
        preview = self.preview()
        self.db.execute('UPDATE reception_shipments SET supplier=? WHERE id=?', ('changed', excluded))
        self.assertEqual(self.preview()['token'], preview['token'])
        self.assertEqual(self.apply(preview)['applied_count'], 1)

    def test_batch_failure_rolls_back_all_closures_and_audits_but_preserves_outer_work(self):
        self.db.executescript('''CREATE TRIGGER qa_block_second BEFORE UPDATE OF app_status
            ON reception_shipments WHEN NEW.id=2 AND NEW.app_status='CERRADO'
            BEGIN SELECT RAISE(ABORT,'QA batch failure'); END;''')
        self.shipment()
        self.shipment()
        self.db.execute("UPDATE reception_shipments SET supplier='outer work' WHERE id=1")
        before = self.rows('reception_shipments')
        with self.assertRaisesRegex(sqlite3.IntegrityError, 'QA batch failure'):
            self.apply()
        self.assertEqual(self.rows('reception_shipments'), before)
        self.assertEqual(self.rows('reception_history'), [])
        self.assertTrue(self.db.in_transaction)

    def test_fresh_retry_excludes_closed_ids_and_never_duplicates_audit(self):
        self.shipment()
        preview = self.preview()
        self.apply(preview)
        history = self.rows('reception_history')
        with self.assertRaises(ValueError):
            self.apply(preview)
        retry = self.preview()
        self.assertEqual(retry['eligible_ids'], [])
        self.assertEqual(retry['counts']['by_reason'], {'ALREADY_CLOSED': 1})
        self.assertEqual(self.apply(retry)['applied_count'], 0)
        self.assertEqual(self.rows('reception_history'), history)

    def test_excel_refresh_and_startup_keep_cutoff_closed_with_zero_physical_quantities(self):
        book = self.workbook()
        reception.import_reception_workbook(self.db, book, 'qa.xlsx', 'qa.sync', ADMIN)
        self.apply()
        book.active.cell(2, 6, '2026-09-01')
        reception.import_reception_workbook(self.db, book, 'qa-refresh.xlsx', 'qa.sync', ADMIN)
        accounting = Workbook()
        self.addCleanup(accounting.close)
        accounting.active.title = reception.ACCOUNTING_SHEET_NAME
        accounting.active.append(['AWB', 'IP', 'FACTURA RESERVA', 'ENTRADA DE MERCANCIA',
                                  'FECHA DE RECEPCION', 'FECHA DE EM'])
        accounting.active.append(['QA-EXCEL', 'IP-QA', 'FR-QA', '', '', ''])
        reception.import_reception_accounting_workbook(
            self.db, accounting, 'qa-accounting.xlsx', 'qa.sync', ADMIN)
        reception.init_reception_schema(self.db)
        shipment = self.rows('reception_shipments')[0]
        self.assertEqual((shipment['app_status'], shipment['condition_status'],
                          shipment['scheduled_date'], shipment['accounting_status']),
                         ('CERRADO', 'HISTORICO POR CORTE', '2026-09-01', 'PENDIENTE EM'))
        self.assertEqual(shipment['received_packages'], 0)
        self.assertIsNone(shipment['first_arrival_at'])
        self.assertIsNone(shipment['completed_arrival_at'])
        self.assertEqual(self.rows('reception_lines')[0]['received_qty'], 0)
        for table in ('reception_receipts', 'reception_attentions', 'inventory_lots'):
            self.assertEqual(self.rows(table), [], table)
        self.assertEqual(reception_work_report(self.db, view='arrival', now=STAMP)['items'], [])
        self.assertEqual(reception_work_report(self.db, now=STAMP)['items'], [])
        history = reception_work_report(self.db, view='history', now=STAMP)['items']
        self.assertEqual([row['work_status'] for row in history], ['CERRADO'])

    def test_source_confirmed_date_wins_and_later_imports_never_auto_close_old_dates(self):
        reception.import_reception_workbook(self.db, self.workbook(confirmed='2026-08-01'),
                                           'qa.xlsx', 'qa.sync', ADMIN)
        self.assertEqual(self.preview()['eligible_ids'], [])
        self.assertEqual(self.preview()['excluded'][0]['reason'], 'AFTER_CUTOFF')
        reception.import_reception_workbook(self.db, self.workbook(bl='QA-OLDER'),
                                           'qa2.xlsx', 'qa.sync', ADMIN)
        self.assertEqual([row['app_status'] for row in self.rows('reception_shipments')],
                         ['PROGRAMADO', 'PROGRAMADO'])

    def test_incomplete_attention_cannot_hide_behind_closed_master_condition(self):
        sid = self.shipment()
        self.db.execute('''INSERT INTO reception_attentions
            (shipment_id,sequence_no,app_status,condition_status,created_at,updated_at)
            VALUES (?,1,'REVISION SISTEMA','COMPLETADO',?,?)''', (sid, STAMP, STAMP))
        self.assertEqual(self.preview()['eligible_ids'], [])

    def test_physical_history_requires_completed_attentions_and_existing_closure_proof(self):
        sid = self.shipment(status='ARRIBADO')
        self.db.execute('UPDATE reception_shipments SET received_packages=5 WHERE id=?', (sid,))
        self.db.execute('UPDATE reception_lines SET received_qty=9 WHERE shipment_id=?', (sid,))
        receipt = self.db.execute('''INSERT INTO reception_receipts
            (shipment_id,sequence_no,received_packages,username,received_at)
            VALUES (?,1,5,'qa',?)''', (sid, STAMP)).lastrowid
        attention = self.db.execute('''INSERT INTO reception_attentions
            (shipment_id,receipt_id,sequence_no,app_status,condition_status,
             created_at,updated_at,completed_at)
            VALUES (?,?,1,'CERRADO','COMPLETADO',?,?,?)''',
            (sid, receipt, STAMP, STAMP, STAMP)).lastrowid
        self.db.execute('''INSERT INTO reception_attention_lines
            (attention_id,reception_line_id,planned_qty,verified_qty) VALUES (?,1,9,9)''', (attention,))
        self.assertEqual(self.preview()['eligible_ids'], [])
        reference = self.db.execute('''INSERT INTO reception_accounting_refs
            (shipment_id,source_key,ip_reference,fr_number,em_number,created_at,updated_at)
            VALUES (?,'QA-REF','IP-QA','FR-QA','EM-QA',?,?)''', (sid, STAMP, STAMP)).lastrowid
        self.db.execute('''INSERT INTO reception_accounting_ems
            (accounting_ref_id,ip_reference_key,em_number,username,created_at)
            VALUES (?,'IPQA','EM-QA','qa',?)''', (reference, STAMP))
        self.assertEqual(self.preview()['eligible_ids'], [sid])
        self.db.execute('UPDATE reception_attentions SET completed_at=NULL')
        self.assertEqual(self.preview()['eligible_ids'], [])

    def test_inventory_linked_only_through_line_is_also_excluded(self):
        self.shipment()
        self.db.execute('''INSERT INTO inventory_lots
            (lot_code,reception_line_id,item_key,item_code,created_by,created_at,updated_at)
            VALUES ('QA-LOT',1,'NP-QA','NP-QA','qa',?,?)''', (STAMP, STAMP))
        self.assertEqual(self.preview()['excluded'][0]['reason'], 'INVENTORY')

    def test_existing_em_child_changes_invalidate_token_without_updating_parent_timestamp(self):
        sid = self.shipment()
        ref = self.db.execute('''INSERT INTO reception_accounting_refs
            (shipment_id,source_key,ip_reference,fr_number,created_at,updated_at)
            VALUES (?,'QA-REF','IP-QA','FR-QA',?,?)''', (sid, STAMP, STAMP)).lastrowid
        preview = self.preview()
        self.db.execute('''INSERT INTO reception_accounting_ems
            (accounting_ref_id,em_number,username,created_at) VALUES (?,'EM-QA','qa',?)''', (ref, STAMP))
        with self.assertRaises(ValueError):
            self.apply(preview)
        self.assertEqual(self.rows('reception_shipments')[0]['app_status'], 'PROGRAMADO')

    def test_success_in_outer_transaction_remains_rollbackable_by_caller(self):
        self.shipment()
        self.apply()
        self.assertTrue(self.db.in_transaction)
        self.db.rollback()
        self.assertEqual(self.rows('reception_shipments'), [])
        self.assertEqual(self.rows('reception_history'), [])

    def test_writer_reservation_prevents_new_work_during_apply_and_releases_afterward(self):
        self.shipment()
        self.db.commit()
        with tempfile.TemporaryDirectory() as folder:
            primary = sqlite3.connect(str(Path(folder) / 'qa.sqlite'))
            primary.row_factory = sqlite3.Row
            other = sqlite3.connect(str(Path(folder) / 'qa.sqlite'), timeout=0)
            try:
                self.db.backup(primary)
                lock_errors = []

                def probe_writer():
                    try:
                        other.execute('UPDATE reception_shipments SET received_packages=1')
                        other.commit()
                    except sqlite3.OperationalError as error:
                        lock_errors.append(str(error))
                        other.rollback()
                    return 1

                primary.create_function('qa_probe_writer', 0, probe_writer)
                primary.execute('''CREATE TRIGGER qa_probe BEFORE UPDATE OF app_status
                    ON reception_shipments BEGIN SELECT qa_probe_writer(); END''')
                preview = preview_historical_cutoff(primary, role=ADMIN)
                apply_historical_cutoff(primary, {'confirmed': True, 'reason': REASON,
                    'token': preview['token']}, 'qa.admin', ADMIN)
                self.assertEqual(lock_errors, ['database is locked'])
                self.assertEqual(primary.execute('SELECT app_status,received_packages '
                                                'FROM reception_shipments').fetchone()[:], ('CERRADO', 0))
                other.execute("UPDATE reception_shipments SET supplier='writer released'")
                other.commit()
            finally:
                primary.close()
                other.close()


if __name__ == '__main__':
    unittest.main()
