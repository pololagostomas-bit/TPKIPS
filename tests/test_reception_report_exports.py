"""Filtered indicators and bounded history, disposable data only."""
from io import BytesIO
from contextlib import closing
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import time
import tracemalloc
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import urlopen
from openpyxl import load_workbook
from backend import app
from backend.services import reception,report_jobs
from tests import test_report_jobs as fixtures
from tests.report_job_client import wait_response


class ReceptionExportTests(unittest.TestCase):
    setUp=fixtures.ReportJobTests.setUp
    tearDown=fixtures.ReportJobTests.tearDown
    request=fixtures.ReportJobTests.request

    def test_indicator_export_matches_filter_and_is_not_full_audit(self):
        with app.db() as db:
            db.execute("INSERT INTO reception_shipments(bl_awb,app_status,created_at,updated_at) VALUES ('OTHER','PROGRAMADO','2026-10-01','2026-10-01')")
        with wait_response(self.request('/api/receptions/report/export?search=QA-REPORT&state=CERRADO')) as response:
            self.assertIn('spreadsheetml',response.headers['Content-Type'])
            book=load_workbook(BytesIO(response.read()),read_only=True)
        self.assertIn('Indicadores BL',book.sheetnames)
        self.assertNotIn('Eventos auditoria',book.sheetnames)
        rows=list(book['Indicadores BL'].values)
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[1][0],'QA-REPORT')
        book.close()

    def test_different_filters_get_different_artifacts_identical_filters_share(self):
        with patch.object(report_jobs,'ensure_worker'):
            first=report_jobs.submit(self.source,'reception-dashboard-xlsx',{'search':'QA-REPORT'})
            same=report_jobs.submit(self.source,'reception-dashboard-xlsx',{'search':'qa-report'})
            other=report_jobs.submit(self.source,'reception-dashboard-xlsx',{'search':'OTHER'})
        self.assertEqual(first['id'],same['id'])
        self.assertNotEqual(first['id'],other['id'])

    def test_invalid_filter_or_worker_role_cannot_enqueue_export(self):
        for suffix in ('?state=INVALID','?from=2026-10-20&to=2026-10-01'):
            with self.assertRaises(HTTPError) as error:urlopen(self.request('/api/receptions/report/export'+suffix))
            self.assertEqual(error.exception.code,400)
        with self.assertRaises(HTTPError) as error:
            urlopen(self.request('/api/receptions/report/export',role='ASISTENTE_RECEPCION'))
        self.assertEqual(error.exception.code,403)

    def test_history_preview_memory_is_bounded_by_preview_not_full_audit(self):
        with app.db() as db:
            db.executemany('''INSERT INTO reception_history(shipment_id,event_type,field_name,username,reason,created_at)
                VALUES (?,'QA','test','qa.admin',?,'2026-10-02 10:00:00')''',((self.sid,'X'*8192) for _ in range(2000)))
            tracemalloc.start()
            try:
                report=reception.reception_history_report(db)
                peak=tracemalloc.get_traced_memory()[1]
            finally:tracemalloc.stop()
        self.assertEqual(report['summary']['events'],2601)
        self.assertEqual(len(report['events']),200)
        self.assertLess(peak,8*1024*1024,'Preview must not retain the full audit payload')
        print('QA preview2601 events: peak_bytes='+str(peak))

    def test_history_event_sink_retains_no_duplicate_event_list(self):
        count=0
        def receive(event):
            nonlocal count
            count+=1
        with app.db() as db:
            report=reception.reception_history_report(db,preview_limit=None,event_sink=receive)
        self.assertEqual(count,601)
        self.assertEqual(report['summary']['events'],601)
        self.assertEqual(report['events'],[])

    def test_xml_scratch_uses_report_volume_and_is_cleaned_after_failure(self):
        roots=[]
        def fail(*args,**kwargs):
            roots.append(Path(tempfile.gettempdir()))
            self.assertEqual(Path(os.environ['SQLITE_TMPDIR']),roots[-1])
            raise RuntimeError('synthetic failure')
        with patch.object(report_jobs,'ensure_worker'):
            job=report_jobs.submit(self.source,'reception-history-xlsx')
        with patch.object(reception,'reception_history_export_bytes',side_effect=fail):
            report_jobs._build(self.source,job)
        self.assertEqual(roots[0].parent,report_jobs.directory(self.source))
        self.assertFalse(roots[0].exists())

    def test_snapshot_history_order_is_indexed_without_changing_operational_schema(self):
        target=Path(self.temp.name)/'indexed-snapshot.sqlite3'
        report_jobs._snapshot(self.source,target,time.monotonic()+30)
        with closing(sqlite3.connect(target)) as db:
            plan=' '.join(str(tuple(row)) for row in db.execute(
                'EXPLAIN QUERY PLAN SELECT * FROM reception_history WHERE shipment_id=? ORDER BY id',(self.sid,)))
            self.assertNotIn('TEMP B-TREE',plan)
        with app.db() as db:
            self.assertIsNone(db.execute("SELECT 1 FROM sqlite_master WHERE name='idx_report_history_order'").fetchone())

    def test_worker_scratch_restores_sqlite_temp_environment_after_failure(self):
        with tempfile.TemporaryDirectory() as root,patch.dict(os.environ,{'SQLITE_TMPDIR':'QA-previous'}):
            with self.assertRaises(RuntimeError):
                with report_jobs._xml_scratch(Path(root),'a'*32):
                    self.assertNotEqual(os.environ['SQLITE_TMPDIR'],'QA-previous')
                    raise RuntimeError('QA failed')
            self.assertEqual(os.environ['SQLITE_TMPDIR'],'QA-previous')

    def test_audit_body_uses_default_style_and_preserves_literal_cells(self):
        with app.db() as db:
            body=reception.reception_history_export_bytes(db)
        book=load_workbook(BytesIO(body),read_only=True)
        try:
            cell=book['Eventos auditoria'].cell(2,15)
            self.assertEqual(cell.font.name,'Calibri')
            self.assertEqual(cell.data_type,'s')
            self.assertTrue(book['Eventos auditoria'].cell(1,1).font.bold)
        finally:book.close()

    def test_empty_filtered_export_is_a_valid_workbook_and_preserves_date_filters(self):
        from backend.services import reception_report_export as export
        options=export.query_options('from=2026-10-01&to=2026-10-31&search=missing&state=CERRADO')
        self.assertEqual((options['date_from'],options['date_to']),('2026-10-01','2026-10-31'))
        with app.db() as db:body=export.build(db,options)
        book=load_workbook(BytesIO(body))
        try:
            self.assertEqual(len(list(book['Indicadores BL'].values)),1)
            self.assertEqual(book['Indicadores BL'].auto_filter.ref,'A1:O1')
            self.assertIn('missing',str(list(book['Resumen'].values)))
        finally:book.close()


class CompleteFilteredCutTests(unittest.TestCase):
    def test_export_filter_does_not_silently_truncate_at_ui_page_size(self):
        report={'items':[{'id':n,'bl_awb':'QA','truck_guides':'','report_started_at':None,
            'app_status':'PROGRAMADO','sla_status':'SIN INICIO'} for n in range(135)]}
        filtered=reception.filter_reception_report(report,limit=None,search='QA')
        self.assertEqual(len(filtered['items']),135)
        self.assertEqual(filtered['total'],135)
