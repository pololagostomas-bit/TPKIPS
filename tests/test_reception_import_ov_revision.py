"""Provisional OV revisions preserve the selected physical work; synthetic data only."""
from io import BytesIO
import unittest
from openpyxl import load_workbook

from backend import app
from backend.services import reception
from tests import test_reception_import_identity as fixtures


class ImportOVRevisionTests(unittest.TestCase):
    setUp = fixtures.ImportLineIdentityTests.setUp
    load = fixtures.ImportLineIdentityTests.load
    workbook = fixtures.ImportLineIdentityTests.workbook
    rows = fixtures.ImportLineIdentityTests.rows

    def test_provisional_reference_becomes_real_ov_on_original_id_once(self):
        for provisional in ('POR REG', 'POR REGULARIZAR', 'COT 123456', 'COTI 123456', 'Cotizacion 123456'):
            with self.subTest(provisional=provisional):
                self.setUp()
                self.load(self.workbook([('54321', provisional, 'QA-IP')]))
                old_id = self.rows()[0]['id']
                for _ in range(2):
                    result = self.load(self.workbook([('54321', '159001', 'QA-IP')]))
                    self.assertEqual((result['lines_created'], result['lines_updated']), (0, 1))
                    self.assertEqual([(r['id'], r['ov_number']) for r in self.rows()], [(old_id, '159001')])
                with app.db() as c:
                    changes = c.execute('SELECT old_value,new_value FROM reception_history WHERE field_name=?',
                                        (f'linea:{old_id}:ov_number',)).fetchall()
                self.assertEqual([tuple(row) for row in changes], [(provisional, '159001')])

    def test_moved_excel_row_still_updates_same_line(self):
        self.load(self.workbook([('54321', 'POR REG', 'QA-IP')]))
        old_id = self.rows()[0]['id']
        book = self.workbook([('54321', '159001', 'QA-IP')])
        book.active.insert_rows(2, 6)
        self.load(book)
        self.assertEqual([(r['id'], r['source_row']) for r in self.rows()], [(old_id, 8)])

    def test_quote_number_revision_keeps_original_line_and_audits_only_once(self):
        self.load(self.workbook([('54321', 'COT 001234', 'QA-IP')]))
        old_id = self.rows()[0]['id']
        for _ in range(2):
            result = self.load(self.workbook([('54321', 'COT 005678', 'QA-IP')]))
            self.assertEqual((result['lines_created'],result['lines_updated']), (0,1))
            self.assertEqual([(r['id'],r['ov_number']) for r in self.rows()], [(old_id,'COT 005678')])
        with app.db() as c:
            changes = c.execute('SELECT old_value,new_value FROM reception_history WHERE field_name=?',
                                (f'linea:{old_id}:ov_number',)).fetchall()
        self.assertEqual([tuple(r) for r in changes], [('COT 001234','COT 005678')])

    def test_two_quotes_still_present_in_source_are_separate_orders(self):
        self.load(self.workbook([('54321', 'COT 001234', 'QA-IP')]))
        old_id = self.rows()[0]['id']
        self.load(self.workbook([('54321', 'COT 005678', 'QA-IP'),('54321', 'COT 001234', 'QA-IP')]))
        self.assertEqual(len(self.rows()),2)
        self.assertEqual(next(r['id'] for r in self.rows() if r['ov_number']=='COT 001234'),old_id)

    def test_ambiguous_quote_revision_rejects_without_creating_a_duplicate(self):
        self.load(self.workbook([('54321', 'COT 001234', 'QA-IP'),('54321', 'COT 005678', 'QA-IP')]))
        before = self.rows()
        with self.assertRaisesRegex(ValueError,'OV ambigua'):
            self.load(self.workbook([('54321', 'COT 009999', 'QA-IP')]))
        self.assertEqual(self.rows(),before)

    def test_existing_real_or_stock_orders_do_not_become_another_ov(self):
        for old_ov in ('158001', 'S/N', ''):
            with self.subTest(old_ov=old_ov):
                self.setUp()
                self.load(self.workbook([('54321', old_ov, 'QA-IP')]))
                self.assertEqual(self.load(self.workbook([('54321', '159001', 'QA-IP')]))['lines_created'], 1)
                self.assertEqual(len(self.rows()), 2)

    def test_provisional_still_in_source_and_real_ov_remain_separate(self):
        self.load(self.workbook([('54321', 'POR REG', 'QA-IP')]))
        old_id = self.rows()[0]['id']
        self.load(self.workbook([('54321', '159001', 'QA-IP'), ('54321', 'POR REG', 'QA-IP')]))
        self.assertEqual(len(self.rows()), 2)
        self.assertEqual(next(r['id'] for r in self.rows() if r['ov_number'] == 'POR REG'), old_id)

    def test_ambiguous_old_or_incoming_ovs_fail_before_writes_in_either_order(self):
        for previous, incoming in [
            ([('54321', 'POR REG', 'QA-IP'), ('54321', 'COT 1', 'QA-IP')], [('54321', '159001', 'QA-IP')]),
            ([('54321', 'POR REG', 'QA-IP')], [('54321', '159001', 'QA-IP'), ('54321', '159002', 'QA-IP')]),
            ([('54321', 'POR REG', 'QA-IP'), ('54321', '158001', 'QA-IP')], [('54321', '159001', 'QA-IP')]),
        ]:
            self.setUp()
            self.load(self.workbook(previous))
            before = self.rows()
            for records in (incoming, list(reversed(incoming))):
                with self.subTest(records=records), self.assertRaisesRegex(ValueError, 'OV ambigua'):
                    self.load(self.workbook(records))
                self.assertEqual(self.rows(), before)

    def test_missing_ip_or_oc_cannot_reassign_provisional_work(self):
        for oc, ip in [('', 'QA-IP'), ('54321', ''), ('', '')]:
            self.setUp()
            self.load(self.workbook([(oc, 'POR REG', ip)]))
            before = self.rows()
            with self.subTest(oc=oc, ip=ip), self.assertRaisesRegex(ValueError, 'OV ambigua'):
                self.load(self.workbook([(oc, '159001', ip)]))
            self.assertEqual(self.rows(), before)

    def test_other_ip_oc_np_bl_or_sheet_are_not_inferred_from_position(self):
        for column, value in [(1, 'QA-OTHER-BL'), (2, 'QA-OTHER-NP'), (3, 'QA-OTHER-IP'), (5, '65432')]:
            self.setUp()
            self.load(self.workbook([('54321', 'POR REG', 'QA-IP')]))
            book = self.workbook([('54321', '159001', 'QA-IP')])
            book.active.cell(2, column, value)
            self.assertEqual(self.load(book)['lines_created'], 1)
            self.assertEqual(len(self.rows()), 2)
        self.setUp()
        self.load(self.workbook([('54321', 'POR REG', 'QA-IP')]))
        book = self.workbook([('54321', '159001', 'QA-IP')])
        book.active.title = 'OTHER BUYER'
        self.assertEqual(self.load(book)['lines_created'], 1)

    def test_exact_real_order_is_reserved_before_other_pending_ov_revision(self):
        self.load(self.workbook([('54321', 'POR REG', 'QA-IP'), ('54321', '158001', 'QA-IP')]))
        before = {r['ov_number']: r['id'] for r in self.rows()}
        result = self.load(self.workbook([('54321', '159001', 'QA-IP'), ('54321', '158001', 'QA-IP')]))
        self.assertEqual(result['lines_created'], 0)
        self.assertEqual({r['ov_number']: r['id'] for r in self.rows()},
                         {'159001': before['POR REG'], '158001': before['158001']})

    def test_repeated_provisional_occurrences_cannot_swap_work_after_one_gets_ov(self):
        self.load(self.workbook([('54321', 'POR REG', 'QA-IP'), ('54321', 'POR REG', 'QA-IP')]))
        before = self.rows()
        records = [('54321', '159001', 'QA-IP'), ('54321', 'POR REG', 'QA-IP')]
        with self.assertRaisesRegex(ValueError, 'OV ambigua'):
            self.load(self.workbook(records))
        self.assertEqual(self.rows(), before)

    def test_attention_scan_quantity_destination_detail_and_report_follow_current_ov(self):
        self.load(self.workbook([('54321', 'POR REG', 'QA-IP')]))
        old = self.rows()[0]
        with app.db() as c:
            c.execute("UPDATE reception_shipments SET app_status='REVISION SISTEMA',received_packages=1 WHERE id=?", (old['shipment_id'],))
            shipment = c.execute('SELECT * FROM reception_shipments WHERE id=?', (old['shipment_id'],)).fetchone()
            attention = reception._create_reception_attention(c, shipment, None, 'qa.admin', 'REVISION SISTEMA')
            c.execute("""UPDATE reception_attention_lines SET verified_qty=3,np_quantity_saved=1,
                final_location='QA-BIN',count_location_saved=1,np_scanned_at='2026-10-23T09:00:00',
                np_scanned_by='qa.admin' WHERE attention_id=?""", (attention['id'],))
            physical_before = tuple(c.execute('SELECT * FROM reception_attention_lines WHERE attention_id=?', (attention['id'],)).fetchone())
        book = self.workbook([('54321', '159001', 'QA-IP')])
        book.active.cell(2, 6, 9)
        self.load(book)
        with app.db() as c:
            self.assertEqual(tuple(c.execute('SELECT * FROM reception_attention_lines WHERE attention_id=?', (attention['id'],)).fetchone()), physical_before)
            detail = reception.reception_detail(c, old['shipment_id'], attention_id=attention['id'])
            selected = [r for r in detail['lines'] if r.get('in_selected_attention')]
            self.assertEqual([(r['id'], r['ov_number'], r['attention_planned_qty'], r['attention_verified_qty'], r['final_location']) for r in selected],
                             [(old['id'], '159001', 5, 3, 'QA-BIN')])
            content = reception.reception_operational_report_bytes(c, old['shipment_id'], attention_id=attention['id'])
        with BytesIO(content) as stream:
            exported = load_workbook(stream, data_only=True)
            try:
                rows = list(exported['Packing List SAP'].iter_rows(min_row=2, values_only=True))
                self.assertEqual([(r[4], r[5], r[6]) for r in rows], [('159001', 5, 3)])
            finally:
                exported.close()
