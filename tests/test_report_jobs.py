"""No production DB: concurrency, read-only snapshots, queue limits and real child process."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from backend import app
from backend.services import reception, report_jobs
from tests.report_job_client import wait_response


class ReportJobTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='wms-report-qa-')
        self.source = Path(self.temp.name) / 'qa.sqlite3'
        self.env = patch.dict(os.environ, {'TRITON_AUTH_MODE': 'demo'})
        self.env.start()
        self.database = patch.object(app, 'DB_PATH', self.source)
        self.database.start()
        app.init_db()
        with app.db() as connection:
            self.sid = connection.execute('''INSERT INTO reception_shipments
                (bl_awb,app_status,transport_type,created_at,updated_at)
                VALUES ('QA-REPORT','CERRADO','MARITIMO','2026-10-01 08:00:00','2026-10-02 08:00:00')''').lastrowid
            for index in range(601):
                connection.execute('''INSERT INTO reception_history
                    (shipment_id,event_type,field_name,new_value,username,reason,created_at)
                    VALUES (?,'QA','test',?,'qa.admin',?,'2026-10-01 09:00:00')''',
                    (self.sid, str(index), '=HYPERLINK("https://example.invalid")'))
        self.server = app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        self.thread = threading.Thread(target=self.server.serve_forever,daemon=True)
        self.thread.start()
        self.base = f'http://127.0.0.1:{self.server.server_port}'

    def tearDown(self):
        self.server.shutdown(); self.server.server_close(); self.thread.join()
        child = report_jobs._children.pop(str(self.source.resolve()), None)
        if child:
            child.wait(timeout=30)
        self.database.stop(); self.env.stop(); self.temp.cleanup()

    def request(self,path,role='ADMINISTRADOR',user='qa.admin'):
        return Request(self.base+path, headers={'X-Role':role,'X-User':user})

    def test_two_admins_share_one_background_job_and_live_requests_remain_available(self):
        with patch.object(report_jobs,'ensure_worker'):
            start=time.monotonic()
            with ThreadPoolExecutor(max_workers=2) as pool:
                def enqueue(user):
                    with urlopen(self.request('/api/receptions/history/export',user=user)) as response:
                        self.assertEqual(response.status,202)
                        return json.load(response)
                jobs=list(pool.map(enqueue,['qa.admin','qa.other']))
            self.assertLess(time.monotonic()-start,2)
            self.assertEqual(jobs[0]['id'],jobs[1]['id'])
            with report_jobs.queue(self.source) as queue:
                self.assertEqual(queue.execute('SELECT COUNT(*) FROM jobs').fetchone()[0],1)
            with app.db() as connection:
                connection.execute("UPDATE reception_shipments SET supplier='Concurrent write' WHERE id=?",(self.sid,))
            with urlopen(self.request('/health')) as response:
                self.assertEqual(json.load(response)['status'],'ok')
        with wait_response(self.request('/api/receptions/history/export')) as response:
            self.assertIn('spreadsheetml',response.headers['Content-Type'])
            book=load_workbook(BytesIO(response.read()),read_only=True)
        self.assertEqual(sum(1 for _ in book['Eventos auditoria'].rows),602)
        self.assertEqual(book['Eventos auditoria'].cell(2,15).data_type,'s')
        book.close()
        with report_jobs.queue(self.source) as queue:
            self.assertEqual(queue.execute('SELECT COUNT(*) FROM jobs').fetchone()[0],1)
        self.assertEqual(list(report_jobs.directory(self.source).glob('*.snapshot.sqlite3*')),[])
        with app.db() as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM reception_history').fetchone()[0],601)

    def test_wal_snapshot_is_consistent_and_cannot_write_back(self):
        original=sqlite3.connect(self.source)
        try:
            original.execute('PRAGMA journal_mode=WAL')
            original.execute("UPDATE reception_shipments SET supplier='WAL commit' WHERE id=?",(self.sid,))
            original.commit()
            target=Path(self.temp.name)/'snapshot.sqlite3'
            report_jobs._snapshot(self.source,target,time.monotonic()+10)
            copy=sqlite3.connect(target.as_uri()+'?mode=ro',uri=True)
            self.assertEqual(copy.execute('SELECT supplier FROM reception_shipments').fetchone()[0],'WAL commit')
            with self.assertRaises(sqlite3.OperationalError):
                copy.execute("UPDATE reception_shipments SET supplier='bad'")
            copy.close()
        finally:
            original.close()

    def test_live_health_and_writes_while_large_excel_runs_in_child(self):
        with app.db() as connection:
            connection.executemany('''INSERT INTO reception_history
                (shipment_id,event_type,field_name,new_value,username,created_at)
                VALUES (?,'QA','load',?,'qa.admin','2026-10-01 10:00:00')''',
                ((self.sid,str(index)) for index in range(10000)))
        started=time.monotonic()
        with urlopen(self.request('/api/receptions/history/export')) as response:
            job=json.load(response)
        self.assertLess(time.monotonic()-started,2)
        deadline=time.monotonic()+20
        while time.monotonic()<deadline:
            state=report_jobs.lookup(self.source,job['id'])
            if state['snapshot_at']:
                break
            time.sleep(.02)
        self.assertIsNotNone(state['snapshot_at'])
        self.assertEqual(state['status'],'RUNNING','A real export is in progress')
        def health(_):
            start=time.monotonic()
            with urlopen(self.request('/health'),timeout=3) as response:
                self.assertEqual(json.load(response)['status'],'ok')
            return time.monotonic()-start
        with ThreadPoolExecutor(max_workers=2) as pool:
            latencies=list(pool.map(health,range(6)))
        with app.db() as connection:
            connection.execute('''INSERT INTO reception_history
                (shipment_id,event_type,field_name,new_value,username,created_at)
                VALUES (?,'AFTER CUT','test','later','qa.admin','2026-10-01 11:00:00')''',(self.sid,))
        self.assertLess(max(latencies),2)
        with wait_response(self.request('/api/receptions/history/export')) as response:
            book=load_workbook(BytesIO(response.read()),read_only=True)
        self.assertEqual(sum(1 for _ in book['Eventos auditoria'].rows),10602)
        book.close()
        print('QA background 10601 events: health_max_ms='+str(round(max(latencies)*1000,1)))

    def test_report_status_and_artifact_require_current_admin_permission(self):
        with wait_response(self.request('/api/receptions/report?state=CERRADO')) as response:
            data=json.load(response)
        self.assertEqual(data['total'],1)
        self.assertEqual(data['items'][0]['app_status'],'CERRADO')
        with report_jobs.queue(self.source) as queue:
            job=queue.execute('SELECT id FROM jobs').fetchone()[0]
        for role in ['ASISTENTE_RECEPCION','AUXILIAR_RECEPCION','PICKER']:
            for path in ['/api/receptions/history/export','/api/report-jobs/'+job,'/api/report-jobs/'+job+'/result']:
                with self.assertRaises(HTTPError) as error:
                    urlopen(self.request(path,role=role))
                self.assertEqual(error.exception.code,403)
        with self.assertRaises(HTTPError) as error:
            urlopen(self.request('/api/report-jobs/not-a-job/result'))
        self.assertEqual(error.exception.code,400)

    def test_failed_job_is_not_silently_retried_and_snapshot_is_removed(self):
        with patch.object(report_jobs,'ensure_worker'):
            job=report_jobs.submit(self.source,'reception-history-xlsx')
        with patch.object(reception,'reception_history_export_bytes',side_effect=RuntimeError('PRIVATE DETAIL')):
            report_jobs._build(self.source,job)
        failed=report_jobs.lookup(self.source,job['id'])
        self.assertEqual(failed['status'],'FAILED')
        self.assertNotIn('PRIVATE',failed['error'])
        self.assertEqual(list(report_jobs.directory(self.source).glob('*.snapshot.sqlite3*')),[])
        with patch.object(report_jobs,'ensure_worker'):
            retry=report_jobs.submit(self.source,'reception-history-xlsx')
        self.assertNotEqual(retry['id'],job['id'])

    def test_export_builds_history_once_and_does_not_truncate_to_preview(self):
        with app.db() as connection:
            with patch.object(reception,'reception_history_report',wraps=reception.reception_history_report) as build:
                body=reception.reception_history_export_bytes(connection)
            self.assertEqual(build.call_count,1)
            self.assertIsNone(build.call_args.kwargs['preview_limit'])
        book=load_workbook(BytesIO(body),read_only=True)
        self.assertEqual(sum(1 for _ in book['Eventos auditoria'].rows),602)
        book.close()

    def test_expiry_and_queue_backpressure(self):
        with patch.object(report_jobs,'ensure_worker'):
            job=report_jobs.submit(self.source,'reception-dashboard')
            self.assertEqual(report_jobs.submit(self.source,'reception-dashboard')['id'],job['id'])
            with patch.object(report_jobs,'MAX_QUEUE',1):
                with self.assertRaises(ValueError):
                    report_jobs.submit(self.source,'reception-history-xlsx')
            with report_jobs.queue(self.source) as connection:
                connection.execute("UPDATE jobs SET status='READY',finished=? WHERE id=?",
                                   (time.time()-report_jobs.RETENTION_SECONDS-1,job['id']))
            with self.assertRaises(ValueError):
                report_jobs.lookup(self.source,job['id'])


if __name__=='__main__':
    unittest.main()
