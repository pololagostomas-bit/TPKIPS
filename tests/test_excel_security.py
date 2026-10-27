"""Small, synthetic adversarial workbooks; never contact an external entity."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from backend.services import safe_excel


def workbook_bytes():
    book = Workbook()
    book.active.append(['NP', 'Cantidad', 'Total'])
    book.active.append(['NP-TEST', 3, '=B2*2'])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def archive(members):
    stream = io.BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as zipped:
        for name, content in members:
            zipped.writestr(name, content)
    return stream.getvalue()


class ExcelSecurityTests(unittest.TestCase):
    def test_valid_stream_keeps_formula_and_reader_options(self):
        stream = io.BytesIO(workbook_bytes())
        book = safe_excel.load_workbook(stream, read_only=True, data_only=False)
        try:
            self.assertTrue(book.read_only)
            self.assertEqual(list(book.active.values)[1], ('NP-TEST', 3, '=B2*2'))
        finally:
            book.close()
        self.assertFalse(stream.closed)

    def test_path_and_data_only_remain_supported(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'normal.xlsx'
            path.write_bytes(workbook_bytes())
            book = safe_excel.load_workbook(path, data_only=True)
            try:
                self.assertEqual(book.active['B2'].value, 3)
                self.assertIsNone(book.active['C2'].value)
            finally:
                book.close()

    def test_expansion_limit_rejects_highly_compressible_content_before_parser(self):
        content = archive([('xl/worksheets/sheet1.xml', b'A'*4096)])
        with patch.object(safe_excel, 'MAX_EXPANDED_BYTES', 2048), \
                patch.object(safe_excel, '_load_workbook', side_effect=AssertionError('Parser reached')):
            with self.assertRaisesRegex(ValueError, 'descomprimir'):
                safe_excel.load_workbook(io.BytesIO(content))

    def test_individual_member_limit_is_not_only_a_total_limit(self):
        content = archive([('xl/worksheets/sheet1.xml', b'A'*4096)])
        with patch.object(safe_excel, 'MAX_MEMBER_BYTES', 2048), \
                patch.object(safe_excel, '_load_workbook', side_effect=AssertionError('Parser reached')):
            with self.assertRaisesRegex(ValueError, 'hoja|entrada'):
                safe_excel.load_workbook(io.BytesIO(content))

    def test_too_many_members_rejected(self):
        content = archive([(f'xl/{i}.xml', b'x') for i in range(4)])
        with patch.object(safe_excel, 'MAX_ENTRIES', 3):
            with self.assertRaisesRegex(ValueError, 'entradas'):
                safe_excel.load_workbook(io.BytesIO(content))

    def test_parser_cannot_silently_fall_back_to_unprotected_xml(self):
        with patch.object(safe_excel, 'DEFUSEDXML', False):
            with self.assertRaisesRegex(ValueError, 'XML protegido'):
                safe_excel.load_workbook(io.BytesIO(workbook_bytes()))

    def test_compressed_limit_and_invalid_zip_leave_stream_open_at_original_position(self):
        for content, limit in [(workbook_bytes(), 16), (b'not a zip archive', 50000)]:
            stream = io.BytesIO(content)
            stream.seek(2)
            with patch.object(safe_excel, 'MAX_FILE_BYTES', limit):
                with self.assertRaises(ValueError):
                    safe_excel.load_workbook(stream)
            self.assertFalse(stream.closed)
            self.assertEqual(stream.tell(), 2)

    def test_worksheet_xml_entities_are_rejected_when_loaded_or_read(self):
        with ZipFile(io.BytesIO(workbook_bytes())) as zipped:
            members = [(entry.filename, zipped.read(entry)) for entry in zipped.infolist()]
        attack = (b'<?xml version="1.0"?><!DOCTYPE worksheet [<!ENTITY internal "UNSAFE">]>'
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&internal;</t></is>'
            b'</c></row></sheetData></worksheet>')
        members = [(name, attack if name == 'xl/worksheets/sheet1.xml' else value)
                   for name, value in members]
        book = None
        try:
            with self.assertRaises(ValueError):
                book = safe_excel.load_workbook(io.BytesIO(archive(members)), read_only=True)
                list(book.active.values)
        finally:
            if book:
                book.close()

    def test_rejected_daily_import_preserves_previous_stock_and_import_history(self):
        from backend import app
        with tempfile.TemporaryDirectory() as folder, \
                patch.object(app, 'DB_PATH', Path(folder)/'synthetic.db'):
            app.init_db()
            with app.db() as connection:
                connection.execute('''INSERT INTO inventory_stock
                    (item_key,warehouse_key,item_code,warehouse,snapshot_qty,snapshot_day,
                     snapshot_at,is_current,updated_at) VALUES
                    ('OLD','1','OLD','1',17,'2026-09-01','2026-09-01T09:00:00',1,'2026-09-01T09:00:00')''')
            book = Workbook()
            book.active.append(['Número de artículo', 'En stock', 'Almacén'])
            book.active.append(['NEW', 7, '01'])
            content = io.BytesIO()
            book.save(content)
            book.close()
            with patch.object(safe_excel, 'MAX_EXPANDED_BYTES', 1):
                with self.assertRaisesRegex(ValueError, 'descomprimir'):
                    app.import_daily_excel(content.getvalue(), 'synthetic.xlsx', 'stock',
                        '2026-10-01T09:00:00', 'qa', 'ADMINISTRADOR', reconcile_delivered=False)
            with app.db() as connection:
                self.assertEqual([tuple(row) for row in connection.execute(
                    'SELECT item_code,snapshot_qty,is_current FROM inventory_stock')], [('OLD',17,1)])
                self.assertEqual(connection.execute('SELECT COUNT(*) FROM data_imports').fetchone()[0], 0)


if __name__ == '__main__':
    unittest.main()
