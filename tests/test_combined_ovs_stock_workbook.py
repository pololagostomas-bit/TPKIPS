"""Regresión: el query diario puede traer OVs y stock en hojas separadas."""

<<<<<<< HEAD
import io
=======
>>>>>>> 3b9f04f67883bd897fae4700181dda909c5f0312
import unittest
from pathlib import Path

from openpyxl import Workbook

import app


class CombinedOvsStockWorkbookTest(unittest.TestCase):
<<<<<<< HEAD
    def test_imports_stock_only_without_replacing_orders(self):
        original_db = app.DB_PATH
        db_path = Path(__file__).parent / ".stock-only-cutoff.db"
        db_path.unlink(missing_ok=True)
        app.DB_PATH = db_path
        try:
            app.init_db()
            book = Workbook()
            sheet = book.active
            sheet.title = "Stock almacén"
            sheet.append(["Número de artículo", "En stock"])
            sheet.append(["SKU-SOLO-STOCK", 12])
            content = io.BytesIO()
            book.save(content)
            book.close()

            result = app.import_daily_excel(
                content.getvalue(), "stock-almacen.xlsx", "stock",
                "2026-09-08T16:00:00", "admin.test", "ADMINISTRADOR",
            )
            with app.db() as connection:
                row = connection.execute(
                    "SELECT snapshot_qty, is_current FROM inventory_stock WHERE item_code = ?",
                    ("SKU-SOLO-STOCK",),
                ).fetchone()
                order_count = connection.execute("SELECT COUNT(*) FROM orders").fetchone()[0]
            self.assertTrue(result["stock_updated"])
            self.assertEqual(result["inventory_skus"], 1)
            self.assertEqual(row["snapshot_qty"], 12)
            self.assertEqual(row["is_current"], 1)
            self.assertEqual(order_count, 0)
        finally:
            app.DB_PATH = original_db
            db_path.unlink(missing_ok=True)

=======
>>>>>>> 3b9f04f67883bd897fae4700181dda909c5f0312
    def test_imports_stock_sheet_as_global_cutoff(self):
        original_db = app.DB_PATH
        test_dir = Path(__file__).parent
        app.DB_PATH = test_dir / ".combined-ovs-stock.db"
        workbook_path = test_dir / ".combined-ovs-stock-query.xlsx"
        workbook_path.unlink(missing_ok=True)
        app.DB_PATH.unlink(missing_ok=True)
        try:
                book = Workbook()
                ovs = book.active
                ovs.title = "Detalle diario"
                ovs.append([
                    "STATUS", "Número de documento", "Status de Documento", "Grupo de Articulo",
                    "Almacen", "N° de Línea", "Número de artículo", "Descripción artículo/serv.",
                    "Cantidad", "Cant. Pendiente", "Fecha de contabilización",
                ])
                ovs.append(["OV ABIERTA", "OV-1", "Abierto", "REPUESTOS", 1, 1, "SKU-CORTE", "Repuesto", 2, 2, "2026-09-08"])
                stock = book.create_sheet("Corte SAP")
                stock.append([])
                stock.append(["Número de artículo", "Descripción del artículo", "Unidad de medida de inventario", "Primera ubicación", "En stock", "Comprometido", "Solicitado", "Disponible"])
                stock.append(["Almacén", 1])
                stock.append(["SKU-CORTE", "Repuesto", "UNIDAD", "01-01-A", 4, 3, 0, 1])
                book.save(workbook_path)
                book.close()

                app.init_db()
                result = app.import_excel(workbook_path, reset=True, cutoff_at="2026-09-08T16:00:00")
                payload = app.order_payload("OV-1")

                self.assertTrue(result["stock_updated"])
                self.assertEqual(result["stock_sheet"], "Corte SAP")
                self.assertEqual(payload["lines"][0]["stock_source_qty"], 4)
                self.assertEqual(payload["lines"][0]["stock_free_qty"], 4)
        finally:
            app.DB_PATH.unlink(missing_ok=True)
            workbook_path.unlink(missing_ok=True)
            app.DB_PATH = original_db


if __name__ == "__main__":
    unittest.main()
