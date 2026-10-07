"""Accounting sheet rename regressions on an isolated in-memory database."""
import hashlib
from io import BytesIO
import sqlite3
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from backend.services import reception as service
from backend.services.traceability import init_traceability_schema


NEW_SHEET = "FACTURAS IMPORTACIONES Repuesto"
HEADERS = ["AWB", "IP", "FACTURA RESERVA", "ENTRADA DE MERCANCIA", "FECHA DE RECEPCION", "FECHA DE EM"]


class AccountingSheetNamesTest(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.addCleanup(self.db.close)
        service.init_reception_schema(self.db)
        init_traceability_schema(self.db)
        self.db.execute("CREATE TABLE order_importation_refs (bl_awb TEXT, sap_ov TEXT, transport_type TEXT, ip_reference TEXT, oc_number TEXT)")
        self.db.execute("CREATE TABLE attentions (id INTEGER PRIMARY KEY, sap_ov TEXT)")
        lots = patch.object(service, "advanced_lots_enabled", return_value=False)
        lots.start()
        self.addCleanup(lots.stop)
        self.shipment_id = service.create_reception(self.db, {"bl_awb":"AWB-SHEET-1"}, "admin", "ADMINISTRADOR")["id"]
        self.db.execute("INSERT INTO reception_lines (shipment_id,np_code,ip_reference,expected_qty) VALUES (?, 'NP-1', 'IP-1', 2)", (self.shipment_id,))

    def workbook(self, title=NEW_SHEET, fr="FR-1"):
        workbook = Workbook()
        self.addCleanup(workbook.close)
        sheet = workbook.active
        sheet.title = title
        sheet.append(HEADERS)
        sheet.append(["AWB-SHEET-1", "IP-1", fr, "", "", ""])
        return workbook

    def load(self, workbook):
        return service.import_reception_accounting_workbook(self.db, workbook, "cortes.xlsx", "admin", "ADMINISTRADOR")

    def test_new_name_and_read_only_excel_upload(self):
        workbook = self.workbook()
        content = BytesIO()
        workbook.save(content)
        result = service.import_reception_accounting_excel_bytes(self.db, content.getvalue(), "cortes.xlsx", "admin", "ADMINISTRADOR")
        self.assertEqual(result["sheet"], NEW_SHEET)
        self.assertEqual(result["matched_records"], 1)
        self.assertEqual(result["references_created"], 1)

    def test_legacy_name_remains_supported(self):
        result = self.load(self.workbook("FACTURAS DHL"))
        self.assertEqual(result["sheet"], "FACTURAS DHL")
        self.assertEqual(result["matched_records"], 1)

    def test_case_is_normalized_and_reload_is_idempotent(self):
        workbook = self.workbook(NEW_SHEET.lower())
        self.load(workbook)
        result = self.load(workbook)
        self.assertEqual(result["references_created"], 0)
        self.assertEqual(result["references_unchanged"], 1)

    def test_new_sheet_takes_precedence_over_old_copy(self):
        workbook = self.workbook()
        old = workbook.create_sheet("FACTURAS DHL")
        old.append(HEADERS)
        old.append(["AWB-SHEET-1", "IP-1", "FR-STALE", "", "", ""])
        result = self.load(workbook)
        self.assertEqual(result["records_read"], 1)
        row = self.db.execute("SELECT fr_number,source_sheet FROM reception_accounting_refs").fetchone()
        self.assertEqual(tuple(row), ("FR-1", NEW_SHEET))

    def test_rename_updates_existing_reference_instead_of_duplicating(self):
        workbook = self.workbook("FACTURAS DHL")
        self.load(workbook)
        before = self.db.execute("SELECT id,source_key FROM reception_accounting_refs").fetchone()
        workbook.active.title = NEW_SHEET
        workbook.active.cell(2, 3, "FR-UPDATED")
        result = self.load(workbook)
        after = self.db.execute("SELECT id,source_key,fr_number,source_sheet FROM reception_accounting_refs").fetchone()
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM reception_accounting_refs").fetchone()[0], 1)
        self.assertEqual((after["id"],after["source_key"]), tuple(before))
        self.assertEqual(after["fr_number"], "FR-UPDATED")
        self.assertEqual(after["source_sheet"], NEW_SHEET)
        self.assertEqual(result["references_created"], 0)
        self.assertEqual(result["references_updated"], 1)

    def test_historical_mixed_case_keys_are_reused(self):
        workbook = self.workbook("Facturas DHL")
        self.load(workbook)
        old_id = self.db.execute("SELECT id FROM reception_accounting_refs").fetchone()[0]
        old_key = hashlib.sha256(b"Facturas DHL|AWBSHEET1|IP1|1").hexdigest()
        self.db.execute("UPDATE reception_accounting_refs SET source_key=?", (old_key,))
        workbook.active.title = NEW_SHEET
        first = self.load(workbook)
        second = self.load(workbook)
        self.assertEqual(first["references_created"], 0)
        self.assertEqual(second["references_created"], 0)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM reception_accounting_refs").fetchone()[0], 1)
        self.assertEqual(self.db.execute("SELECT id FROM reception_accounting_refs").fetchone()[0], old_id)

    def test_invalid_sheet_is_rejected_with_current_name(self):
        with self.assertRaisesRegex(ValueError, NEW_SHEET):
            self.load(self.workbook("OTRA HOJA"))
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM reception_accounting_refs").fetchone()[0], 0)

    def test_non_admin_cannot_import_new_sheet(self):
        with self.assertRaises(PermissionError):
            service.import_reception_accounting_workbook(self.db, self.workbook(), "cortes.xlsx", "aux", "AUXILIAR_RECEPCION")
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM reception_accounting_refs").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
