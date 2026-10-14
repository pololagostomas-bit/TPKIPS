"""Extract MARCAS in-cell logos without importing inventory rows or altering images."""
import argparse
import hashlib
import json
import posixpath
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.services.reception_brands import BRANDS, MAX_LOGO_BYTES

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RICH = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE = "http://schemas.openxmlformats.org/package/2006/relationships"
RICH_REL = "http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel"


def read_part(archive, path, limit=2 * 1024 * 1024):
    with archive.open(path) as stream:
        data = stream.read(limit + 1)
    if len(data) > limit:
        raise ValueError("Parte del Excel demasiado grande: " + path)
    return data


def xml(archive, path):
    return ET.fromstring(read_part(archive, path))


def relationships(archive, path, base, kind="image"):
    result = {}
    for element in xml(archive, path).findall("{" + PACKAGE + "}Relationship"):
        if element.get("Type") != REL + "/" + kind:
            continue
        if element.get("TargetMode") == "External":
            continue
        target = element.get("Target", "")
        target = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base, target))
        if not target.startswith("xl/") or "\\" in target:
            raise ValueError("Relacion de imagen fuera del libro")
        result[element.get("Id")] = target
    return result


def indexed(items, index):
    index = int(index)
    if index < 0 or index >= len(items):
        raise ValueError("Indice de metadatos fuera de rango")
    return items[index]


def local_image_part(archive, cell, metadata, rich_values, structures, rich_links, targets):
    vm = int(cell.get("vm", "0")) - 1
    values = metadata.findall("{" + MAIN + "}valueMetadata/{" + MAIN + "}bk")
    if vm < 0 or vm >= len(values):
        raise ValueError("La celda no tiene metadatos de imagen locales")
    types = metadata.findall("{" + MAIN + "}metadataTypes/{" + MAIN + "}metadataType")
    record = next((record for record in values[vm] if indexed(types, int(record.get("t")) - 1).get("name") == "XLRICHVALUE"), None)
    future = next((item for item in metadata if item.get("name") == "XLRICHVALUE"), None)
    if record is None or future is None:
        raise ValueError("Tipo de metadatos de imagen no compatible")
    # Follow the cell binding and named structure key, not the media-file order.
    future_record = indexed(future.findall("{" + MAIN + "}bk"), record.get("v"))
    binding = future_record.find(".//{" + RICH + "}rvb")
    if binding is None:
        raise ValueError("Vinculo de imagen ausente")
    value = indexed(rich_values.findall("{" + RICH + "}rv"), binding.get("i"))
    structure = indexed(structures.findall("{" + RICH + "}s"), value.get("s"))
    if structure.get("t") != "_localImage":
        raise ValueError("Se requiere una imagen local en celda, no un enlace remoto")
    keys = [key.get("n") for key in structure.findall("{" + RICH + "}k")]
    index = int(value.findall("{" + RICH + "}v")[keys.index("_rvRel:LocalImageIdentifier")].text)
    link = indexed(rich_links.findall("{" + RICH_REL + "}rel"), index)
    target = targets[link.get("{" + REL + "}id")]
    if not target.startswith("xl/media/"):
        raise ValueError("La relacion no apunta a una imagen del libro")
    return target


def extract(workbook, output):
    from openpyxl import load_workbook
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = next(sheet for sheet in book if sheet.title.strip().upper() == "MARCAS")
        rows = list(sheet.iter_rows(min_row=6, max_row=15, min_col=2, max_col=2, values_only=True))
        brands = {"E" + str(index): str(row[0] or "").strip().upper() for index, row in enumerate(rows, 6)}
        sheet_name = sheet.title
    finally:
        book.close()
    logos = []
    with zipfile.ZipFile(workbook) as archive:
        workbook_xml = xml(archive, "xl/workbook.xml")
        worksheet = next(item for item in workbook_xml.findall("{" + MAIN + "}sheets/{" + MAIN + "}sheet") if item.get("name") == sheet_name)
        sheets = relationships(archive, "xl/_rels/workbook.xml.rels", "xl", "worksheet")
        sheet_xml = xml(archive, sheets[worksheet.get("{" + REL + "}id")])
        metadata = xml(archive, "xl/metadata.xml")
        values = xml(archive, "xl/richData/rdrichvalue.xml")
        structures = xml(archive, "xl/richData/rdrichvaluestructure.xml")
        links = xml(archive, "xl/richData/richValueRel.xml")
        targets = relationships(archive, "xl/richData/_rels/richValueRel.xml.rels", "xl/richData")
        for cell in sheet_xml.findall(".//{" + MAIN + "}c"):
            brand = brands.get(cell.get("r"))
            if not brand:
                continue
            if brand not in BRANDS:
                raise ValueError("Marca no admitida: " + brand)
            target = local_image_part(archive, cell, metadata, values, structures, links, targets)
            data = read_part(archive, target, MAX_LOGO_BYTES)
            extension = ".png" if data.startswith(b"\x89PNG\r\n\x1a\n") else ".jpg" if data.startswith(b"\xff\xd8\xff") else ""
            if not extension:
                raise ValueError("Solo se admiten PNG/JPEG locales")
            logos.append((brand, brand.lower() + extension, data))
    if set(brand for brand, _, _ in logos) != set(BRANDS) or len(logos) != len(BRANDS):
        raise ValueError("Se requieren las diez marcas, una imagen por marca")
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    for _, filename, data in logos:
        destination = output / filename
        if destination.is_symlink() or (destination.exists() and destination.read_bytes() != data):
            raise ValueError("El destino ya tiene otro logo; no se sobrescribe: " + filename)
    for _, filename, data in logos:
        (output / filename).write_bytes(data)
    manifest = [{"brand": brand, "file": filename, "sha256": hashlib.sha256(data).hexdigest()} for brand, filename, data in logos]
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(json.dumps({"logos": extract(args.workbook, args.output)}, ensure_ascii=True))
