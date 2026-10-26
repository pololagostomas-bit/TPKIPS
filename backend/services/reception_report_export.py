"""The filtered indicator grid, not the full audit workbook."""
from io import BytesIO
import json
from urllib.parse import parse_qs
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
from backend.services import reception

DEFAULTS={'date_from':'','date_to':'','arrival':'ALL','sla':'ALL','state':'ALL',
          'transport':'ALL','search':'','sort':'default','direction':'asc'}


def normalize(options=None):
    options=options or {}
    if not isinstance(options,dict) or set(options)-set(DEFAULTS):raise ValueError('Filtro de reporte invalido')
    values={**DEFAULTS,**options}
    if any(not isinstance(value,str) for value in values.values()):raise ValueError('Filtro de reporte invalido')
    values['search']=values['search'].strip().casefold()
    reception.filter_reception_report({'items':[]},**values)
    return values


def query_options(query):
    parsed=parse_qs(query)
    return normalize({key:parsed.get({'date_from':'from','date_to':'to'}.get(key,key),[default])[0]
                      for key,default in DEFAULTS.items()})


def build(connection,options):
    report=reception.filter_reception_report(reception.reception_report(connection),limit=None,**normalize(options))
    book=Workbook(write_only=True)
    summary=book.create_sheet('Resumen')
    sheet=book.create_sheet('Indicadores BL')
    fill=PatternFill('solid',fgColor='3E4A52')
    font=Font(bold=True,color='FFFFFF')
    def append(target,values,header=False):
        cells=[]
        for value in values:
            cell=WriteOnlyCell(target,value='' if value is None else value)
            if isinstance(value,str):cell.data_type='s'
            if header:cell.fill=fill;cell.font=font
            cells.append(cell)
        target.append(cells)
    append(summary,['Triton WMS - Indicadores de recepcion'],True)
    append(summary,['Generado',report['generated_at']])
    append(summary,['BL en la base',report['summary']['total_bl']])
    append(summary,['BL del resultado filtrado',report['total']])
    append(summary,['BL parciales (segun filtros)',report['filtered_summary']['partial_bl']])
    append(summary,['BL rojas (segun filtros)',report['filtered_summary']['semaforizacion']['ROJO']])
    append(summary,['BL amarillas (segun filtros)',report['filtered_summary']['semaforizacion']['AMARILLO']])
    append(summary,['Filtros',json.dumps(normalize(options),ensure_ascii=False)])
    fields=[('BL / AWB','bl_awb'),('Camion','truck_guides'),('SKUs','sku_count'),
            ('Unidades','total_units'),('Etapa','app_status'),('Transporte','transport_type'),
            ('Inicio','report_started_at'),('Asistente','current_assistant'),('Auxiliar','current_auxiliary'),
            ('Fin','report_closed_at'),('Horas transcurridas','elapsed_hours'),('Horas SLA','sla_hours'),
            ('Horas restantes','remaining_hours'),('SLA','sla_status'),('Entregas','delivery_count')]
    sheet.freeze_panes='A2'
    for index in range(len(fields)):
        sheet.column_dimensions[chr(65+index)].width=24
    append(sheet,[name for name,key in fields],True)
    for item in report['items']:append(sheet,[item.get(key) for name,key in fields])
    sheet.auto_filter.ref=f'A1:{get_column_letter(len(fields))}{len(report["items"])+1}'
    output=BytesIO()
    book.save(output)
    return output.getvalue()
