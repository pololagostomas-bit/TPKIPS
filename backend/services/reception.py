"""Dominio inicial del módulo Recepción de TRITON WMS.

El módulo comparte la base SQLite del piloto, pero usa tablas y estados propios.
Esto permite desplegarlo junto a Despacho sin mezclar sus tiempos ni historiales.
"""

import hashlib
import json
import math
import os
import re
import unicodedata
import warnings
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from backend.services.daily_operations import local_now, advanced_lots_enabled

from backend.services.traceability import (
    generate_lots_for_shipment,
    lots_for_shipment,
    sync_shipment_lot_references,
)


RECEPTION_ROLES = {"ADMINISTRADOR", "ASISTENTE_RECEPCION", "AUXILIAR_RECEPCION"}
RECEPTION_STATES = (
    "PROGRAMADO",
    "ARRIBADO",
    "REVISION SISTEMA",
    "EM",
    "UBICACION",
    "VALIDACION",
    "SOLICITUD TRANSFERENCIA",
    "CERRADO",
)
RECEPTION_STATE_INDEX = {state: index for index, state in enumerate(RECEPTION_STATES)}


def reception_now():
    return local_now()


def init_reception_schema(connection):
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS reception_shipments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            bl_awb TEXT NOT NULL UNIQUE,
            transport_type TEXT NOT NULL DEFAULT 'AEREO',
            supplier TEXT,
            brand_summary TEXT,
            country_origin TEXT,
            source_sheets TEXT,
            ip_reference TEXT,
            primary_oc TEXT,
            primary_ov TEXT,
            expected_packages REAL NOT NULL DEFAULT 0,
            received_packages REAL NOT NULL DEFAULT 0,
            app_status TEXT NOT NULL DEFAULT 'PROGRAMADO',
            condition_status TEXT NOT NULL DEFAULT 'POR ARRIBAR',
            current_assistant TEXT,
            current_auxiliary TEXT,
            fr_number TEXT,
            em_number TEXT,
            accounting_status TEXT NOT NULL DEFAULT 'PENDIENTE CONTABILIDAD',
            accounting_conflict_count INTEGER NOT NULL DEFAULT 0,
            accounting_conflict_detail TEXT,
            accounting_email_date TEXT,
            source_reception_date TEXT,
            em_date TEXT,
            costing_date TEXT,
            scheduled_date TEXT,
            first_arrival_at TEXT,
            completed_arrival_at TEXT,
            physical_initialized INTEGER NOT NULL DEFAULT 0,
            system_quantities_initialized INTEGER NOT NULL DEFAULT 0,
            bultos_closed INTEGER NOT NULL DEFAULT 0,
            location_text TEXT,
            validation_sample_json TEXT,
            transfer_assistant_checked INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS reception_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL REFERENCES reception_shipments(id) ON DELETE CASCADE,
            source_row INTEGER,
            source_sheet TEXT,
            source_key TEXT,
            np_code TEXT,
            description TEXT,
            oc_number TEXT,
            ov_number TEXT,
            expected_qty REAL NOT NULL DEFAULT 0,
            requested_qty REAL NOT NULL DEFAULT 0,
            invoiced_qty REAL NOT NULL DEFAULT 0,
            pending_qty REAL NOT NULL DEFAULT 0,
            received_qty REAL NOT NULL DEFAULT 0,
            purchase_type TEXT,
            brand TEXT,
            applicant TEXT,
            source_status TEXT,
            source_invoice TEXT,
            ip_reference TEXT,
            is_replacement INTEGER NOT NULL DEFAULT 0,
            replacement_np TEXT,
            blocked_reason TEXT
        );

        CREATE TABLE IF NOT EXISTS reception_receipts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL REFERENCES reception_shipments(id) ON DELETE CASCADE,
            sequence_no INTEGER NOT NULL,
            received_packages REAL NOT NULL,
            notes TEXT,
            location_text TEXT,
            username TEXT NOT NULL,
            received_at TEXT NOT NULL,
            UNIQUE(shipment_id, sequence_no)
        );

        -- Una BL puede llegar en más de una recepción. Cada llegada que se
        -- trabaja después de cerrar la anterior tiene su propia atención,
        -- responsable, estado y cantidades remanentes.
        CREATE TABLE IF NOT EXISTS reception_attentions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL REFERENCES reception_shipments(id) ON DELETE CASCADE,
            receipt_id INTEGER REFERENCES reception_receipts(id) ON DELETE SET NULL,
            sequence_no INTEGER NOT NULL,
            app_status TEXT NOT NULL DEFAULT 'ARRIBADO',
            condition_status TEXT NOT NULL DEFAULT 'EN PROCESO',
            current_assistant TEXT,
            current_auxiliary TEXT,
            location_text TEXT,
            system_quantities_initialized INTEGER NOT NULL DEFAULT 0,
            transfer_assistant_checked INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            completed_at TEXT,
            UNIQUE(shipment_id, sequence_no),
            UNIQUE(receipt_id)
        );

        CREATE TABLE IF NOT EXISTS reception_attention_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            attention_id INTEGER NOT NULL REFERENCES reception_attentions(id) ON DELETE CASCADE,
            reception_line_id INTEGER NOT NULL REFERENCES reception_lines(id) ON DELETE CASCADE,
            planned_qty REAL NOT NULL DEFAULT 0,
            verified_qty REAL NOT NULL DEFAULT 0,
            UNIQUE(attention_id, reception_line_id)
        );

        CREATE TABLE IF NOT EXISTS reception_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL REFERENCES reception_shipments(id) ON DELETE CASCADE,
            event_type TEXT NOT NULL,
            field_name TEXT NOT NULL,
            old_value TEXT,
            new_value TEXT,
            username TEXT NOT NULL,
            reason TEXT,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS reception_accounting_refs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL REFERENCES reception_shipments(id) ON DELETE CASCADE,
            source_key TEXT NOT NULL UNIQUE,
            source_sheet TEXT,
            source_row INTEGER,
            awb_reference TEXT,
            ip_reference TEXT,
            fr_number TEXT,
            em_number TEXT,
            email_date TEXT,
            reception_date TEXT,
            em_date TEXT,
            costing_date TEXT,
            import_status TEXT,
            source_file TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS reception_notifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_id INTEGER NOT NULL REFERENCES reception_shipments(id) ON DELETE CASCADE,
            notification_type TEXT NOT NULL,
            recipient TEXT NOT NULL,
            subject TEXT NOT NULL,
            body TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'PENDIENTE ENVIO',
            created_at TEXT NOT NULL,
            sent_at TEXT,
            error TEXT
        );

        CREATE INDEX IF NOT EXISTS idx_reception_shipments_status
            ON reception_shipments(app_status);
        CREATE INDEX IF NOT EXISTS idx_reception_lines_np
            ON reception_lines(np_code);
        CREATE INDEX IF NOT EXISTS idx_reception_accounting_shipment
            ON reception_accounting_refs(shipment_id);
        CREATE INDEX IF NOT EXISTS idx_reception_notifications_shipment
            ON reception_notifications(shipment_id);
        """
    )
    _ensure_column(connection, "reception_shipments", "brand_summary", "TEXT")
    _ensure_column(connection, "reception_shipments", "country_origin", "TEXT")
    _ensure_column(connection, "reception_shipments", "source_sheets", "TEXT")
    _ensure_column(connection, "reception_shipments", "physical_initialized", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_shipments", "system_quantities_initialized", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_shipments", "bultos_closed", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_shipments", "location_text", "TEXT")
    _ensure_column(connection, "reception_shipments", "validation_sample_json", "TEXT")
    _ensure_column(connection, "reception_shipments", "transfer_assistant_checked", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_shipments", "accounting_email_date", "TEXT")
    _ensure_column(connection, "reception_shipments", "source_reception_date", "TEXT")
    _ensure_column(connection, "reception_shipments", "em_date", "TEXT")
    _ensure_column(connection, "reception_shipments", "costing_date", "TEXT")
    _ensure_column(
        connection,
        "reception_shipments",
        "accounting_status",
        "TEXT NOT NULL DEFAULT 'PENDIENTE CONTABILIDAD'",
    )
    _ensure_column(
        connection,
        "reception_shipments",
        "accounting_conflict_count",
        "INTEGER NOT NULL DEFAULT 0",
    )
    _ensure_column(connection, "reception_shipments", "accounting_conflict_detail", "TEXT")
    _ensure_column(connection, "reception_receipts", "location_text", "TEXT")
    _ensure_column(connection, "reception_lines", "source_sheet", "TEXT")
    _ensure_column(connection, "reception_lines", "source_key", "TEXT")
    _ensure_column(connection, "reception_lines", "requested_qty", "REAL NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_lines", "invoiced_qty", "REAL NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_lines", "pending_qty", "REAL NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_lines", "purchase_type", "TEXT")
    _ensure_column(connection, "reception_lines", "brand", "TEXT")
    _ensure_column(connection, "reception_lines", "applicant", "TEXT")
    _ensure_column(connection, "reception_lines", "source_status", "TEXT")
    _ensure_column(connection, "reception_lines", "source_invoice", "TEXT")
    _ensure_column(connection, "reception_lines", "ip_reference", "TEXT")
    _ensure_column(connection, "reception_lines", "validation_sap_checked", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_lines", "validation_location_checked", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_lines", "validation_comment_checked", "INTEGER NOT NULL DEFAULT 0")
    _ensure_column(connection, "reception_lines", "validation_comment", "TEXT")
    connection.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_reception_lines_source_key ON reception_lines(source_key)"
    )
    legacy_physical_review = connection.execute(
        """SELECT id, app_status, current_assistant, current_auxiliary
           FROM reception_shipments
           WHERE app_status = 'REVISION FISICA'"""
    ).fetchall()
    for shipment in legacy_physical_review:
        has_assistant = bool(str(shipment["current_assistant"] or "").strip())
        has_auxiliary = bool(str(shipment["current_auxiliary"] or "").strip())
        target_status = "REVISION SISTEMA" if has_assistant and has_auxiliary else "ARRIBADO"
        target_condition = "EN PROCESO" if target_status == "REVISION SISTEMA" else "ARRIBO REGISTRADO"
        connection.execute(
            "UPDATE reception_shipments SET app_status = ?, condition_status = ?, updated_at = ? WHERE id = ?",
            (target_status, target_condition, reception_now(), shipment["id"]),
        )
        _write_history(
            connection,
            shipment["id"],
            "MIGRACION",
            "app_status",
            "REVISION FISICA",
            target_status,
            "sistema.migracion",
            "Se retiró REVISION FISICA como estado; el conteo físico continúa fuera del app y el sistema valida en REVISION SISTEMA",
        )
    legacy_system_review = connection.execute(
        """SELECT id, physical_initialized FROM reception_shipments
           WHERE system_quantities_initialized = 0
             AND app_status IN ('REVISION SISTEMA', 'EM', 'UBICACION', 'VALIDACION',
                                'SOLICITUD TRANSFERENCIA', 'CERRADO')"""
    ).fetchall()
    for shipment in legacy_system_review:
        migration_value = "PRESERVADAS" if shipment["physical_initialized"] else "100%"
        if not shipment["physical_initialized"]:
            connection.execute(
                "UPDATE reception_lines SET received_qty = expected_qty WHERE shipment_id = ?",
                (shipment["id"],),
            )
        connection.execute(
            "UPDATE reception_shipments SET system_quantities_initialized = 1 WHERE id = ?",
            (shipment["id"],),
        )
        _write_history(
            connection,
            shipment["id"],
            "MIGRACION",
            "cantidades_revision_sistema",
            0,
            migration_value,
            "sistema.migracion",
            "Habilitación de cantidades para expedientes que ya habían iniciado revisión de sistema",
        )
    for shipment in connection.execute("SELECT id FROM reception_shipments").fetchall():
        _refresh_accounting_summary(
            connection,
            shipment["id"],
            "sistema.migracion",
            "recalculo de estados contables al iniciar",
        )
    _migrate_reception_attentions(connection)


def _ensure_column(connection, table, column, definition):
    columns = {row["name"] for row in connection.execute(f"PRAGMA table_info({table})")}
    if column not in columns:
        connection.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def _attention_rows(connection, shipment_id):
    return connection.execute(
        "SELECT * FROM reception_attentions WHERE shipment_id = ? ORDER BY sequence_no",
        (shipment_id,),
    ).fetchall()


def _active_reception_attention(connection, shipment_id):
    """Return the current arrival attention without changing legacy data."""
    rows = _attention_rows(connection, shipment_id)
    if not rows:
        return None
    pending = [row for row in rows if row["app_status"] != "CERRADO"]
    return (pending[-1] if pending else rows[-1])


def _attention_verified_total(connection, reception_line_id):
    row = connection.execute(
        """SELECT COALESCE(SUM(verified_qty), 0) AS total
             FROM reception_attention_lines
            WHERE reception_line_id = ?""",
        (reception_line_id,),
    ).fetchone()
    return float(row["total"] or 0) if row else 0.0


def _create_reception_attention(connection, shipment, receipt_id, username,
                                initial_status="ARRIBADO"):
    """Create an attention for a later arrival using the unverified balance.

    The first receipt creates Atención 1. A later receipt creates Atención 2
    only after the previous attention was closed, so two receipts entered on
    the same day are still treated as one operational task.
    """
    shipment_id = int(shipment["id"])
    sequence = connection.execute(
        "SELECT COALESCE(MAX(sequence_no), 0) + 1 FROM reception_attentions WHERE shipment_id = ?",
        (shipment_id,),
    ).fetchone()[0]
    timestamp = reception_now()
    cursor = connection.execute(
        """INSERT INTO reception_attentions
           (shipment_id, receipt_id, sequence_no, app_status, condition_status,
            current_assistant, current_auxiliary, created_at, updated_at)
           VALUES (?, ?, ?, ?, 'EN PROCESO', ?, ?, ?, ?)""",
        (
            shipment_id,
            receipt_id,
            sequence,
            initial_status,
            shipment["current_assistant"],
            shipment["current_auxiliary"],
            timestamp,
            timestamp,
        ),
    )
    attention_id = cursor.lastrowid
    for line in connection.execute(
        "SELECT * FROM reception_lines WHERE shipment_id = ? ORDER BY id", (shipment_id,)
    ).fetchall():
        previous = _attention_verified_total(connection, line["id"])
        planned = max(0.0, float(line["expected_qty"] or 0) - previous)
        connection.execute(
            """INSERT INTO reception_attention_lines
               (attention_id, reception_line_id, planned_qty, verified_qty)
               VALUES (?, ?, ?, 0)""",
            (attention_id, line["id"], planned),
        )
    _write_history(
        connection,
        shipment_id,
        "ATENCION",
        "attention",
        "",
        f"Atención {sequence} · recepción {receipt_id}",
        username,
        "Nueva llegada física separada de la atención anterior",
    )
    return connection.execute(
        "SELECT * FROM reception_attentions WHERE id = ?", (attention_id,)
    ).fetchone()


def _migrate_reception_attentions(connection):
    """Add attention tracking to existing databases without deleting data."""
    shipments = connection.execute(
        "SELECT * FROM reception_shipments WHERE EXISTS "
        "(SELECT 1 FROM reception_receipts r WHERE r.shipment_id = reception_shipments.id)"
    ).fetchall()
    for shipment in shipments:
        if connection.execute(
            "SELECT 1 FROM reception_attentions WHERE shipment_id = ? LIMIT 1",
            (shipment["id"],),
        ).fetchone():
            continue
        receipts = connection.execute(
            "SELECT * FROM reception_receipts WHERE shipment_id = ? ORDER BY sequence_no",
            (shipment["id"],),
        ).fetchall()
        for index, receipt in enumerate(receipts):
            status = shipment["app_status"] if index == len(receipts) - 1 else "CERRADO"
            timestamp = receipt["received_at"] or reception_now()
            cursor = connection.execute(
                """INSERT INTO reception_attentions
                   (shipment_id, receipt_id, sequence_no, app_status, condition_status,
                    current_assistant, current_auxiliary, created_at, updated_at, completed_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    shipment["id"], receipt["id"], receipt["sequence_no"], status,
                    "COMPLETADO" if status == "CERRADO" else "EN PROCESO",
                    shipment["current_assistant"], shipment["current_auxiliary"],
                    timestamp, timestamp, timestamp if status == "CERRADO" else None,
                ),
            )
            attention_id = cursor.lastrowid
            for line in connection.execute(
                "SELECT * FROM reception_lines WHERE shipment_id = ? ORDER BY id",
                (shipment["id"],),
            ).fetchall():
                planned = float(line["expected_qty"] or 0) if index == 0 else 0.0
                verified = float(line["received_qty"] or 0) if status == "CERRADO" and index == len(receipts) - 1 else 0.0
                connection.execute(
                    """INSERT INTO reception_attention_lines
                       (attention_id, reception_line_id, planned_qty, verified_qty)
                       VALUES (?, ?, ?, ?)""",
                    (attention_id, line["id"], planned, verified),
                )


def _ensure_reception_attention(connection, shipment, username):
    """Create a compatibility attention for older rows that have no receipt."""
    active = _active_reception_attention(connection, shipment["id"])
    if active:
        return active
    if shipment["app_status"] == "PROGRAMADO" and not float(shipment["received_packages"] or 0):
        return None
    return _create_reception_attention(connection, shipment, None, username, shipment["app_status"])


def _sync_reception_line_totals(connection, shipment_id):
    """Mirror the sum of all arrival-attention quantities for legacy reports."""
    for line in connection.execute(
        "SELECT id FROM reception_lines WHERE shipment_id = ?", (shipment_id,)
    ).fetchall():
        total = _attention_verified_total(connection, line["id"])
        connection.execute(
            "UPDATE reception_lines SET received_qty = ? WHERE id = ?",
            (total, line["id"]),
        )


def require_reception_access(role):
    if role not in RECEPTION_ROLES:
        raise PermissionError("Tu usuario no tiene acceso al módulo Recepción")


def _is_assigned_to(shipment, username):
    current_user = str(username or "").strip().casefold()
    if not current_user:
        return False
    return current_user in {
        str(shipment["current_assistant"] or "").strip().casefold(),
        str(shipment["current_auxiliary"] or "").strip().casefold(),
    }


def _default_reception_user(connection, role, setting_name):
    """Resolve the default responsible without hard-coding a person.

    TI can pin the operational users with environment variables. If they are
    not configured, the first active user of the exact reception role is used.
    This keeps imported BLs out of an unassigned queue while preserving the
    administrator's ability to change the assignment later.
    """
    users_table = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'users'"
    ).fetchone()
    if not users_table:
        return ""
    configured = str(os.getenv(setting_name, "") or "").strip().lower()
    if configured:
        selected = connection.execute(
            """SELECT username FROM users
               WHERE lower(username) = ? AND active = 1 AND role = ?
               LIMIT 1""",
            (configured, role),
        ).fetchone()
        if selected:
            return selected["username"]
    selected = connection.execute(
        """SELECT username FROM users
           WHERE active = 1 AND role = ?
           ORDER BY datetime(created_at) ASC, username ASC
           LIMIT 1""",
        (role,),
    ).fetchone()
    return selected["username"] if selected else ""


def _ensure_default_reception_assignees(connection, shipment_id, username,
                                        reason="Asignación automática de recepción"):
    """Assign the default assistant and auxiliary only when a slot is empty."""
    shipment = connection.execute(
        "SELECT current_assistant, current_auxiliary FROM reception_shipments WHERE id = ?",
        (shipment_id,),
    ).fetchone()
    if not shipment:
        return
    changes = {}
    if not str(shipment["current_assistant"] or "").strip():
        changes["current_assistant"] = _default_reception_user(
            connection, "ASISTENTE_RECEPCION", "TRITON_DEFAULT_RECEPTION_ASSISTANT"
        )
    if not str(shipment["current_auxiliary"] or "").strip():
        changes["current_auxiliary"] = _default_reception_user(
            connection, "AUXILIAR_RECEPCION", "TRITON_DEFAULT_RECEPTION_AUXILIARY"
        )
    for field, value in changes.items():
        if not value:
            continue
        old_value = str(shipment[field] or "")
        connection.execute(
            f"UPDATE reception_shipments SET {field} = ?, updated_at = ? WHERE id = ?",
            (value, reception_now(), shipment_id),
        )
        _write_history(connection, shipment_id, "ASIGNACION AUTOMATICA", field,
                       old_value, value, username, reason)


def _get_reception_shipment(connection, shipment_id, username, role):
    require_reception_access(role)
    shipment = connection.execute(
        "SELECT * FROM reception_shipments WHERE id = ?", (shipment_id,)
    ).fetchone()
    if not shipment:
        raise ValueError("BL/AWB no encontrada")
    if role != "ADMINISTRADOR" and not _is_assigned_to(shipment, username):
        raise PermissionError("Esta BL/AWB no está asignada a tu usuario")
    return shipment


def _as_dict(row):
    return dict(row) if row else None


def _write_history(connection, shipment_id, event_type, field_name, old_value,
                   new_value, username, reason=""):
    connection.execute(
        """INSERT INTO reception_history
           (shipment_id, event_type, field_name, old_value, new_value,
            username, reason, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            shipment_id,
            event_type,
            field_name,
            "" if old_value is None else str(old_value),
            "" if new_value is None else str(new_value),
            username,
            reason,
            reception_now(),
        ),
    )


def _notification_recipients(setting_name, default):
    configured = os.getenv(setting_name, default)
    return [item.strip() for item in configured.split(",") if item.strip()]


def _enqueue_notification(connection, shipment_id, notification_type, subject, body,
                          setting_name, default_recipient):
    recipients = _notification_recipients(setting_name, default_recipient)
    for recipient in recipients:
        duplicate = connection.execute(
            """SELECT 1 FROM reception_notifications
               WHERE shipment_id = ? AND notification_type = ?
                 AND recipient = ? AND subject = ? AND body = ?
               LIMIT 1""",
            (shipment_id, notification_type, recipient, subject, body),
        ).fetchone()
        if duplicate:
            continue
        connection.execute(
            """INSERT INTO reception_notifications
               (shipment_id, notification_type, recipient, subject, body, status, created_at)
               VALUES (?, ?, ?, ?, ?, 'PENDIENTE ENVIO', ?)""",
            (shipment_id, notification_type, recipient, subject, body, reception_now()),
        )


def _ip_tokens(value):
    """Devuelve IP individuales sin separar guiones que forman parte del código."""
    raw = str(value or "").replace("\r", "\n")
    tokens = re.split(r"[,;\n]|\s+-\s+", raw)
    result = []
    for token in tokens:
        normalized = _normalized_identifier(token)
        if normalized and normalized not in result:
            result.append(normalized)
    return result


def _accounting_rows_by_ip(rows):
    by_ip = {}
    for row in rows:
        for ip_key in _ip_tokens(row["ip_reference"]):
            by_ip.setdefault(ip_key, []).append(row)
    return by_ip


def _shipment_expected_ips(connection, shipment_id, shipment=None):
    if shipment is None:
        shipment = connection.execute(
            "SELECT * FROM reception_shipments WHERE id = ?", (shipment_id,)
        ).fetchone()
    line_ips = set()
    for row in connection.execute(
        "SELECT ip_reference FROM reception_lines WHERE shipment_id = ?", (shipment_id,)
    ).fetchall():
        line_ips.update(_ip_tokens(row["ip_reference"]))
    return line_ips or set(_ip_tokens(shipment["ip_reference"] if shipment else ""))


def _accounting_status_for_shipment(connection, shipment_id, shipment=None, rows=None):
    if shipment is None:
        shipment = connection.execute(
            "SELECT * FROM reception_shipments WHERE id = ?", (shipment_id,)
        ).fetchone()
    if rows is None:
        rows = connection.execute(
            "SELECT * FROM reception_accounting_refs WHERE shipment_id = ?",
            (shipment_id,),
        ).fetchall()
    if int(shipment["accounting_conflict_count"] or 0):
        return "CONFLICTO IP"
    by_ip = _accounting_rows_by_ip(rows)
    expected_ips = _shipment_expected_ips(connection, shipment_id, shipment)
    if not rows and str(shipment["fr_number"] or "").strip():
        return "EM REGISTRADA" if str(shipment["em_number"] or "").strip() else "PENDIENTE EM"
    if not expected_ips:
        if not rows:
            if not str(shipment["fr_number"] or "").strip():
                return "PENDIENTE CONTABILIDAD"
            return "EM REGISTRADA" if str(shipment["em_number"] or "").strip() else "PENDIENTE EM"
        if rows and any(not str(row["fr_number"] or "").strip() for row in rows):
            return "PENDIENTE FR"
        if rows and any(not str(row["em_number"] or "").strip() for row in rows):
            return "PENDIENTE EM"
        return "EM REGISTRADA" if str(shipment["em_number"] or "").strip() or rows else "PENDIENTE CONTABILIDAD"
    missing_fr = 0
    missing_em = 0
    any_fr = False
    for ip_key in expected_ips:
        refs = by_ip.get(ip_key, [])
        has_fr = any(str(row["fr_number"] or "").strip() for row in refs)
        has_em = any(str(row["em_number"] or "").strip() for row in refs)
        any_fr = any_fr or has_fr
        if not has_fr:
            missing_fr += 1
        elif not has_em:
            missing_em += 1
    if missing_fr:
        return "PENDIENTE FR" if any_fr else "PENDIENTE CONTABILIDAD"
    if missing_em:
        return "PENDIENTE EM"
    return "EM REGISTRADA"


def _line_has_complete_accounting(line, accounting_by_ip):
    line_ips = _ip_tokens(line["ip_reference"])
    if not line_ips:
        return False
    return all(
        any(
            str(row["fr_number"] or "").strip()
            and str(row["em_number"] or "").strip()
            for row in accounting_by_ip.get(ip_key, [])
        )
        for ip_key in line_ips
    )


def _accounting_warning(accounting_status):
    if accounting_status in {"PENDIENTE CONTABILIDAD", "PENDIENTE FR"}:
        return (
            "Falta factura de reserva para una o más referencias. "
            "Puedes registrar el arribo, asignar responsables y revisar el sistema; "
            "debes completar la FR antes de entrar a EM."
        )
    if accounting_status == "CONFLICTO IP":
        return (
            "Hay una IP relacionada con más de una FR/EM. "
            "Puedes continuar los pasos previos a EM; corrige el conflicto antes de entrar a EM."
        )
    return ""


def list_receptions(connection, search="", role="ADMINISTRADOR", username="", limit=10,
                    arrival_date="", arrival_date_end="", state=""):
    require_reception_access(role)
    query = """SELECT s.*,
                       (SELECT GROUP_CONCAT(DISTINCT r.sap_ov)
                          FROM order_importation_refs r
                         WHERE lower(COALESCE(r.bl_awb, '')) = lower(COALESCE(s.bl_awb, ''))) AS linked_ovs,
                       (SELECT GROUP_CONCAT(DISTINCT r.transport_type)
                          FROM order_importation_refs r
                         WHERE lower(COALESCE(r.bl_awb, '')) = lower(COALESCE(s.bl_awb, ''))) AS linked_transport,
                       (SELECT COUNT(*) FROM reception_lines l WHERE l.shipment_id = s.id) AS line_count,
                       (SELECT COUNT(*) FROM reception_receipts r WHERE r.shipment_id = s.id) AS receipt_count,
                       (SELECT COUNT(*) FROM reception_attentions a WHERE a.shipment_id = s.id) AS attention_count,
                       (SELECT MAX(a.sequence_no) FROM reception_attentions a
                         WHERE a.shipment_id = s.id AND a.app_status <> 'CERRADO') AS active_attention_sequence,
                       (SELECT COUNT(*) FROM reception_receipts r
                         WHERE r.shipment_id = s.id AND TRIM(COALESCE(r.location_text, '')) = '') AS location_missing_count
               FROM reception_shipments s"""
    conditions = []
    params = []
    if arrival_date:
        from datetime import date
        arrival_date = date.fromisoformat(str(arrival_date)).isoformat()
        if arrival_date_end:
            arrival_date_end = date.fromisoformat(str(arrival_date_end)).isoformat()
            if arrival_date_end < arrival_date:
                raise ValueError("La fecha final no puede ser anterior a la fecha inicial")
            conditions.append("substr(COALESCE(s.scheduled_date, ''), 1, 10) BETWEEN ? AND ?")
            params.extend([arrival_date, arrival_date_end])
        else:
            conditions.append("substr(COALESCE(s.scheduled_date, ''), 1, 10) = ?")
            params.append(arrival_date)
    elif arrival_date_end:
        from datetime import date
        arrival_date_end = date.fromisoformat(str(arrival_date_end)).isoformat()
        conditions.append("substr(COALESCE(s.scheduled_date, ''), 1, 10) <= ?")
        params.append(arrival_date_end)
    if state and state != "ALL":
        conditions.append("s.app_status = ?")
        params.append(state)
    normalized = _normalized_identifier(search)
    if normalized:
        # BL/AWB no tiene índice intencionalmente: sus formatos cambian entre
        # courier, aéreo y marítimo. Para una búsqueda numérica se usa la
        # clave normalizada y sus últimos cuatro caracteres.
        if normalized.isdigit() and len(normalized) >= 4:
            tail = f"%{normalized[-4:]}%"
            conditions.append("""REPLACE(REPLACE(REPLACE(REPLACE(UPPER(COALESCE(s.bl_awb, '')), '-', ''), ' ', ''), '/', ''), '.', '') LIKE ?""")
            params.append(tail)
        else:
            token = f"%{str(search or '').strip().casefold()}%"
            conditions.append("""(lower(COALESCE(s.bl_awb, '')) LIKE ?
                             OR lower(COALESCE(s.supplier, '')) LIKE ?
                             OR lower(COALESCE(s.ip_reference, '')) LIKE ?
                             OR lower(COALESCE(s.primary_oc, '')) LIKE ?
                             OR lower(COALESCE(s.primary_ov, '')) LIKE ?
                             OR lower(COALESCE(s.fr_number, '')) LIKE ?
                             OR lower(COALESCE(s.em_number, '')) LIKE ?
                             OR EXISTS (
                                 SELECT 1 FROM order_importation_refs r
                                 WHERE lower(COALESCE(r.bl_awb, '')) = lower(COALESCE(s.bl_awb, ''))
                                   AND (lower(COALESCE(r.sap_ov, '')) LIKE ?
                                        OR lower(COALESCE(r.bl_awb, '')) LIKE ?
                                        OR lower(COALESCE(r.ip_reference, '')) LIKE ?
                                        OR lower(COALESCE(r.oc_number, '')) LIKE ?)
                             )
                             OR EXISTS (
                                SELECT 1 FROM reception_accounting_refs a
                                WHERE a.shipment_id = s.id
                                  AND (lower(COALESCE(a.awb_reference, '')) LIKE ?
                                       OR lower(COALESCE(a.ip_reference, '')) LIKE ?
                                       OR lower(COALESCE(a.fr_number, '')) LIKE ?
                                       OR lower(COALESCE(a.em_number, '')) LIKE ?)
                            )
                             OR EXISTS (
                                SELECT 1 FROM reception_lines l
                                WHERE l.shipment_id = s.id
                                  AND lower(COALESCE(l.np_code, '')) LIKE ?
                             ))""")
            params.extend([token] * 16)
    if role != "ADMINISTRADOR" and username:
        current_user = str(username).strip().casefold()
        conditions.append("""(lower(COALESCE(s.current_assistant, '')) = ?
                              OR lower(COALESCE(s.current_auxiliary, '')) = ?)""")
        params.extend([current_user, current_user])
    if conditions:
        query += " WHERE " + " AND ".join(conditions)
    query += """ ORDER BY CASE
                             -- FR completa y pendiente de EM: trabajo listo para avanzar.
                             WHEN s.accounting_status = 'PENDIENTE EM' THEN 0
                             -- Próximas llegadas: se muestran antes que el resto de la cola.
                             WHEN s.app_status = 'PROGRAMADO' AND COALESCE(s.scheduled_date, '') <> '' THEN 1
                             WHEN s.app_status = 'PROGRAMADO' THEN 2
                             -- Referencias ya cerradas contablemente, pero aún visibles para seguimiento.
                             WHEN s.accounting_status = 'EM REGISTRADA' THEN 3
                             WHEN s.app_status = 'CERRADO' THEN 5
                             ELSE 4
                         END,
                         CASE WHEN s.accounting_status = 'PENDIENTE EM' THEN COALESCE(s.scheduled_date, '9999-12-31') END ASC,
                         CASE WHEN s.app_status = 'PROGRAMADO' THEN COALESCE(s.scheduled_date, '9999-12-31') END ASC,
                         s.updated_at DESC
                         LIMIT ?"""
    try:
        # El límite aplica después de búsqueda, fecha, etapa y permisos.
        safe_limit = max(1, min(int(limit), 10))
    except (TypeError, ValueError):
        safe_limit = 10
    params.append(safe_limit)
    shipment_rows = [_as_dict(row) for row in connection.execute(query, params).fetchall()]
    rows = []
    for shipment_row in shipment_rows:
        shipment_id = int(shipment_row["id"])
        attention_rows = connection.execute(
            """SELECT id, sequence_no, app_status, condition_status,
                      current_assistant, current_auxiliary, receipt_id
                 FROM reception_attentions
                WHERE shipment_id = ?
                ORDER BY sequence_no""",
            (shipment_id,),
        ).fetchall()
        attention_count = len(attention_rows)
        active_id = None
        active_attention = _active_reception_attention(connection, shipment_id)
        if active_attention:
            active_id = int(active_attention["id"])

        # Una BL puede recibirse en varias llegadas. Cada llegada es una
        # atención independiente en la cola para que el equipo pueda ver,
        # por ejemplo, Atención 1/2 y Atención 2/2. El límite SQL se aplica
        # primero a BLs, por lo que una BL con dos atenciones no desplaza a
        # otras BLs del corte inicial.
        queue_items = attention_rows or [None]
        for attention in queue_items:
            item = dict(shipment_row)
            item["accounting_warning"] = _accounting_warning(item.get("accounting_status"))
            if attention is None:
                item["attention_id"] = None
                item["attention_sequence"] = None
                item["attention_total"] = 0
                item["attention_label"] = ""
                item["attention_is_active"] = False
            else:
                attention_id = int(attention["id"])
                item["attention_id"] = attention_id
                item["attention_sequence"] = int(attention["sequence_no"])
                item["attention_total"] = attention_count
                item["attention_label"] = (
                    f"Atención {int(attention['sequence_no'])}/{attention_count}"
                )
                item["attention_is_active"] = attention_id == active_id
                # La cola representa el trabajo de esta llegada, no el
                # resumen histórico de la BL completa.
                item["app_status"] = attention["app_status"]
                item["condition_status"] = attention["condition_status"]
                item["current_assistant"] = attention["current_assistant"] or item.get("current_assistant")
                item["current_auxiliary"] = attention["current_auxiliary"] or item.get("current_auxiliary")
            if not state or state == "ALL" or item.get("app_status") == state:
                rows.append(item)
    return rows


def reception_links(connection, shipment_id, role="ADMINISTRADOR", username=""):
    """Carga bajo demanda la relación única OV/OC de una BL."""
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    rows = connection.execute(
        """SELECT DISTINCT ov_number, oc_number
             FROM reception_lines
            WHERE shipment_id = ?
              AND (TRIM(COALESCE(ov_number, '')) <> ''
                   OR TRIM(COALESCE(oc_number, '')) <> '')
            ORDER BY ov_number, oc_number""",
        (shipment_id,),
    ).fetchall()
    return {
        "shipment_id": shipment_id,
        "bl_awb": shipment["bl_awb"],
        "rows": [_as_dict(row) for row in rows],
    }


def _validation_sample_ids(connection, shipment_id):
    """Return a stable 30% sample of the shipment's NPs.

    The sample is persisted on the shipment so refreshing the page never changes
    what the assistant must validate.
    """
    lines = connection.execute(
        "SELECT id FROM reception_lines WHERE shipment_id = ? ORDER BY id", (shipment_id,)
    ).fetchall()
    line_ids = [int(row["id"]) for row in lines]
    shipment = connection.execute(
        "SELECT validation_sample_json FROM reception_shipments WHERE id = ?", (shipment_id,)
    ).fetchone()
    try:
        stored = [int(value) for value in json.loads(shipment["validation_sample_json"] or "[]")]
    except (TypeError, ValueError, json.JSONDecodeError):
        stored = []
    if stored and set(stored).issubset(set(line_ids)):
        sample = stored
    elif not line_ids:
        return []
    else:
        sample_size = max(1, math.ceil(len(line_ids) * 0.30))
        ranked = sorted(
            line_ids,
            key=lambda line_id: hashlib.sha256(f"{shipment_id}:{line_id}".encode()).hexdigest(),
        )
        sample = ranked[:sample_size]
        connection.execute(
            "UPDATE reception_shipments SET validation_sample_json = ? WHERE id = ?",
            (json.dumps(sample), shipment_id),
        )
    # La validación operativa solo requiere confirmar la ubicación. En una
    # muestra nueva el check parte marcado para que el usuario solo lo quite
    # si detecta una ubicación incorrecta; una decisión ya guardada no se
    # sobreescribe al refrescar la pantalla.
    for line_id in sample:
        saved = connection.execute(
            "SELECT 1 FROM reception_history WHERE shipment_id = ? "
            "AND field_name = ? LIMIT 1",
            (shipment_id, f"linea:{line_id}:muestra"),
        ).fetchone()
        if not saved:
            connection.execute(
                "UPDATE reception_lines SET validation_location_checked = 1 WHERE id = ?",
                (line_id,),
            )
    return sample


def reception_detail(connection, shipment_id, role="ADMINISTRADOR", username=""):
    require_reception_access(role)
    shipment = connection.execute(
        "SELECT * FROM reception_shipments WHERE id = ?", (shipment_id,)
    ).fetchone()
    if not shipment:
        return None
    if role != "ADMINISTRADOR" and username and not _is_assigned_to(shipment, username):
        raise PermissionError("Esta BL/AWB no está asignada a tu usuario")
    payload = _as_dict(shipment)
    payload["accounting_warning"] = _accounting_warning(payload.get("accounting_status"))
    linked = connection.execute(
        """SELECT GROUP_CONCAT(DISTINCT sap_ov) AS linked_ovs,
                         GROUP_CONCAT(DISTINCT transport_type) AS linked_transport
              FROM order_importation_refs
             WHERE lower(COALESCE(bl_awb, '')) = lower(COALESCE(?, ''))""",
        (shipment["bl_awb"],),
    ).fetchone()
    payload["linked_ovs"] = linked["linked_ovs"] if linked else ""
    payload["linked_transport"] = linked["linked_transport"] if linked else ""
    if payload["linked_transport"] and str(payload.get("transport_type") or "SIN DEFINIR") == "SIN DEFINIR":
        payload["transport_type"] = payload["linked_transport"]
    payload["lines"] = [
        _as_dict(row)
        for row in connection.execute(
            "SELECT * FROM reception_lines WHERE shipment_id = ? ORDER BY id", (shipment_id,)
        ).fetchall()
    ]
    attention_rows = _attention_rows(connection, shipment_id)
    active_attention = _active_reception_attention(connection, shipment_id)
    payload["active_attention_id"] = int(active_attention["id"]) if active_attention else None
    payload["active_attention_sequence"] = (
        int(active_attention["sequence_no"]) if active_attention else None
    )
    payload["attention_count"] = len(attention_rows)
    payload["attentions"] = []
    for attention in attention_rows:
        attention_payload = _as_dict(attention)
        attention_payload["label"] = f"Atención {attention['sequence_no']}/{len(attention_rows)}"
        attention_payload["lines"] = [
            _as_dict(row)
            for row in connection.execute(
                """SELECT al.*, l.np_code, l.description
                     FROM reception_attention_lines al
                     JOIN reception_lines l ON l.id = al.reception_line_id
                    WHERE al.attention_id = ? ORDER BY l.id""",
                (attention["id"],),
            ).fetchall()
        ]
        payload["attentions"].append(attention_payload)
    if active_attention:
        active_lines = {
            int(row["reception_line_id"]): row
            for row in connection.execute(
                "SELECT * FROM reception_attention_lines WHERE attention_id = ?",
                (active_attention["id"],),
            ).fetchall()
        }
        for line in payload["lines"]:
            attention_line = active_lines.get(int(line["id"]))
            if attention_line:
                line["attention_id"] = int(active_attention["id"])
                line["attention_sequence"] = int(active_attention["sequence_no"])
                line["attention_system_initialized"] = int(
                    active_attention["system_quantities_initialized"] or 0
                )
                line["attention_planned_qty"] = float(attention_line["planned_qty"] or 0)
                line["attention_verified_qty"] = float(attention_line["verified_qty"] or 0)
                line["attention_remaining_qty"] = max(
                    0.0,
                    line["attention_planned_qty"] - line["attention_verified_qty"],
                )
            else:
                line["attention_system_initialized"] = 1
                line["attention_planned_qty"] = float(line["expected_qty"] or 0)
                line["attention_verified_qty"] = float(line["received_qty"] or 0)
    if role == "ADMINISTRADOR" or RECEPTION_STATE_INDEX.get(payload["app_status"], -1) >= RECEPTION_STATE_INDEX["REVISION SISTEMA"]:
        accounting_rows = connection.execute(
            "SELECT * FROM reception_accounting_refs WHERE shipment_id = ? ORDER BY reception_date, source_row, id",
            (shipment_id,),
        ).fetchall()
        accounting_by_ip = _accounting_rows_by_ip(accounting_rows)
        for line in payload["lines"]:
            line_ips = set(_ip_tokens(line.get("ip_reference")))
            matching_refs = [
                row for row in accounting_rows
                if line_ips.intersection(_ip_tokens(row["ip_reference"]))
            ]
            line["fr_number"] = _joined_accounting_values(matching_refs, "fr_number") if matching_refs else ""
            line["em_number"] = _joined_accounting_values(matching_refs, "em_number") if matching_refs else ""
            line["accounting_locked"] = int(_line_has_complete_accounting(line, accounting_by_ip))
            line["accounting_lock_reason"] = (
                "FR y EM registrados para el IP; línea no requiere atención"
                if line["accounting_locked"] else ""
            )
    payload["receipts"] = [
        _as_dict(row)
        for row in connection.execute(
            """SELECT r.*, a.sequence_no AS attention_sequence
                 FROM reception_receipts r
                 LEFT JOIN reception_attentions a ON a.receipt_id = r.id
                WHERE r.shipment_id = ? ORDER BY r.sequence_no""", (shipment_id,)
        ).fetchall()
    ]
    payload["validation_sample_ids"] = _validation_sample_ids(connection, shipment_id)
    payload["lots"] = lots_for_shipment(connection, shipment_id)
    if role == "ADMINISTRADOR" or RECEPTION_STATE_INDEX.get(payload["app_status"], -1) >= RECEPTION_STATE_INDEX["REVISION SISTEMA"]:
        payload["accounting_refs"] = [
            _as_dict(row)
            for row in connection.execute(
                """SELECT * FROM reception_accounting_refs
                   WHERE shipment_id = ? ORDER BY reception_date, source_row, id""",
                (shipment_id,),
            ).fetchall()
        ]
    return payload


def _hours_since(value):
    if not value:
        return None
    try:
        started = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if started.tzinfo:
            started = started.replace(tzinfo=None)
        return max(0.0, (datetime.now() - started).total_seconds() / 3600)
    except (TypeError, ValueError):
        return None


def reception_report(connection, role="ADMINISTRADOR"):
    require_reception_access(role)
    if role != "ADMINISTRADOR":
        raise PermissionError("La reportería de Recepción es solo para el administrador")
    rows = connection.execute(
        """SELECT s.*,
                  (SELECT COUNT(*) FROM reception_lines l WHERE l.shipment_id = s.id) AS line_count,
                  (SELECT COUNT(*) FROM reception_receipts r WHERE r.shipment_id = s.id) AS receipt_count
           FROM reception_shipments s
           WHERE s.app_status <> 'CERRADO'
           ORDER BY CASE WHEN COALESCE(s.scheduled_date, '') = '' THEN 1 ELSE 0 END,
                    s.scheduled_date ASC, s.updated_at DESC"""
    ).fetchall()
    items = []
    by_status = {}
    by_transport = {}
    overdue = 0
    due_soon = 0
    for row in rows:
        item = _as_dict(row)
        transport = str(item.get("transport_type") or "SIN DEFINIR")
        sla_hours = 96 if transport == "MARITIMO" else 72
        elapsed = _hours_since(item.get("first_arrival_at"))
        if elapsed is None:
            sla_status = "SIN ARRIBO"
            remaining = None
        else:
            remaining = round(sla_hours - elapsed, 1)
            if remaining < 0:
                sla_status = "VENCIDO"
                overdue += 1
            elif remaining <= 12:
                sla_status = "VENCE PRONTO"
                due_soon += 1
            else:
                sla_status = "DENTRO SLA"
        item["sla_hours"] = sla_hours
        item["elapsed_hours"] = round(elapsed, 1) if elapsed is not None else None
        item["remaining_hours"] = remaining
        item["sla_status"] = sla_status
        item["work_bucket"] = (
            "CONTROL ADMINISTRATIVO"
            if item.get("accounting_status") in {"PENDIENTE CONTABILIDAD", "PENDIENTE FR", "CONFLICTO IP"}
            else "TRABAJO RECEPCIÓN"
        )
        items.append(item)
        by_status[item["app_status"]] = by_status.get(item["app_status"], 0) + 1
        by_transport[transport] = by_transport.get(transport, 0) + 1
    return {
        "generated_at": reception_now(),
        "summary": {
            "total_pendientes": len(items),
            "sin_arribo": sum(item["sla_status"] == "SIN ARRIBO" for item in items),
            "arribados_abiertos": sum(item["sla_status"] != "SIN ARRIBO" for item in items),
            "vencen_pronto": due_soon,
            "vencidos": overdue,
        },
        "by_status": by_status,
        "by_transport": by_transport,
        "items": items,
    }


def reception_operational_report_bytes(connection, shipment_id, role="ADMINISTRADOR", username=""):
    """Export the system review and final controls for one BL/AWB."""
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    lines = connection.execute(
        "SELECT * FROM reception_lines WHERE shipment_id = ? ORDER BY id", (shipment_id,)
    ).fetchall()
    refs = connection.execute(
        "SELECT * FROM reception_accounting_refs WHERE shipment_id = ? ORDER BY id", (shipment_id,)
    ).fetchall()
    refs_by_ip = {}
    for ref in refs:
        for ip in _ip_tokens(ref["ip_reference"]):
            refs_by_ip.setdefault(ip, []).append(ref)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["TRITON WMS · Reporte operativo de recepción"])
    summary.append(["BL / AWB", shipment["bl_awb"]])
    summary.append(["Estado", shipment["app_status"]])
    summary.append(["Bultos", f"{shipment['received_packages']} / {shipment['expected_packages']}"])
    summary.append(["Ubicación", shipment["location_text"] or ""])
    summary.append(["Generado", reception_now()])
    summary.append([])
    summary.append(["El reporte consolida el control de sistema, diferencias, observaciones y referencias contables."])
    lines_sheet = workbook.create_sheet("Packing List SAP")
    lines_sheet.append([
        "IP", "NP", "Descripción", "OC", "OV", "Esperado", "Verificado",
        "Diferencia", "Resultado", "Observación", "Comentario validación", "FR", "EM",
    ])
    for line in lines:
        expected = float(line["expected_qty"] or 0)
        verified = float(line["received_qty"] or 0)
        line_refs = []
        for ip in _ip_tokens(line["ip_reference"]):
            line_refs.extend(refs_by_ip.get(ip, []))
        unique_refs = {int(ref["id"]): ref for ref in line_refs}
        fr_values = ", ".join(sorted({str(ref["fr_number"] or "") for ref in unique_refs.values() if ref["fr_number"]}))
        em_values = ", ".join(sorted({str(ref["em_number"] or "") for ref in unique_refs.values() if ref["em_number"]}))
        difference = verified - expected
        result_label = "EXCEDENTE" if difference > 0 else "FALTANTE" if difference < 0 else "CONFORME"
        lines_sheet.append([
            line["ip_reference"] or "", line["np_code"] or "", line["description"] or "",
            line["oc_number"] or "", line["ov_number"] or "", expected, verified,
            difference, result_label, line["blocked_reason"] or "",
            line["validation_comment"] or "", fr_values, em_values,
        ])
    refs_sheet = workbook.create_sheet("FR EM")
    refs_sheet.append(["IP", "AWB", "Factura reserva", "EM", "Fecha recepción", "Fecha EM", "Estado"])
    for ref in refs:
        refs_sheet.append([
            ref["ip_reference"] or "", ref["awb_reference"] or "", ref["fr_number"] or "",
            ref["em_number"] or "", ref["reception_date"] or "", ref["em_date"] or "",
            ref["import_status"] or "",
        ])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="3E4A52")
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        for column in sheet.columns:
            values = list(column)
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in values[:100]) + 2))
            sheet.column_dimensions[values[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _history_datetime(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.replace(tzinfo=None) if parsed.tzinfo else parsed
    except (TypeError, ValueError):
        return None


def reception_history_report(connection, role="ADMINISTRADOR", preview_limit=200):
    """Builds operational timing data only when the administrator requests it."""
    require_reception_access(role)
    if role != "ADMINISTRADOR":
        raise PermissionError("El historial de Recepción es solo para el administrador")
    shipments = connection.execute(
        "SELECT * FROM reception_shipments ORDER BY created_at DESC, id DESC"
    ).fetchall()
    history_rows = connection.execute(
        """SELECT h.*, s.bl_awb, s.transport_type, s.current_assistant,
                         s.current_auxiliary, s.app_status AS current_app_status,
                         s.first_arrival_at
             FROM reception_history h
             JOIN reception_shipments s ON s.id = h.shipment_id
            ORDER BY h.shipment_id, h.id"""
    ).fetchall()
    by_shipment = {}
    for row in history_rows:
        by_shipment.setdefault(row["shipment_id"], []).append(row)

    timing_rows = []
    event_rows = []
    stage_names = list(RECEPTION_STATES)
    total_work_hours = 0.0
    by_user = {}
    by_event = {}
    now_value = reception_now()
    now_dt = _history_datetime(now_value) or datetime.now()

    for shipment in shipments:
        events = by_shipment.get(shipment["id"], [])
        stage_started = _history_datetime(shipment["created_at"])
        if stage_started is None:
            stage_started = now_dt
        current_stage = "PROGRAMADO"
        durations = {stage: 0.0 for stage in stage_names}
        previous_event_dt = None
        arrival_dt = _history_datetime(shipment["first_arrival_at"])

        for event in events:
            event_dt = _history_datetime(event["created_at"])
            if event_dt is None:
                continue
            stage_before = current_stage
            stage_closed_hours = None
            if event["field_name"] == "app_status" and event["new_value"]:
                stage_closed_hours = max(0.0, (event_dt - stage_started).total_seconds() / 3600)
                durations[current_stage] = durations.get(current_stage, 0.0) + stage_closed_hours
                current_stage = str(event["new_value"])
                stage_started = event_dt
            since_previous_minutes = (
                max(0.0, (event_dt - previous_event_dt).total_seconds() / 60)
                if previous_event_dt else None
            )
            since_arrival_hours = (
                max(0.0, (event_dt - arrival_dt).total_seconds() / 3600)
                if arrival_dt else None
            )
            event_rows.append({
                "shipment_id": shipment["id"],
                "bl_awb": event["bl_awb"],
                "transport_type": event["transport_type"],
                "event_type": event["event_type"],
                "field_name": event["field_name"],
                "old_value": event["old_value"] or "",
                "new_value": event["new_value"] or "",
                "stage_before": stage_before,
                "stage_after": current_stage,
                "stage_duration_hours": round(stage_closed_hours, 2) if stage_closed_hours is not None else "",
                "minutes_since_previous_event": round(since_previous_minutes, 2) if since_previous_minutes is not None else "",
                "hours_since_first_arrival": round(since_arrival_hours, 2) if since_arrival_hours is not None else "",
                "username": event["username"],
                "reason": event["reason"] or "",
                "created_at": event["created_at"],
            })
            previous_event_dt = event_dt
            by_user[event["username"] or "Sin usuario"] = by_user.get(event["username"] or "Sin usuario", 0) + 1
            by_event[event["event_type"] or "SIN EVENTO"] = by_event.get(event["event_type"] or "SIN EVENTO", 0) + 1

        current_stage = shipment["app_status"] or current_stage
        current_duration = max(0.0, (now_dt - stage_started).total_seconds() / 3600)
        durations[current_stage] = durations.get(current_stage, 0.0) + current_duration
        total_elapsed = sum(durations.values())
        total_work_hours += total_elapsed
        timing = {
            "shipment_id": shipment["id"],
            "bl_awb": shipment["bl_awb"],
            "transport_type": shipment["transport_type"],
            "current_status": shipment["app_status"],
            "assistant": shipment["current_assistant"] or "",
            "auxiliary": shipment["current_auxiliary"] or "",
            "created_at": shipment["created_at"],
            "first_arrival_at": shipment["first_arrival_at"] or "",
            "total_elapsed_hours": round(total_elapsed, 2),
            "events_count": len(events),
        }
        for stage in stage_names:
            timing[f"hours_{stage.lower().replace(' ', '_')}"] = round(durations.get(stage, 0.0), 2)
        timing_rows.append(timing)

    event_rows.sort(key=lambda row: (row["created_at"], row["shipment_id"]), reverse=True)
    timing_rows.sort(key=lambda row: (row["created_at"], row["shipment_id"]), reverse=True)
    if preview_limit is None:
        preview_events = event_rows
    else:
        preview_events = event_rows[:max(1, min(int(preview_limit), 500))]
    return {
        "generated_at": now_value,
        "summary": {
            "shipments": len(timing_rows),
            "events": len(event_rows),
            "total_elapsed_hours": round(total_work_hours, 2),
            "users": len(by_user),
        },
        "by_user": dict(sorted(by_user.items(), key=lambda item: (-item[1], item[0]))),
        "by_event": dict(sorted(by_event.items(), key=lambda item: (-item[1], item[0]))),
        "timings": timing_rows,
        "events": preview_events,
    }


def reception_history_export_bytes(connection, role="ADMINISTRADOR"):
    """Exports the complete audit and timing history in a readable workbook."""
    report = reception_history_report(connection, role, preview_limit=500)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    timings_sheet = workbook.create_sheet("Tiempos por BL")
    events_sheet = workbook.create_sheet("Eventos auditoria")
    sheets = [summary_sheet, timings_sheet, events_sheet]
    header_fill = PatternFill("solid", fgColor="3E4A52")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="343A40")

    summary_sheet.append(["TRITON WMS · Historial de Recepción"])
    summary_sheet.append(["Generado", report["generated_at"]])
    summary_sheet.append(["Expedientes BL/AWB", report["summary"]["shipments"]])
    summary_sheet.append(["Eventos de auditoría", report["summary"]["events"]])
    summary_sheet.append(["Horas acumuladas de seguimiento", report["summary"]["total_elapsed_hours"]])
    summary_sheet.append([])
    summary_sheet.append(["Eventos por usuario"])
    summary_sheet.append(["Usuario", "Eventos"])
    for username, count in report["by_user"].items():
        summary_sheet.append([username, count])
    summary_sheet.append([])
    summary_sheet.append(["Eventos por tipo"])
    summary_sheet.append(["Tipo", "Eventos"])
    for event_type, count in report["by_event"].items():
        summary_sheet.append([event_type, count])

    stage_headers = [
        "shipment_id", "bl_awb", "transport_type", "current_status", "assistant", "auxiliary",
        "created_at", "first_arrival_at", "total_elapsed_hours", "events_count",
    ] + [f"hours_{stage.lower().replace(' ', '_')}" for stage in RECEPTION_STATES]
    timings_sheet.append(stage_headers)
    for row in report["timings"]:
        timings_sheet.append([row.get(header, "") for header in stage_headers])

    event_headers = [
        "shipment_id", "bl_awb", "transport_type", "event_type", "field_name", "old_value",
        "new_value", "stage_before", "stage_after", "stage_duration_hours",
        "minutes_since_previous_event", "hours_since_first_arrival", "username", "reason", "created_at",
    ]
    events_sheet.append(event_headers)
    # The UI preview is capped, but the export always reads the complete event set.
    full_report = reception_history_report(connection, role, preview_limit=None)
    for row in full_report["events"]:
        events_sheet.append([row.get(header, "") for header in event_headers])

    for sheet in sheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = min(42, max(len(str(cell.value or "")) for cell in column_cells) + 2)
            sheet.column_dimensions[column_letter].width = max(12, max_length)
    summary_sheet.freeze_panes = "A8"
    for sheet in (timings_sheet, events_sheet):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)) and "hours" in str(sheet.cell(1, cell.column).value):
                    cell.number_format = "0.00"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_reception(connection, data, username, role):
    require_reception_access(role)
    if role != "ADMINISTRADOR":
        raise PermissionError("Solo el administrador puede crear o importar una BL")
    bl_awb = str(data.get("bl_awb") or "").strip()
    if not bl_awb:
        raise ValueError("La BL/AWB es obligatoria")
    transport_type = str(data.get("transport_type") or "AEREO").strip().upper()
    if transport_type not in {"AEREO", "MARITIMO", "COURIER"}:
        raise ValueError("El transporte debe ser AEREO, MARITIMO o COURIER")
    expected_packages = 0
    current_assistant = str(data.get("current_assistant") or "").strip()
    current_auxiliary = str(data.get("current_auxiliary") or "").strip()
    timestamp = reception_now()
    cursor = connection.execute(
        """INSERT INTO reception_shipments
           (bl_awb, transport_type, supplier, ip_reference, primary_oc, primary_ov,
            expected_packages, current_assistant, current_auxiliary, scheduled_date,
            created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            bl_awb,
            transport_type,
            str(data.get("supplier") or "").strip(),
            str(data.get("ip_reference") or "").strip(),
            str(data.get("primary_oc") or "").strip(),
            str(data.get("primary_ov") or "").strip(),
            expected_packages,
            current_assistant,
            current_auxiliary,
            str(data.get("scheduled_date") or "").strip(),
            timestamp,
            timestamp,
        ),
    )
    shipment_id = cursor.lastrowid
    _write_history(connection, shipment_id, "CREACION", "bl_awb", "", bl_awb, username)
    _ensure_default_reception_assignees(
        connection,
        shipment_id,
        username,
        "Creación de expediente: responsables predeterminados",
    )
    return reception_detail(connection, shipment_id, role)


def _normalize_header(value):
    value = unicodedata.normalize("NFKD", str(value or "").strip().upper())
    value = "".join(character for character in value if not unicodedata.combining(character))
    return " ".join(value.replace("\n", " ").split())


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    normalized = str(value).strip().replace(" ", "").replace(",", "")
    try:
        return float(normalized)
    except ValueError:
        return 0.0


def _find_column(headers, *aliases):
    for alias in aliases:
        normalized = _normalize_header(alias)
        if normalized in headers:
            return headers.index(normalized)
    return None


def _row_value(row, index):
    return row[index] if index is not None and index < len(row) else None


def _transport_type(value):
    normalized = _normalize_header(value)
    if "MARIT" in normalized:
        return "MARITIMO"
    if "AERE" in normalized:
        return "AEREO"
    if "COURIER" in normalized:
        return "COURIER"
    return normalized or "SIN DEFINIR"


def _unique_summary(records, field, limit=12):
    values = []
    seen = set()
    for record in records:
        value = str(record.get(field) or "").strip()
        if not value or value.upper() in {"S/N", "SIN NUMERO", "NONE"} or value in seen:
            continue
        seen.add(value)
        values.append(value)
    if len(values) > limit:
        return ", ".join(values[:limit]) + f" (+{len(values) - limit})"
    return ", ".join(values)


def _read_reception_workbook(workbook):
    grouped = {}
    valid_sheets = []
    skipped_rows = 0
    for worksheet in workbook.worksheets:
        header_row = None
        indexes = None
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 20), values_only=True), 1
        ):
            headers = [_normalize_header(value) for value in row]
            guide_index = _find_column(headers, "GUIA DE IMPORTACION")
            code_index = _find_column(headers, "CODIGO")
            if guide_index is None or code_index is None:
                continue
            header_row = row_number
            indexes = {
                "guide": guide_index,
                "code": code_index,
                "description": _find_column(headers, "DESCRIPCION"),
                "ov": _find_column(headers, "OV", "ORDEN DE VENTA"),
                "oc": _find_column(headers, "ORDEN DE COMPRA"),
                "requested": _find_column(headers, "CANTIDAD SOLICITADA"),
                "invoiced": _find_column(headers, "CANTIDAD FACTURADA"),
                "pending": _find_column(headers, "CANTIDAD PENDIENTE"),
                "purchase_type": _find_column(headers, "TIPO DE COMPRA"),
                "brand": _find_column(headers, "MARCA"),
                "applicant": _find_column(headers, "SOLICITANTE"),
                "ip": _find_column(headers, "PEDIDO TRITON"),
                "status": _find_column(headers, "STATUS"),
                "status_by_np": _find_column(headers, "ESTADO POR NP"),
                "invoice": _find_column(headers, "N° FACTURA", "Nº FACTURA", "N FACTURA"),
                "transport": _find_column(headers, "MODO DE TRANSPORTE"),
                "country": _find_column(headers, "PAIS DE ORIGEN"),
                "confirmed_date": _find_column(headers, "FECHA CONFIRMADA TRITON"),
                "estimated_date": _find_column(headers, "FECHA TRITON ESTIMADA"),
            }
            break
        if header_row is None:
            continue
        valid_sheets.append(worksheet.title)
        duplicate_counter = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            guide = _cell_text(_row_value(row, indexes["guide"]))
            np_code = _cell_text(_row_value(row, indexes["code"]))
            if not guide or not np_code:
                skipped_rows += 1
                continue
            requested = _number(_row_value(row, indexes["requested"]))
            invoiced = _number(_row_value(row, indexes["invoiced"]))
            pending = _number(_row_value(row, indexes["pending"]))
            record = {
                "guide": guide,
                "source_sheet": worksheet.title,
                "source_row": row_number,
                "np_code": np_code,
                "description": _cell_text(_row_value(row, indexes["description"])),
                "ov": _cell_text(_row_value(row, indexes["ov"])),
                "oc": _cell_text(_row_value(row, indexes["oc"])),
                "requested_qty": requested,
                "invoiced_qty": invoiced,
                "pending_qty": pending,
                "expected_qty": invoiced if invoiced > 0 else requested,
                "purchase_type": _cell_text(_row_value(row, indexes["purchase_type"])),
                "brand": _cell_text(_row_value(row, indexes["brand"])),
                "applicant": _cell_text(_row_value(row, indexes["applicant"])),
                "ip": _cell_text(_row_value(row, indexes["ip"])),
                "source_status": _cell_text(_row_value(row, indexes["status"])),
                "status_by_np": _cell_text(_row_value(row, indexes["status_by_np"])),
                "source_invoice": _cell_text(_row_value(row, indexes["invoice"])),
                "transport_type": _transport_type(_row_value(row, indexes["transport"])),
                "country": _cell_text(_row_value(row, indexes["country"])),
                "confirmed_date": _cell_text(_row_value(row, indexes["confirmed_date"])),
                "estimated_date": _cell_text(_row_value(row, indexes["estimated_date"])),
            }
            record["scheduled_date"] = record["confirmed_date"] or record["estimated_date"]
            identity = "|".join(
                (
                    worksheet.title,
                    guide,
                    record["ip"],
                    record["oc"],
                    record["ov"],
                    np_code,
                )
            )
            duplicate_counter[identity] = duplicate_counter.get(identity, 0) + 1
            record["source_key"] = hashlib.sha256(
                f"{identity}|{duplicate_counter[identity]}".encode("utf-8")
            ).hexdigest()
            grouped.setdefault(guide, []).append(record)
    if not valid_sheets:
        raise ValueError(
            "No encontré hojas con las columnas GUIA DE IMPORTACION y CODIGO en las primeras 20 filas"
        )
    return grouped, valid_sheets, skipped_rows


def _source_confirms_arrival(records):
    """Return whether the import source confirms physical arrival.

    A confirmed date is deliberately required; an estimated date must not move
    a BL from PROGRAMADO. The source used by Importaciones marks each line as
    DISPONIBLE and COMPLETE when the received import is ready for operation.
    A BL is considered arrived only when every imported line has no pending
    quantity and carries one of those source completion states.
    """
    if not records or not all(record.get("confirmed_date") for record in records):
        return False
    valid_statuses = {"DISPONIBLE", "COMPLETO", "COMPLETADO"}
    for record in records:
        source_status = _normalized_identifier(record.get("source_status"))
        line_status = _normalized_identifier(record.get("status_by_np"))
        if float(record.get("pending_qty") or 0) > 0:
            return False
        if source_status not in valid_statuses and line_status not in valid_statuses:
            return False
    return True


def import_reception_workbook(connection, workbook, filename, username, role):
    require_reception_access(role)
    if role != "ADMINISTRADOR":
        raise PermissionError("Solo el administrador puede cargar el Excel de Recepción")
    grouped, valid_sheets, skipped_rows = _read_reception_workbook(workbook)
    timestamp = reception_now()
    shipments_created = 0
    shipments_updated = 0
    lines_created = 0
    lines_updated = 0
    for guide, records in grouped.items():
        transport = _unique_summary(records, "transport_type", 2) or "SIN DEFINIR"
        brands = _unique_summary(records, "brand", 5)
        country = _unique_summary(records, "country", 3)
        source_sheets = _unique_summary(records, "source_sheet", 12)
        ip_reference = _unique_summary(records, "ip", 12)
        primary_oc = _unique_summary(records, "oc", 12)
        primary_ov = _unique_summary(records, "ov", 12)
        scheduled_date = next((record["scheduled_date"] for record in records if record["scheduled_date"]), "")
        shipment = connection.execute(
            "SELECT * FROM reception_shipments WHERE bl_awb = ?", (guide,)
        ).fetchone()
        if shipment:
            shipment_id = shipment["id"]
            connection.execute(
                """UPDATE reception_shipments
                   SET transport_type = ?, brand_summary = ?, country_origin = ?, source_sheets = ?,
                       ip_reference = ?, primary_oc = ?, primary_ov = ?,
                       scheduled_date = CASE WHEN ? <> '' THEN ? ELSE scheduled_date END,
                       updated_at = ?
                   WHERE id = ?""",
                (
                    transport,
                    brands,
                    country,
                    source_sheets,
                    ip_reference,
                    primary_oc,
                    primary_ov,
                    scheduled_date,
                    scheduled_date,
                    timestamp,
                    shipment_id,
                ),
            )
            shipments_updated += 1
        else:
            cursor = connection.execute(
                """INSERT INTO reception_shipments
                   (bl_awb, transport_type, brand_summary, country_origin, source_sheets,
                    ip_reference, primary_oc, primary_ov, scheduled_date,
                    created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    guide,
                    transport,
                    brands,
                    country,
                    source_sheets,
                    ip_reference,
                    primary_oc,
                    primary_ov,
                    scheduled_date,
                    timestamp,
                    timestamp,
                ),
            )
            shipment_id = cursor.lastrowid
            shipments_created += 1
        _ensure_default_reception_assignees(
            connection,
            shipment_id,
            username,
            f"{filename}: responsable predeterminado para evitar pendientes de asignación",
        )
        if _source_confirms_arrival(records):
            current = connection.execute(
                "SELECT app_status, condition_status, first_arrival_at FROM reception_shipments WHERE id = ?",
                (shipment_id,),
            ).fetchone()
            current_status = str(current["app_status"] or "PROGRAMADO")
            if RECEPTION_STATE_INDEX.get(current_status, 0) < RECEPTION_STATE_INDEX["ARRIBADO"]:
                arrival_date = next(
                    (record["confirmed_date"] for record in records if record.get("confirmed_date")),
                    "",
                )
                connection.execute(
                    """UPDATE reception_shipments
                       SET app_status = 'ARRIBADO', condition_status = 'ARRIBO REGISTRADO',
                           first_arrival_at = COALESCE(first_arrival_at, ?), updated_at = ?
                       WHERE id = ?""",
                    (arrival_date, timestamp, shipment_id),
                )
                _write_history(
                    connection,
                    shipment_id,
                    "IMPORTACION EXCEL",
                    "app_status",
                    current_status,
                    "ARRIBADO",
                    username,
                    "STATUS=DISPONIBLE, fecha confirmada y todas las líneas completas sin pendiente",
                )
        shipment_lines_created = 0
        shipment_lines_updated = 0
        for record in records:
            existing_line = connection.execute(
                "SELECT id FROM reception_lines WHERE source_key = ?", (record["source_key"],)
            ).fetchone()
            values = (
                shipment_id,
                record["source_row"],
                record["source_sheet"],
                record["source_key"],
                record["np_code"],
                record["description"],
                record["oc"],
                record["ov"],
                record["expected_qty"],
                record["requested_qty"],
                record["invoiced_qty"],
                record["pending_qty"],
                record["purchase_type"],
                record["brand"],
                record["applicant"],
                record["source_status"],
                record["source_invoice"],
                record["ip"],
            )
            if existing_line:
                connection.execute(
                    """UPDATE reception_lines
                       SET shipment_id = ?, source_row = ?, source_sheet = ?, source_key = ?,
                           np_code = ?, description = ?, oc_number = ?, ov_number = ?,
                           expected_qty = ?, requested_qty = ?, invoiced_qty = ?, pending_qty = ?,
                           purchase_type = ?, brand = ?, applicant = ?, source_status = ?,
                           source_invoice = ?, ip_reference = ?
                       WHERE id = ?""",
                    values + (existing_line["id"],),
                )
                lines_updated += 1
                shipment_lines_updated += 1
            else:
                connection.execute(
                    """INSERT INTO reception_lines
                       (shipment_id, source_row, source_sheet, source_key, np_code, description,
                        oc_number, ov_number, expected_qty, requested_qty, invoiced_qty, pending_qty,
                        purchase_type, brand, applicant, source_status, source_invoice, ip_reference)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    values,
                )
                lines_created += 1
                shipment_lines_created += 1
        _write_history(
            connection,
            shipment_id,
            "IMPORTACION EXCEL",
            "lineas",
            shipment_lines_updated,
            shipment_lines_created + shipment_lines_updated,
            username,
            f"{filename}: {shipment_lines_created} nuevas y {shipment_lines_updated} actualizadas",
        )
        # Si ya existía una carga contable, vuelve a evaluar la BL después de
        # refrescar sus líneas. Así una BL completa puede cerrarse y sus
        # cantidades quedar alineadas sin esperar al siguiente reinicio.
        _refresh_accounting_summary(connection, shipment_id, username, filename)
    return {
        "filename": filename,
        "sheets": valid_sheets,
        "shipments": len(grouped),
        "shipments_created": shipments_created,
        "shipments_updated": shipments_updated,
        "lines_created": lines_created,
        "lines_updated": lines_updated,
        "skipped_rows": skipped_rows,
    }


def import_reception_excel_bytes(connection, content, filename, username, role):
    if not str(filename or "").lower().endswith(".xlsx"):
        raise ValueError("El archivo de Recepción debe ser .xlsx")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Conditional Formatting extension is not supported")
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        try:
            return import_reception_workbook(connection, workbook, filename, username, role)
        finally:
            workbook.close()


def import_reception_excel_path(connection, path, username="sistema.local", role="ADMINISTRADOR"):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Conditional Formatting extension is not supported")
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            return import_reception_workbook(connection, workbook, str(path), username, role)
        finally:
            workbook.close()


def _normalized_identifier(value):
    normalized = unicodedata.normalize("NFKD", str(value or "").strip().upper())
    normalized = "".join(character for character in normalized if not unicodedata.combining(character))
    return "".join(character for character in normalized if character.isalnum())


def _accounting_header_indexes(headers):
    required = {
        "fr_number": _find_column(headers, "FACTURA RESERVA"),
        "em_number": _find_column(headers, "ENTRADA DE MERCANCIA"),
        "reception_date": _find_column(headers, "FECHA DE RECEPCION"),
        "em_date": _find_column(headers, "FECHA DE EM"),
        "ip_reference": _find_column(headers, "IP"),
        "awb_reference": _find_column(headers, "AWB"),
    }
    if any(index is None for index in required.values()):
        return None
    return {
        **required,
        "email_date": _find_column(headers, "FECHA DE CORREO"),
        "costing_date": _find_column(headers, "FECHA DE COSTEO"),
        "import_status": _find_column(headers, "ESTADO IMPORTACION", "ESTADO"),
    }


def _read_accounting_workbook(workbook):
    canonical = [
        worksheet for worksheet in workbook.worksheets
        if _normalize_header(worksheet.title) == "FACTURAS DHL"
    ]
    if not canonical:
        raise ValueError("No encontré la hoja FACTURAS DHL en el archivo contable")

    records = []
    blocks = []
    skipped_rows = 0
    duplicate_counter = {}
    for worksheet in canonical:
        indexes = None
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
            headers = [_normalize_header(value) for value in row]
            detected = _accounting_header_indexes(headers)
            if detected:
                indexes = detected
                blocks.append({"sheet": worksheet.title, "header_row": row_number})
                continue
            if indexes is None:
                continue

            awb = _cell_text(_row_value(row, indexes["awb_reference"]))
            ip_reference = _cell_text(_row_value(row, indexes["ip_reference"]))
            fr_number = _cell_text(_row_value(row, indexes["fr_number"]))
            em_number = _cell_text(_row_value(row, indexes["em_number"]))
            if not awb and not ip_reference and not fr_number and not em_number:
                continue
            if not (awb or ip_reference) or not (fr_number or em_number):
                skipped_rows += 1
                continue

            pair = (_normalized_identifier(awb), _normalized_identifier(ip_reference))
            duplicate_counter[pair] = duplicate_counter.get(pair, 0) + 1
            source_key = hashlib.sha256(
                f"{worksheet.title}|{pair[0]}|{pair[1]}|{duplicate_counter[pair]}".encode("utf-8")
            ).hexdigest()
            records.append(
                {
                    "source_key": source_key,
                    "source_sheet": worksheet.title,
                    "source_row": row_number,
                    "awb_reference": awb,
                    "ip_reference": ip_reference,
                    "fr_number": fr_number,
                    "em_number": em_number,
                    "email_date": _cell_text(_row_value(row, indexes["email_date"])),
                    "reception_date": _cell_text(_row_value(row, indexes["reception_date"])),
                    "em_date": _cell_text(_row_value(row, indexes["em_date"])),
                    "costing_date": _cell_text(_row_value(row, indexes["costing_date"])),
                    "import_status": _cell_text(_row_value(row, indexes["import_status"])),
                }
            )
    if not blocks:
        raise ValueError(
            "No encontré un bloque con AWB, IP, FACTURA RESERVA, ENTRADA DE MERCANCIA, "
            "FECHA DE RECEPCION y FECHA DE EM"
        )
    if not records:
        raise ValueError("El archivo contable no contiene filas FR/EM utilizables")
    return records, blocks, skipped_rows


def _identifier_candidates(value):
    raw = str(value or "").strip()
    candidates = []
    whole = _normalized_identifier(raw)
    if whole:
        candidates.append(whole)
    if "(" in raw and ")" in raw:
        outside = raw.split("(", 1)[0]
        inside = raw.split("(", 1)[1].split(")", 1)[0]
        for part in (outside, inside):
            normalized = _normalized_identifier(part)
            if normalized and normalized not in candidates:
                candidates.append(normalized)
    return candidates


def _shipment_identifier_maps(connection):
    awb_map = {}
    ip_map = {}
    shipments = connection.execute(
        "SELECT id, bl_awb, ip_reference FROM reception_shipments"
    ).fetchall()
    for shipment in shipments:
        awb = _normalized_identifier(shipment["bl_awb"])
        if awb:
            awb_map.setdefault(awb, set()).add(shipment["id"])
        for ip_reference in str(shipment["ip_reference"] or "").replace(";", ",").split(","):
            ip_key = _normalized_identifier(ip_reference)
            if ip_key:
                ip_map.setdefault(ip_key, set()).add(shipment["id"])
    for line in connection.execute(
        "SELECT shipment_id, ip_reference FROM reception_lines WHERE ip_reference IS NOT NULL"
    ).fetchall():
        ip_key = _normalized_identifier(line["ip_reference"])
        if ip_key:
            ip_map.setdefault(ip_key, set()).add(line["shipment_id"])
    return awb_map, ip_map


def _match_accounting_record(record, awb_map, ip_map):
    candidates = set()
    for awb_key in _identifier_candidates(record["awb_reference"]):
        candidates.update(awb_map.get(awb_key, set()))
    method = "AWB"
    if not candidates:
        raw_ip = str(record["ip_reference"] or "").replace(";", ",").replace(" - ", ",")
        for ip_reference in raw_ip.split(","):
            ip_key = _normalized_identifier(ip_reference)
            candidates.update(ip_map.get(ip_key, set()))
        method = "IP"
    if len(candidates) == 1:
        return next(iter(candidates)), method
    if len(candidates) > 1:
        return None, "AMBIGUO"
    return None, "SIN COINCIDENCIA"


def _joined_accounting_values(rows, field, limit=60):
    values = []
    seen = set()
    for row in rows:
        value = str(row[field] or "").strip()
        if not value or value in seen:
            continue
        seen.add(value)
        values.append(value)
    if len(values) > limit:
        return ", ".join(values[:limit]) + f" (+{len(values) - limit})"
    return ", ".join(values)


def _refresh_accounting_summary(connection, shipment_id, username, filename):
    shipment = connection.execute(
        "SELECT * FROM reception_shipments WHERE id = ?", (shipment_id,)
    ).fetchone()
    rows = connection.execute(
        """SELECT * FROM reception_accounting_refs
           WHERE shipment_id = ? ORDER BY reception_date, source_row, id""",
        (shipment_id,),
    ).fetchall()
    # Una línea con FR y EM completos ya fue atendida por el circuito contable.
    # Algunas fuentes de importación traen received_qty=0; para evitar que una
    # BL cerrada aparezca con cantidades pendientes, inicializamos únicamente
    # esas líneas en la cantidad esperada. No sobrescribimos diferencias
    # registradas manualmente durante la revisión de sistema.
    accounting_by_ip = _accounting_rows_by_ip(rows)
    for line in connection.execute(
        "SELECT * FROM reception_lines WHERE shipment_id = ? ORDER BY id",
        (shipment_id,),
    ).fetchall():
        expected_qty = float(line["expected_qty"] or 0)
        received_qty = float(line["received_qty"] or 0)
        if (
            expected_qty > 0
            and received_qty == 0
            and _line_has_complete_accounting(line, accounting_by_ip)
        ):
            connection.execute(
                "UPDATE reception_lines SET received_qty = ? WHERE id = ?",
                (expected_qty, line["id"]),
            )
            _write_history(
                connection,
                shipment_id,
                "IMPORTACION FR/EM",
                f"linea:{line['id']}:received_qty",
                received_qty,
                expected_qty,
                username,
                "FR y EM completas para la IP; se toma la cantidad esperada como atendida",
            )
    new_accounting_status = _accounting_status_for_shipment(
        connection, shipment_id, shipment, rows
    )
    old_accounting_status = str(shipment["accounting_status"] or "")
    state_changes = 0
    if new_accounting_status != old_accounting_status:
        connection.execute(
            "UPDATE reception_shipments SET accounting_status = ?, updated_at = ? WHERE id = ?",
            (new_accounting_status, reception_now(), shipment_id),
        )
        _write_history(
            connection,
            shipment_id,
            "IMPORTACION FR/EM",
            "accounting_status",
            old_accounting_status,
            new_accounting_status,
            username,
            filename,
        )
    old_app_status = str(shipment["app_status"] or "")
    if new_accounting_status == "EM REGISTRADA" and old_app_status in {"PROGRAMADO", "ARRIBADO"}:
        connection.execute(
            """UPDATE reception_shipments
               SET app_status = 'CERRADO', condition_status = 'COMPLETADO', updated_at = ?
               WHERE id = ?""",
            (reception_now(), shipment_id),
        )
        _write_history(
            connection,
            shipment_id,
            "IMPORTACION FR/EM",
            "app_status",
            old_app_status,
            "CERRADO",
            username,
            "FR y EM completas para las referencias contables de la BL",
        )
        state_changes += 1
    elif new_accounting_status != "EM REGISTRADA" and old_app_status == "CERRADO":
        reopened_status = "ARRIBADO" if float(shipment["received_packages"] or 0) > 0 else "PROGRAMADO"
        reopened_condition = "ARRIBO REGISTRADO" if reopened_status == "ARRIBADO" else "POR ARRIBAR"
        connection.execute(
            """UPDATE reception_shipments
               SET app_status = ?, condition_status = ?, updated_at = ?
               WHERE id = ?""",
            (reopened_status, reopened_condition, reception_now(), shipment_id),
        )
        _write_history(
            connection,
            shipment_id,
            "IMPORTACION FR/EM",
            "app_status",
            old_app_status,
            reopened_status,
            username,
            "La referencia contable dejó de estar completa; se reabre para control",
        )
        state_changes += 1
    mappings = {
        "fr_number": "fr_number",
        "em_number": "em_number",
        "accounting_email_date": "email_date",
        "source_reception_date": "reception_date",
        "em_date": "em_date",
        "costing_date": "costing_date",
    }
    changed_fields = 0
    for shipment_field, reference_field in mappings.items():
        new_value = _joined_accounting_values(rows, reference_field)
        old_value = str(shipment[shipment_field] or "")
        if new_value == old_value:
            continue
        connection.execute(
            f"UPDATE reception_shipments SET {shipment_field} = ?, updated_at = ? WHERE id = ?",
            (new_value, reception_now(), shipment_id),
        )
        _write_history(
            connection,
            shipment_id,
            "IMPORTACION FR/EM",
            shipment_field,
            old_value,
            new_value,
            username,
            filename,
        )
        changed_fields += 1
    sync_shipment_lot_references(connection, shipment_id)
    return changed_fields + int(new_accounting_status != old_accounting_status) + state_changes


def _refresh_manual_accounting_status(connection, shipment_id, username):
    shipment = connection.execute(
        "SELECT * FROM reception_shipments WHERE id = ?", (shipment_id,)
    ).fetchone()
    if not shipment:
        return
    new_status = _accounting_status_for_shipment(connection, shipment_id, shipment)
    old_status = str(shipment["accounting_status"] or "")
    if new_status != old_status:
        connection.execute(
            "UPDATE reception_shipments SET accounting_status = ?, updated_at = ? WHERE id = ?",
            (new_status, reception_now(), shipment_id),
        )
        _write_history(
            connection, shipment_id, "MODIFICACION", "accounting_status",
            old_status, new_status, username,
        )


def import_reception_accounting_workbook(connection, workbook, filename, username, role):
    require_reception_access(role)
    if role != "ADMINISTRADOR":
        raise PermissionError("Solo el administrador puede cargar el Excel de FR/EM")
    records, blocks, skipped_rows = _read_accounting_workbook(workbook)
    awb_map, ip_map = _shipment_identifier_maps(connection)
    timestamp = reception_now()
    created = 0
    updated = 0
    unchanged = 0
    matched_by_awb = 0
    matched_by_ip = 0
    ambiguous = []
    unmatched = []
    affected_shipments = set()
    matched_records = []
    reference_fields = (
        "source_sheet", "source_row", "awb_reference", "ip_reference", "fr_number",
        "em_number", "email_date", "reception_date", "em_date", "costing_date",
        "import_status", "source_file",
    )
    for record in records:
        shipment_id, method = _match_accounting_record(record, awb_map, ip_map)
        if shipment_id is None:
            item = {
                "row": record["source_row"],
                "awb": record["awb_reference"],
                "ip": record["ip_reference"],
            }
            (ambiguous if method == "AMBIGUO" else unmatched).append(item)
            continue
        affected_shipments.add(shipment_id)
        matched_records.append((record, shipment_id, method))

    ip_assignments = {}
    for record, shipment_id, _method in matched_records:
        for ip_key in _ip_tokens(record["ip_reference"]):
            ip_assignments.setdefault((shipment_id, ip_key), set()).add(
                (
                    _normalized_identifier(record["fr_number"]),
                    _normalized_identifier(record["em_number"]),
                )
            )
    conflicting_ip_keys = {
        key for key, assignments in ip_assignments.items() if len(assignments) > 1
    }
    conflicts = []
    conflict_by_shipment = {}
    valid_records = []
    for record, shipment_id, method in matched_records:
        record_ip_keys = [
            (shipment_id, ip_key) for ip_key in _ip_tokens(record["ip_reference"])
        ]
        conflict_keys = [key for key in record_ip_keys if key in conflicting_ip_keys]
        if conflict_keys:
            conflict_ips = [key[1] for key in conflict_keys]
            for _shipment, ip_key in conflict_keys:
                conflict_by_shipment.setdefault(shipment_id, set()).add(ip_key)
            conflicts.append(
                {
                    "row": record["source_row"],
                    "awb": record["awb_reference"],
                    "ip": record["ip_reference"],
                    "fr": record["fr_number"],
                    "em": record["em_number"],
                    "conflicting_ips": conflict_ips,
                }
            )
            continue
        valid_records.append((record, shipment_id, method))

    for record, shipment_id, method in valid_records:
        matched_by_awb += int(method == "AWB")
        matched_by_ip += int(method == "IP")
        values = {
            **{field: record.get(field, "") for field in reference_fields if field != "source_file"},
            "source_file": filename,
        }
        existing = connection.execute(
            "SELECT * FROM reception_accounting_refs WHERE source_key = ?",
            (record["source_key"],),
        ).fetchone()
        if existing:
            changed = existing["shipment_id"] != shipment_id or any(
                str(existing[field] or "") != str(values[field] or "") for field in reference_fields
            )
            if changed:
                assignments = ", ".join(f"{field} = ?" for field in reference_fields)
                connection.execute(
                    f"""UPDATE reception_accounting_refs
                        SET shipment_id = ?, {assignments}, updated_at = ? WHERE id = ?""",
                    (shipment_id, *(values[field] for field in reference_fields), timestamp, existing["id"]),
                )
                updated += 1
            else:
                unchanged += 1
        else:
            columns = ", ".join(reference_fields)
            placeholders = ", ".join("?" for _ in reference_fields)
            connection.execute(
                f"""INSERT INTO reception_accounting_refs
                    (shipment_id, source_key, {columns}, created_at, updated_at)
                    VALUES (?, ?, {placeholders}, ?, ?)""",
                (
                    shipment_id,
                    record["source_key"],
                    *(values[field] for field in reference_fields),
                    timestamp,
                    timestamp,
                ),
            )
            created += 1

    changed_summary_fields = 0
    conflict_shipments = set(affected_shipments)
    for shipment_id in conflict_shipments:
        conflict_ips = sorted(conflict_by_shipment.get(shipment_id, set()))
        new_count = len(conflict_ips)
        new_detail = ", ".join(conflict_ips[:20])
        if len(conflict_ips) > 20:
            new_detail += f" (+{len(conflict_ips) - 20})"
        current = connection.execute(
            "SELECT accounting_conflict_count, accounting_conflict_detail FROM reception_shipments WHERE id = ?",
            (shipment_id,),
        ).fetchone()
        old_count = int(current["accounting_conflict_count"] or 0)
        old_detail = str(current["accounting_conflict_detail"] or "")
        if old_count == new_count and old_detail == new_detail:
            continue
        connection.execute(
            """UPDATE reception_shipments
               SET accounting_conflict_count = ?, accounting_conflict_detail = ?, updated_at = ?
               WHERE id = ?""",
            (new_count, new_detail, reception_now(), shipment_id),
        )
        _write_history(
            connection,
            shipment_id,
            "IMPORTACION FR/EM",
            "accounting_conflict_detail",
            old_detail,
            new_detail,
            username,
            "IP con más de una combinación FR/EM; filas conflictivas no aplicadas",
        )
        changed_summary_fields += 1
    for shipment_id in affected_shipments:
        changed_summary_fields += _refresh_accounting_summary(
            connection, shipment_id, username, filename
        )
    return {
        "filename": filename,
        "sheet": "FACTURAS DHL",
        "blocks": blocks,
        "records_read": len(records),
        "matched_records": matched_by_awb + matched_by_ip,
        "matched_by_awb": matched_by_awb,
        "matched_by_ip": matched_by_ip,
        "shipments_updated": len(affected_shipments),
        "references_created": created,
        "references_updated": updated,
        "references_unchanged": unchanged,
        "summary_fields_changed": changed_summary_fields,
        "unmatched_count": len(unmatched),
        "unmatched_examples": unmatched[:20],
        "ambiguous_count": len(ambiguous),
        "ambiguous_examples": ambiguous[:20],
        "skipped_rows": skipped_rows,
        "conflict_count": len(conflicts),
        "conflict_ip_count": len(conflicting_ip_keys),
        "conflict_examples": conflicts[:20],
    }


def import_reception_accounting_excel_bytes(connection, content, filename, username, role):
    if not str(filename or "").lower().endswith(".xlsx"):
        raise ValueError("El archivo de FR/EM debe ser .xlsx")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Conditional Formatting extension is not supported")
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        try:
            return import_reception_accounting_workbook(
                connection, workbook, filename, username, role
            )
        finally:
            workbook.close()


def import_reception_accounting_excel_path(
    connection, path, username="sistema.local", role="ADMINISTRADOR"
):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Conditional Formatting extension is not supported")
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            return import_reception_accounting_workbook(
                connection, workbook, str(path), username, role
            )
        finally:
            workbook.close()


def add_physical_receipt(connection, shipment_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    if int(shipment["bultos_closed"] or 0):
        raise PermissionError("Los bultos ya fueron cerrados; si llegó algo adicional, el administrador debe reabrir el expediente")
    quantity = float(data.get("received_packages") or 0)
    if quantity <= 0:
        raise ValueError("La cantidad recibida debe ser mayor que cero")
    sequence_no = connection.execute(
        "SELECT COALESCE(MAX(sequence_no), 0) + 1 FROM reception_receipts WHERE shipment_id = ?",
        (shipment_id,),
    ).fetchone()[0]
    timestamp = reception_now()
    receipt_cursor = connection.execute(
        """INSERT INTO reception_receipts
           (shipment_id, sequence_no, received_packages, notes, location_text, username, received_at)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (
            shipment_id,
            sequence_no,
            quantity,
            str(data.get("notes") or "").strip(),
            str(data.get("location") or "").strip(),
            username,
            timestamp,
        ),
    )
    receipt_id = receipt_cursor.lastrowid
    # Una nueva atención se abre cuando la anterior ya terminó. Esto permite
    # que una BL tenga Atención 1/2 y luego Atención 2/2 sin mezclar cantidades.
    active_attention = _active_reception_attention(connection, shipment_id)
    create_attention = active_attention is None or active_attention["app_status"] == "CERRADO"
    if create_attention:
        _create_reception_attention(connection, shipment, receipt_id, username, "ARRIBADO")
    old_total = float(shipment["received_packages"] or 0)
    new_total = old_total + quantity
    expected = float(shipment["expected_packages"] or 0)
    if expected and new_total < expected:
        condition = "SALDO POR ARRIBAR"
    elif expected and new_total > expected:
        condition = "EXCEDENTE POR VALIDAR"
    elif expected:
        condition = "ARRIBO COMPLETO"
    else:
        condition = "ARRIBO REGISTRADO"
    resulting_status = (
        "ARRIBADO"
        if shipment["app_status"] == "CERRADO"
        else shipment["app_status"]
        if RECEPTION_STATE_INDEX.get(shipment["app_status"], 0) >= RECEPTION_STATE_INDEX["ARRIBADO"]
        else "ARRIBADO"
    )
    connection.execute(
        """UPDATE reception_shipments
               SET received_packages = ?, app_status = ?, condition_status = ?,
                   first_arrival_at = COALESCE(first_arrival_at, ?),
                   completed_arrival_at = CASE WHEN expected_packages > 0 AND ? >= expected_packages THEN ? ELSE completed_arrival_at END,
                   system_quantities_initialized = CASE WHEN ? = 'CERRADO' THEN 0 ELSE system_quantities_initialized END,
                   transfer_assistant_checked = CASE WHEN ? = 'CERRADO' THEN 0 ELSE transfer_assistant_checked END,
                   location_text = CASE WHEN ? = 'CERRADO' THEN '' ELSE location_text END,
                   updated_at = ?
               WHERE id = ?""",
        (
            new_total,
            resulting_status,
            condition,
            timestamp,
            new_total,
            timestamp,
            shipment["app_status"],
            shipment["app_status"],
            shipment["app_status"],
            timestamp,
            shipment_id,
        ),
    )
    _write_history(
        connection,
        shipment_id,
        "ARRIBO",
        "received_packages",
        old_total,
        new_total,
        username,
        f"Recepción física {sequence_no}",
    )
    return reception_detail(connection, shipment_id, role)


def close_physical_receipts(connection, shipment_id, data, username, role):
    """Close the physical-arrival window without changing expected packages."""
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    if shipment["app_status"] == "CERRADO":
        raise PermissionError("La recepción ya está cerrada")
    if int(shipment["bultos_closed"] or 0):
        return reception_detail(connection, shipment_id, role)
    reason = str(data.get("reason") or "No quedan bultos pendientes por recibir").strip()
    timestamp = reception_now()
    connection.execute(
        "UPDATE reception_shipments SET bultos_closed = 1, updated_at = ? WHERE id = ?",
        (timestamp, shipment_id),
    )
    _write_history(
        connection, shipment_id, "ARRIBO", "bultos_closed", 0, 1, username, reason
    )
    return reception_detail(connection, shipment_id, role)


def update_reception_location(connection, shipment_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    # El asistente/auxiliar solo puede registrar la ubicación en el paso 2.
    # El administrador puede corregir una ubicación histórica después de que
    # el expediente avanzó o se cerró, sin reabrir ni borrar la trazabilidad.
    is_admin = role == "ADMINISTRADOR"
    if not is_admin and shipment["app_status"] != "ARRIBADO":
        raise PermissionError("La ubicación de recepción se registra únicamente en ZONA RECEPCIÓN")
    location = str(data.get("location") or "").strip()
    if not location:
        raise ValueError("Indica la ubicación física en texto")
    receipt_id = data.get("receipt_id")
    if receipt_id:
        receipt = connection.execute(
            "SELECT * FROM reception_receipts WHERE id = ? AND shipment_id = ?",
            (int(receipt_id), shipment_id),
        ).fetchone()
    else:
        receipt = connection.execute(
            "SELECT * FROM reception_receipts WHERE shipment_id = ? ORDER BY sequence_no DESC LIMIT 1",
            (shipment_id,),
        ).fetchone()
    if not receipt:
        raise ValueError("Registra primero la llegada de al menos un bulto")
    old_value = str(receipt["location_text"] or "")
    if location == old_value:
        return reception_detail(connection, shipment_id, role)
    timestamp = reception_now()
    latest_receipt = connection.execute(
        "SELECT id FROM reception_receipts WHERE shipment_id = ? ORDER BY sequence_no DESC LIMIT 1",
        (shipment_id,),
    ).fetchone()
    # location_text en reception_shipments es solo un resumen de la última
    # llegada. Al corregir un bulto histórico no debe reemplazar el resumen
    # de una llegada posterior.
    shipment_location = (
        location if latest_receipt and int(latest_receipt["id"]) == int(receipt["id"])
        else str(shipment["location_text"] or "")
    )
    connection.execute(
        "UPDATE reception_receipts SET location_text = ? WHERE id = ?",
        (location, receipt["id"]),
    )
    connection.execute(
        "UPDATE reception_shipments SET location_text = ?, updated_at = ? WHERE id = ?",
        (shipment_location, timestamp, shipment_id),
    )
    attention = connection.execute(
        "SELECT * FROM reception_attentions WHERE receipt_id = ? LIMIT 1",
        (receipt["id"],),
    ).fetchone()
    if attention:
        connection.execute(
            "UPDATE reception_attentions SET location_text = ?, updated_at = ? WHERE id = ?",
            (location, timestamp, attention["id"]),
        )
    _write_history(
        connection, shipment_id, "ZONA RECEPCION", f"bulto:{receipt['sequence_no']}:location_text",
        old_value, location, username,
    )
    return reception_detail(connection, shipment_id, role)


def update_reception_validation(connection, shipment_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    if shipment["app_status"] != "VALIDACION":
        raise PermissionError("La muestra solo se registra durante VALIDACIÓN / TRANSFERENCIA")
    sample_ids = set(_validation_sample_ids(connection, shipment_id))
    checks = data.get("checks") or []
    if not isinstance(checks, list):
        raise ValueError("Formato de muestra no válido")
    known = {int(row["id"]): row for row in connection.execute(
        "SELECT * FROM reception_lines WHERE shipment_id = ?", (shipment_id,)
    ).fetchall()}
    timestamp = reception_now()
    for item in checks:
        try:
            line_id = int(item.get("line_id"))
        except (AttributeError, TypeError, ValueError):
            raise ValueError("La muestra contiene una línea no válida")
        if line_id not in sample_ids or line_id not in known:
            raise ValueError("Solo se puede validar la muestra seleccionada por el sistema")
        values = {
            "validation_sap_checked": 0,
            "validation_location_checked": int(bool(item.get("location_checked"))),
            "validation_comment_checked": 0,
        }
        validation_comment = str(item.get("comment") or "").strip()
        connection.execute(
            """UPDATE reception_lines SET validation_sap_checked = ?,
               validation_location_checked = ?, validation_comment_checked = ?,
               validation_comment = ? WHERE id = ?""",
            (values["validation_sap_checked"], values["validation_location_checked"],
             values["validation_comment_checked"], validation_comment, line_id),
        )
        values["comment"] = validation_comment
        _write_history(
            connection, shipment_id, "VALIDACION", f"linea:{line_id}:muestra",
            "", json.dumps(values), username, "Control de ubicación y comentario de muestra",
        )
    return reception_detail(connection, shipment_id, role)


def confirm_reception_transfer(connection, shipment_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    if shipment["app_status"] != "SOLICITUD TRANSFERENCIA":
        raise PermissionError("La confirmación administrativa solo aplica a una transferencia solicitada")
    if role not in RECEPTION_ROLES:
        raise PermissionError("No tienes permiso para confirmar la solicitud")
    if not bool(data.get("checked")):
        raise ValueError("Marca la solicitud como revisada")
    if int(shipment["transfer_assistant_checked"] or 0):
        return reception_detail(connection, shipment_id, role)
    timestamp = reception_now()
    connection.execute(
        "UPDATE reception_shipments SET transfer_assistant_checked = 1, updated_at = ? WHERE id = ?",
        (timestamp, shipment_id),
    )
    _write_history(connection, shipment_id, "TRANSFERENCIA", "transfer_assistant_checked", 0, 1, username)
    return reception_detail(connection, shipment_id, role)


def update_reception_references(connection, shipment_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    changed = False
    if "expected_packages" in data:
        raise ValueError("Los bultos esperados provienen de la fuente de recepción y no se editan manualmente")
    if (
        "em_number" in data
        and role == "ADMINISTRADOR"
        and shipment["app_status"] == "REVISION SISTEMA"
        and not str(shipment["em_number"] or "").strip()
    ):
        raise PermissionError("Completa la revisión de sistema y pasa a la etapa EM antes de registrar la EM")
    if (
        "em_number" in data
        and role != "ADMINISTRADOR"
        and _is_assigned_to(shipment, username)
        and shipment["app_status"] != "EM"
    ):
        raise PermissionError("La EM solo se registra dentro de la etapa EM")
    allowed = set()
    if role == "ADMINISTRADOR":
        allowed |= {"fr_number", "em_number", "current_assistant", "current_auxiliary"}
    elif _is_assigned_to(shipment, username) and shipment["app_status"] == "EM":
        # La revisión de sistema solo valida cantidades y observaciones.
        # El trabajador registra el número de EM únicamente después de entrar
        # a la etapa EM. La creación automática en SAP se conectará aquí
        # cuando TI entregue las credenciales del Service Layer.
        if "em_number" in data:
            accounting_status = _accounting_status_for_shipment(connection, shipment_id, shipment)
            if accounting_status != "PENDIENTE EM":
                raise PermissionError(
                    "La EM solo puede registrarse cuando todas las IP ya tienen factura de reserva"
                )
            if str(shipment["em_number"] or "").strip() and str(data.get("em_number") or "").strip() != str(shipment["em_number"] or "").strip():
                raise PermissionError("La EM ya registrada solo puede corregirla el administrador")
            allowed.add("em_number")
    for field in allowed:
        if field not in data:
            continue
        value = str(data.get(field) or "").strip()
        old_value = str(shipment[field] or "")
        if value == old_value:
            continue
        connection.execute(
            f"UPDATE reception_shipments SET {field} = ?, updated_at = ? WHERE id = ?",
            (value, reception_now(), shipment_id),
        )
        _write_history(connection, shipment_id, "MODIFICACION", field, old_value, value, username)
        changed = True
    if not changed:
        raise ValueError("No se recibió ningún cambio")
    _refresh_manual_accounting_status(connection, shipment_id, username)
    return reception_detail(connection, shipment_id, role)


def update_reception_line_quantity(connection, shipment_id, line_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    if shipment["app_status"] != "REVISION SISTEMA":
        raise PermissionError("Las cantidades verificadas solo se modifican durante REVISIÓN DE SISTEMA")
    attention = _ensure_reception_attention(connection, shipment, username)
    line = connection.execute(
        "SELECT * FROM reception_lines WHERE id = ? AND shipment_id = ?",
        (line_id, shipment_id),
    ).fetchone()
    if not line:
        raise ValueError("Línea NP no encontrada")
    received_qty = float(data.get("received_qty") or 0)
    if received_qty < 0:
        raise ValueError("La cantidad recibida no puede ser negativa")
    reason = str(data.get("reason") or "").strip()
    attention_line = None
    if attention:
        attention_line = connection.execute(
            "SELECT * FROM reception_attention_lines WHERE attention_id = ? AND reception_line_id = ?",
            (attention["id"], line_id),
        ).fetchone()
    # El esperado es el valor inicial de conteo, no un tope. El físico puede
    # traer faltantes o excedentes y ambos deben registrarse para que el
    # reporte muestre la diferencia y se notifique a Importaciones.
    old_qty = float(attention_line["verified_qty"] or 0) if attention_line else float(line["received_qty"] or 0)
    if attention_line:
        connection.execute(
            "UPDATE reception_attention_lines SET verified_qty = ? WHERE id = ?",
            (received_qty, attention_line["id"]),
        )
        # Una edición explícita confirma que esta atención ya fue inicializada
        # y evita volver a mostrar el esperado como valor predeterminado en
        # una siguiente consulta.
        connection.execute(
            "UPDATE reception_attentions SET system_quantities_initialized = 1 WHERE id = ?",
            (attention["id"],),
        )
    connection.execute(
        "UPDATE reception_lines SET blocked_reason = ? WHERE id = ?",
        (reason, line_id),
    )
    _sync_reception_line_totals(connection, shipment_id)
    _write_history(
        connection,
        shipment_id,
        "REVISION SISTEMA",
        f"linea:{line_id}:received_qty",
        old_qty,
        received_qty,
        username,
        reason,
    )
    if reason or received_qty != float(line["expected_qty"] or 0):
        _enqueue_notification(
            connection,
            shipment_id,
            "OBSERVACION_IMPORTACIONES",
            f"Observación en revisión de sistema · BL/AWB {shipment['bl_awb']}",
            (
                f"BL/AWB: {shipment['bl_awb']}\n"
                f"NP: {line['np_code'] or '—'}\n"
                f"Descripción: {line['description'] or '—'}\n"
                f"OC: {line['oc_number'] or '—'} · OV: {line['ov_number'] or '—'}\n"
                f"Esperado: {line['expected_qty']} · Verificado: {received_qty}\n"
                f"Observación: {reason or 'Diferencia detectada'}\n"
                f"Registrado por: {username}"
            ),
            "TRITON_IMPORTACIONES_EMAILS",
            "sgallo@triton.com.pe;claudia.acedo@triton.com.pe;jorge.cucho@triton.com.pe;pfigueroa@triton.com.pe",
        )
    return reception_detail(connection, shipment_id, role)


def change_reception_status(connection, shipment_id, data, username, role):
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    attention = _ensure_reception_attention(connection, shipment, username)
    old_status = shipment["app_status"]
    new_status = str(data.get("status") or "").strip().upper()
    if new_status not in RECEPTION_STATE_INDEX:
        raise ValueError("Estado de recepción no válido")

    validation_result = str(data.get("validation_result") or "").strip().upper()
    if old_status == "VALIDACION" and validation_result == "NO CONFORME":
        new_status = "UBICACION"
        condition = "REUBICACION REQUERIDA"
        connection.execute(
            "UPDATE reception_shipments SET app_status = ?, condition_status = ?, updated_at = ? WHERE id = ?",
            (new_status, condition, reception_now(), shipment_id),
        )
        if attention:
            connection.execute(
                "UPDATE reception_attentions SET app_status = ?, condition_status = ?, updated_at = ? WHERE id = ?",
                (new_status, condition, reception_now(), attention["id"]),
            )
        _write_history(
            connection, shipment_id, "VALIDACION", "app_status", old_status, new_status,
            username, "Muestra de ubicación no conforme; se exige nueva ubicación y nueva muestra",
        )
        return reception_detail(connection, shipment_id, role)

    expected_next = RECEPTION_STATE_INDEX[old_status] + 1
    if RECEPTION_STATE_INDEX[new_status] != expected_next:
        raise PermissionError("El flujo de recepción debe avanzar una etapa a la vez")
    expected_packages = float(shipment["expected_packages"] or 0)
    received_packages = float(shipment["received_packages"] or 0)
    # El saldo de bultos no bloquea el trabajo previo: pueden existir arribos
    # acumulados y el trabajador puede revisar lo disponible. La ventana de
    # bultos permanece abierta hasta que el usuario la cierre explícitamente.
    # Consultar las referencias actuales evita aceptar una FR parcial por IP
    # o depender de un resumen contable desactualizado. No modifica la BD.
    accounting_status = _accounting_status_for_shipment(connection, shipment_id, shipment)
    if new_status == "REVISION SISTEMA":
        if not str(shipment["current_assistant"] or "").strip() or not str(shipment["current_auxiliary"] or "").strip():
            raise PermissionError(
                "Asigna un asistente y un auxiliar de recepción antes de iniciar la revisión de sistema"
            )
        if data.get("require_location"):
            missing_locations = connection.execute(
                """SELECT COUNT(*) FROM reception_receipts
                    WHERE shipment_id = ? AND TRIM(COALESCE(location_text, '')) = ''""",
                (shipment_id,),
            ).fetchone()[0]
            if missing_locations:
                raise PermissionError("Registra la zona de recepción de cada bulto antes de iniciar el conteo")
    if new_status == "EM" and accounting_status in {"PENDIENTE CONTABILIDAD", "PENDIENTE FR"}:
        raise PermissionError(
            "Registra la factura de reserva de todas las referencias antes de pasar a EM"
        )
    if new_status == "EM" and accounting_status == "CONFLICTO IP":
        raise PermissionError(
            "La BL tiene una IP relacionada con más de una FR/EM; corrige el Excel contable antes de pasar a EM"
        )
    if new_status == "REVISION SISTEMA" and accounting_status == "EM REGISTRADA" and not (
        attention and old_status == "ARRIBADO"
    ):
        raise PermissionError("La BL ya tiene EM registrada; no requiere trabajo de recepción")
    if new_status == "UBICACION" and not str(shipment["em_number"] or "").strip():
        raise PermissionError("Registra la EM antes de iniciar ubicación")
    if old_status == "VALIDACION" and validation_result != "CONFORME":
        raise PermissionError("Indica si la muestra de ubicación fue CONFORME o NO CONFORME")
    if old_status == "VALIDACION" and validation_result == "CONFORME":
        sample_ids = _validation_sample_ids(connection, shipment_id)
        if sample_ids:
            placeholders = ",".join("?" for _ in sample_ids)
            sample_rows = connection.execute(
                f"SELECT validation_sap_checked, validation_location_checked, validation_comment_checked, blocked_reason "
                f"FROM reception_lines WHERE shipment_id = ? AND id IN ({placeholders})",
                [shipment_id, *sample_ids],
            ).fetchall()
            if any(not row["validation_location_checked"] for row in sample_rows):
                raise PermissionError("Completa el check de ubicación de toda la muestra")
    if new_status == "CERRADO" and not int(shipment["transfer_assistant_checked"] or 0):
        raise PermissionError("La solicitud de transferencia debe ser revisada y marcada antes de cerrar")

    condition = "EN PROCESO" if new_status != "CERRADO" else "COMPLETADO"
    # La inicialización pertenece a cada atención, no a la BL completa. Una
    # BL puede tener Atención 1 cerrada y una nueva llegada (Atención 2); la
    # segunda también debe arrancar con su esperado al 100% aunque la BL ya
    # haya pasado antes por revisión de sistema.
    attention_initialized = (
        int(attention["system_quantities_initialized"] or 0)
        if attention else int(shipment["system_quantities_initialized"] or 0)
    )
    initialize_system = new_status == "REVISION SISTEMA" and not attention_initialized
    if initialize_system:
        if attention:
            connection.execute(
                "UPDATE reception_attention_lines SET verified_qty = planned_qty WHERE attention_id = ?",
                (attention["id"],),
            )
            connection.execute(
                "UPDATE reception_attentions SET system_quantities_initialized = 1 WHERE id = ?",
                (attention["id"],),
            )
            _sync_reception_line_totals(connection, shipment_id)
        else:
            connection.execute(
                "UPDATE reception_lines SET received_qty = expected_qty WHERE shipment_id = ?",
                (shipment_id,),
            )
        connection.execute(
            """UPDATE reception_shipments
               SET app_status = ?, condition_status = ?, system_quantities_initialized = 1, updated_at = ?
               WHERE id = ?""",
            (new_status, condition, reception_now(), shipment_id),
        )
        line_count = connection.execute(
            "SELECT COUNT(*) FROM reception_lines WHERE shipment_id = ?", (shipment_id,)
        ).fetchone()[0]
        _write_history(
            connection,
            shipment_id,
            "INICIALIZACION",
            "cantidades_revision_sistema",
            0,
            "100%",
            username,
            f"{line_count} líneas inicializadas al 100% al iniciar revisión de sistema",
        )
    else:
        connection.execute(
            "UPDATE reception_shipments SET app_status = ?, condition_status = ?, updated_at = ? WHERE id = ?",
            (new_status, condition, reception_now(), shipment_id),
        )
    if attention:
        timestamp = reception_now()
        connection.execute(
            """UPDATE reception_attentions
                  SET app_status = ?, condition_status = ?,
                      updated_at = ?, completed_at = CASE WHEN ? = 'CERRADO' THEN ? ELSE completed_at END
                WHERE id = ?""",
            (new_status, condition, timestamp, new_status, timestamp, attention["id"]),
        )
    _write_history(connection, shipment_id, "ESTADO", "app_status", old_status, new_status, username)
    if old_status == "REVISION SISTEMA" and new_status == "EM" and advanced_lots_enabled():
        created_lots = generate_lots_for_shipment(connection, shipment_id, username)
        _write_history(
            connection,
            shipment_id,
            "LOTES",
            "lotes_generados",
            0,
            len(created_lots),
            username,
            "Lotes internos generados con las cantidades verificadas",
        )
    if old_status == "VALIDACION" and new_status == "SOLICITUD TRANSFERENCIA":
        already_requested = connection.execute(
            """SELECT 1 FROM reception_notifications
               WHERE shipment_id = ? AND notification_type = 'SOLICITUD_TRANSFERENCIA'
               LIMIT 1""",
            (shipment_id,),
        ).fetchone()
        if not already_requested:
            _enqueue_notification(
                connection,
                shipment_id,
                "SOLICITUD_TRANSFERENCIA",
                f"Solicitud de transferencia IMP a 01 · BL/AWB {shipment['bl_awb']}",
                (
                    f"BL/AWB: {shipment['bl_awb']}\n"
                    f"EM: {shipment['em_number'] or '—'}\n"
                    "La validación de ubicación fue CONFORME. Solicitud de transferencia de IMP a 01.\n"
                    f"Registrado por: {username}"
                ),
                "TRITON_CONTABILIDAD_EMAILS",
                "mpucurimay@triton.com.pe",
            )
    return reception_detail(connection, shipment_id, role)


def rewind_reception_stage(connection, shipment_id, data, username, role):
    """Reopen a completed reception stage for an administrator.

    Workers can navigate to prior stages only as read-only. This separate
    operation keeps the normal forward-only workflow intact and records the
    administrative reason without deleting the existing audit trail.
    """
    shipment = _get_reception_shipment(connection, shipment_id, username, role)
    if role != "ADMINISTRADOR":
        raise PermissionError("Solo el administrador puede reabrir una etapa")

    current_status = str(shipment["app_status"] or "PROGRAMADO")
    target_status = str(data.get("status") or "").strip().upper()
    if target_status not in RECEPTION_STATE_INDEX:
        raise ValueError("Etapa de reapertura no válida")
    if RECEPTION_STATE_INDEX[target_status] >= RECEPTION_STATE_INDEX[current_status]:
        raise PermissionError("La reapertura debe devolver la BL a una etapa anterior")

    reason = str(data.get("reason") or "").strip()
    if len(reason) < 5:
        raise ValueError("Indica el motivo de la reapertura")
    if target_status == "PROGRAMADO" and (
        float(shipment["received_packages"] or 0) > 0
        or connection.execute(
            "SELECT 1 FROM reception_receipts WHERE shipment_id = ? LIMIT 1",
            (shipment_id,),
        ).fetchone()
    ):
        raise PermissionError(
            "No se puede devolver a PROGRAMADO porque la BL ya tiene bultos registrados"
        )

    timestamp = reception_now()
    target_condition = {
        "PROGRAMADO": "POR ARRIBAR",
        "ARRIBADO": "ARRIBO REGISTRADO",
    }.get(target_status, "EN PROCESO")
    connection.execute(
        """UPDATE reception_shipments
           SET app_status = ?, condition_status = ?,
               validation_sample_json = CASE WHEN ? < ? THEN NULL ELSE validation_sample_json END,
               transfer_assistant_checked = CASE WHEN ? < ? THEN 0 ELSE transfer_assistant_checked END,
               system_quantities_initialized = CASE WHEN ? < ? THEN 0 ELSE system_quantities_initialized END,
               updated_at = ?
         WHERE id = ?""",
        (
            target_status,
            target_condition,
            RECEPTION_STATE_INDEX[target_status],
            RECEPTION_STATE_INDEX["VALIDACION"],
            RECEPTION_STATE_INDEX[target_status],
            RECEPTION_STATE_INDEX["SOLICITUD TRANSFERENCIA"],
            RECEPTION_STATE_INDEX[target_status],
            RECEPTION_STATE_INDEX["REVISION SISTEMA"],
            timestamp,
            shipment_id,
        ),
    )

    # Reset only controls that belong to stages after the selected reopening
    # point. The physical receipt and its manual location remain evidence.
    if RECEPTION_STATE_INDEX[target_status] < RECEPTION_STATE_INDEX["VALIDACION"]:
        connection.execute(
            """UPDATE reception_lines
               SET validation_sap_checked = 0,
                   validation_location_checked = 0,
                   validation_comment_checked = 0
             WHERE shipment_id = ?""",
            (shipment_id,),
        )
    attention = _active_reception_attention(connection, shipment_id)
    if attention:
        target_index = RECEPTION_STATE_INDEX[target_status]
        if target_status == "REVISION SISTEMA":
            attention_system_init = 1
        elif target_index < RECEPTION_STATE_INDEX["REVISION SISTEMA"]:
            attention_system_init = 0
        else:
            attention_system_init = int(attention["system_quantities_initialized"] or 0)
        attention_transfer_check = (
            0
            if RECEPTION_STATE_INDEX[target_status] < RECEPTION_STATE_INDEX["SOLICITUD TRANSFERENCIA"]
            else int(attention["transfer_assistant_checked"] or 0)
        )
        connection.execute(
            """UPDATE reception_attentions
               SET app_status = ?, condition_status = ?,
                   system_quantities_initialized = ?,
                   transfer_assistant_checked = ?,
                   completed_at = NULL, updated_at = ?
             WHERE id = ?""",
            (
                target_status,
                target_condition,
                attention_system_init,
                attention_transfer_check,
                timestamp,
                attention["id"],
            ),
        )
        if RECEPTION_STATE_INDEX[target_status] < RECEPTION_STATE_INDEX["REVISION SISTEMA"]:
            connection.execute(
                "UPDATE reception_attention_lines SET verified_qty = 0 WHERE attention_id = ?",
                (attention["id"],),
            )
        elif target_status == "REVISION SISTEMA":
            connection.execute(
                "UPDATE reception_attention_lines SET verified_qty = planned_qty WHERE attention_id = ?",
                (attention["id"],),
            )

    _write_history(
        connection,
        shipment_id,
        "RETROCESO_ADMIN",
        "app_status",
        current_status,
        target_status,
        username,
        reason,
    )
    return reception_detail(connection, shipment_id, role)
