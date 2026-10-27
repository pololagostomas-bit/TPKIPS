"""Shared, opt-in purchase monitoring; never changes operational WMS data."""
import hashlib
import html
import io
import json
import os
import re
import sqlite3
import string
import threading
import time
import uuid
from contextlib import contextmanager
from datetime import date, datetime, timezone
from email.utils import getaddresses
from pathlib import Path
from zoneinfo import ZoneInfo
from urllib.parse import parse_qs, urlsplit

import requests
from backend.services.safe_excel import load_workbook
from openpyxl.utils.datetime import from_excel

from backend.services import cloud_auth, cloud_excel, purchase_history
from backend.services.reception import _cell_text, _normalize_header

DEFAULTS = {
    'enabled': False, 'sender': '', 'sender_name': '', 'to': [], 'cc': [],
    'interval_minutes': 60, 'daily_time': '10:00', 'weekdays': [0, 1, 2, 3, 4],
    'timezone': 'America/Lima', 'excluded_sheets': ['ANIBAL.23.04'],
    'subject': '[Control de compras] OC por regularizar - {fecha}',
    'new_subject': '[Control de compras] Nuevos casos OC por regularizar - {fecha}',
    'body': ('Estimados:\n\nEn la revision de Importaciones se identificaron pedidos Triton '
             'con una referencia preliminar o una OC de mas de cinco digitos.\n\n'
             'Conforme al lineamiento de Gerencia, no deben efectuarse compras sin una OC '
             'previamente aprobada. Favor validar los casos detallados y actualizar el Excel '
             'con la OC aprobada de cinco digitos.\n\n'
             'El reporte incluira comprador, preliminar y la IP/Pedido Triton.'),
    'new_intro': ('Se adicionaron {nuevos} casos desde el ultimo aviso. '
                  'En total, quedan {pendientes} casos pendientes.'),
}
TOKENS = {'fecha', 'total', 'nuevos', 'pendientes'}
REQUEST_DATE_HEADERS = {
    'FECHA DE SOLICITUD DE IMP OV', 'FECHA SOLICITUD DE IMP OV',
    'FECHA DE SOLICITUD IMP OV', 'FECHA SOLICITUD IMP OV',
}
_guard = threading.Lock()
_worker_started = False


def state_path():
    return Path(os.getenv('WMS_PURCHASE_ALERT_DB', str(cloud_auth.cache_path().parent / 'purchase-alerts.sqlite3')))


def _folder():
    path = state_path().parent
    path.mkdir(parents=True, exist_ok=True, mode=0o700)
    if os.name != 'nt':
        path.chmod(0o700)
    return path


@contextmanager
def db():
    _folder()
    path = state_path()
    descriptor = os.open(path, os.O_CREAT | os.O_RDWR, 0o600)
    os.close(descriptor)
    if os.name != 'nt':
        path.chmod(0o600)
    connection = sqlite3.connect(path, timeout=10)
    connection.row_factory = sqlite3.Row
    try:
        connection.executescript('''
            CREATE TABLE IF NOT EXISTS settings(id INTEGER PRIMARY KEY CHECK(id=1), value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS state(id INTEGER PRIMARY KEY CHECK(id=1), value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS reported(case_key TEXT PRIMARY KEY);
            CREATE TABLE IF NOT EXISTS outbox(
                id TEXT PRIMARY KEY, event_key TEXT UNIQUE NOT NULL, kind TEXT NOT NULL,
                created_at TEXT NOT NULL, sent_at TEXT, status TEXT NOT NULL,
                detail TEXT NOT NULL DEFAULT '', payload TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS audit(at TEXT NOT NULL, actor TEXT NOT NULL, action TEXT NOT NULL);
        ''')
        purchase_history.initialize(connection)
        with connection:
            yield connection
    finally:
        connection.close()


def _get(connection, table, default):
    row = connection.execute(f'SELECT value FROM {table} WHERE id=1').fetchone()
    return json.loads(row['value']) if row else json.loads(json.dumps(default))


def _set(connection, table, value):
    connection.execute(f'INSERT INTO {table}(id,value) VALUES(1,?) '
        'ON CONFLICT(id) DO UPDATE SET value=excluded.value', (json.dumps(value, ensure_ascii=False),))


def settings():
    with db() as connection:
        return {**DEFAULTS, **_get(connection, 'settings', {})}


def _recipients(values):
    if not isinstance(values, list) or len(values) > 50:
        raise ValueError('Usa hasta 50 destinatarios por lista.')
    result, seen = [], set()
    for value in values:
        if not isinstance(value, str) or len(value) > 320 or '\r' in value or '\n' in value:
            raise ValueError('Destinatario no valido.')
        addresses = getaddresses([value])
        if len(addresses) != 1:
            raise ValueError('Ingresa un destinatario por linea.')
        name, address = addresses[0]
        if not re.fullmatch(r'[A-Za-z0-9.!#$%&\x27*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+', address):
            raise ValueError('Revisa las direcciones de correo.')
        key = address.casefold()
        if key not in seen:
            seen.add(key)
            result.append({'name': name.strip(), 'address': address})
    return result


def _addresses(values):
    return [f'{item["name"]} <{item["address"]}>' if item['name'] else item['address'] for item in values]


def validate(data):
    if not isinstance(data, dict) or set(data) - set(DEFAULTS):
        raise ValueError('Configuracion no valida.')
    value = {**settings(), **data}
    if not isinstance(value['enabled'], bool):
        raise ValueError('Estado de automatizacion no valido.')
    for field, maximum in [('sender', 254), ('sender_name', 120), ('subject', 200),
                            ('new_subject', 200), ('body', 10000), ('new_intro', 2000)]:
        text = value[field]
        if not isinstance(text, str) or len(text) > maximum or '\x00' in text:
            raise ValueError('Revisa el texto de la configuracion.')
        if field in {'sender', 'sender_name', 'subject', 'new_subject'} and ('\n' in text or '\r' in text):
            raise ValueError('Los encabezados deben tener una sola linea.')
        if field in {'subject', 'new_subject', 'body', 'new_intro'}:
            try:
                for _, token, spec, conversion in string.Formatter().parse(text):
                    if token is not None and (token not in TOKENS or spec or conversion):
                        raise ValueError('Marcador de texto no permitido.')
            except ValueError as error:
                raise ValueError('Marcadores permitidos: {fecha}, {total}, {nuevos}, {pendientes}.') from error
        value[field] = text.strip()
    if value['sender']:
        parsed = _recipients([value['sender']])
        if parsed[0]['name']:
            raise ValueError('El remitente debe ser solo su direccion de correo.')
        value['sender'] = parsed[0]['address']
    to, cc = _recipients(value['to']), _recipients(value['cc'])
    to_keys = {item['address'].casefold() for item in to}
    value['to'] = _addresses(to)
    value['cc'] = _addresses([item for item in cc if item['address'].casefold() not in to_keys])
    if type(value['interval_minutes']) is not int or not 5 <= value['interval_minutes'] <= 1440:
        raise ValueError('El intervalo debe estar entre 5 y 1440 minutos.')
    if not isinstance(value['daily_time'], str) or not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d', value['daily_time']):
        raise ValueError('Hora no valida; usa HH:MM.')
    if not isinstance(value['weekdays'], list) or not value['weekdays'] or any(type(day) is not int or day not in range(7) for day in value['weekdays']):
        raise ValueError('Selecciona al menos un dia para el resumen.')
    value['weekdays'] = sorted(set(value['weekdays']))
    if value['timezone'] != 'America/Lima':
        raise ValueError('Los horarios se configuran en hora de Lima.')
    excluded = value['excluded_sheets']
    if not isinstance(excluded, list) or len(excluded) > 50 or any(not isinstance(s, str) or not s.strip() or len(s) > 31 for s in excluded):
        raise ValueError('Revisa las hojas excluidas.')
    value['excluded_sheets'] = sorted({sheet.strip() for sheet in excluded})
    if value['enabled'] and (not value['sender'] or not value['to'] or not value['body'] or not value['subject'] or not value['new_subject']):
        raise ValueError('Completa remitente, destinatarios, asunto y mensaje antes de activar.')
    return value


def save(data, actor):
    value = validate(data)
    with db() as connection:
        _set(connection, 'settings', value)
        state = _get(connection, 'state', {})
        state['next_check'] = 0
        _set(connection, 'state', state)
        connection.execute('INSERT INTO audit VALUES(?,?,?)', (_stamp(), actor, 'CONFIGURACION'))
    return status()


def _stamp(now=None):
    return datetime.fromtimestamp(time.time() if now is None else now, timezone.utc).isoformat(timespec='seconds')


@contextmanager
def job_lock():
    if not _guard.acquire(blocking=False):
        yield False
        return
    descriptor, acquired = None, False
    try:
        descriptor = os.open(_folder() / 'purchase-alerts.lock', os.O_CREAT | os.O_RDWR, 0o600)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(descriptor, msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(descriptor, fcntl.LOCK_EX | fcntl.LOCK_NB)
            acquired = True
        except OSError:
            pass
        yield acquired
    finally:
        if descriptor is not None:
            if acquired:
                if os.name == 'nt':
                    import msvcrt
                    os.lseek(descriptor, 0, os.SEEK_SET)
                    msvcrt.locking(descriptor, msvcrt.LK_UNLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(descriptor, fcntl.LOCK_UN)
            os.close(descriptor)
        _guard.release()


def _request_date(value, epoch):
    if type(value) in (int, float):
        try:
            value = from_excel(value, epoch) if value > 0 else None
        except (ValueError, OverflowError):
            return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat() if value.year >= 1900 else None
    if isinstance(value, str):
        for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d'):
            try:
                parsed = datetime.strptime(value.strip(), fmt).date()
                return parsed.isoformat() if parsed.year >= 1900 else None
            except ValueError:
                pass
    return None


def read_cases(content, excluded_sheets):
    """Inspect purchase rows even when no shipment or part code exists yet."""
    books = []
    try:
        for mode in (True, False):
            books.append(load_workbook(io.BytesIO(content), read_only=True, data_only=mode))
        values, formulas = books
        cases, checked, excluded, skipped, ov_sheets = {}, [], [], [], []
        exclude = {_normalize_header(sheet) for sheet in excluded_sheets}
        for sheet in values.worksheets:
            if _normalize_header(sheet.title) in exclude:
                excluded.append(sheet.title)
                continue
            header = None
            for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True), 1):
                headers = [_normalize_header(cell) for cell in row]
                if 'ORDEN DE COMPRA' in headers and 'PEDIDO TRITON' in headers:
                    ov_index = next((headers.index(name) for name in ('OV', 'ORDEN DE VENTA') if name in headers), None)
                    date_index = next((index for index, name in enumerate(headers)
                        if ' '.join(name.replace('/', ' ').split()) in REQUEST_DATE_HEADERS), None)
                    header = (row_number, headers.index('ORDEN DE COMPRA'), headers.index('PEDIDO TRITON'), ov_index, date_index)
                    break
            if header is None:
                skipped.append(sheet.title)
                continue
            checked.append(sheet.title)
            number, oc_index, ip_index, ov_index, date_index = header
            if ov_index is not None:
                ov_sheets.append(sheet.title)
            columns = [oc_index, ip_index] + ([ov_index] if ov_index is not None else [])
            max_col = max(columns + ([date_index] if date_index is not None else [])) + 1
            rows = sheet.iter_rows(min_row=number + 1, max_col=max_col, values_only=True)
            raw_rows = formulas[sheet.title].iter_rows(min_row=number + 1, max_col=max_col)
            for row_number, (row, raw) in enumerate(zip(rows, raw_rows), number + 1):
                # X is an explicit opt-out even if the row has a COT reference.
                ip = _cell_text(row[ip_index])
                if ip.strip().casefold() == 'x':
                    continue
                for index in columns:
                    if raw[index].data_type == 'e' or (raw[index].data_type == 'f' and row[index] is None):
                        raise ValueError('Hay errores o formulas sin resultado guardado en OC/Pedido Triton.')
                oc = _cell_text(row[oc_index])
                ov = _cell_text(row[ov_index]) if ov_index is not None else ''
                rules = []
                if ip and (sum(char.isdigit() for char in oc) > 5 or 'pre' in oc.casefold()):
                    rules.append(('PRELIMINAR', 'REGULARIZAR'))
                if ov.upper().startswith('COT') and (oc or ip):
                    rules.append(('COT', 'GRAVE'))
                request_date = _request_date(row[date_index], values.epoch) if date_index is not None and rules else None
                for issue_type, severity in rules:
                    identity = _normalize_header(ip) if ip else 'OC:' + _normalize_header(oc)
                    history_identity = issue_type + '\x00' + identity
                    if issue_type == 'COT':
                        history_identity += '\x00' + _normalize_header(ov)
                    # Preserve the existing ordinary mail dedupe keys across upgrades.
                    key_text = _normalize_header(oc) + '\x00' + _normalize_header(ip)
                    if issue_type == 'COT':
                        key_text += '\x00COT\x00' + _normalize_header(ov)
                    key = hashlib.sha256(key_text.encode()).hexdigest()
                    item = cases.setdefault(key, {'key': key, 'preliminary': oc, 'ip': ip,
                        'ov': ov, 'issue_type': issue_type, 'severity': severity,
                        'history_key': hashlib.sha256(history_identity.encode()).hexdigest(),
                        'sources': [], 'rows': 0})
                    item['rows'] += 1
                    item['sources'].append({'sheet': sheet.title, 'row': row_number,
                                            'request_date': request_date})
        if not checked:
            raise ValueError('No se encontraron hojas con Orden de compra y Pedido Triton.')
        return {'cases': sorted(cases.values(), key=lambda item: (item['sources'][0]['sheet'], item['ip'], item['preliminary'])),
                'sheets': checked, 'excluded_sheets': excluded, 'skipped_sheets': skipped,
                'ov_sheets': ov_sheets}
    finally:
        for book in books:
            book.close()


def fetch_snapshot(config):
    sources = [item for item in cloud_excel.configured_sources() if item['source_type'] == 'importation']
    if len(sources) != 1:
        raise ValueError('Configura el Excel de Importaciones en Cortes diarios.')
    source = cloud_excel.download_cloud_excels(sources=sources)[0]
    snapshot = read_cases(source['content'], config['excluded_sheets'])
    snapshot.update(modified_at=source.get('modified_at'), checked_at=_stamp())
    return snapshot


def render(config, snapshot, new_keys, kind='daily'):
    cases = snapshot['cases']
    local = datetime.fromisoformat(snapshot['checked_at']).astimezone(ZoneInfo(config['timezone']))
    tokens = {'fecha': local.strftime('%d/%m/%Y %H:%M'), 'total': len(cases),
              'nuevos': len(new_keys), 'pendientes': len(cases)}
    subject = config['new_subject'] if kind == 'incremental' else config['subject']
    message = html.escape(config['body'].format(**tokens)).replace('\n', '<br>')
    intro = config['new_intro'].format(**tokens) if kind == 'incremental' else f"Casos por regularizar: {len(cases)}."
    headings = ['Nivel', 'Comprador / hoja', 'Preliminar / OC registrada',
                'IP / Pedido Triton', 'OV', 'Fecha de solicitud IMP/OV', 'Desde deteccion']

    def table(items):
        rows = []
        for item in items:
            buyers = ', '.join(sorted({source['sheet'] for source in item['sources']}))
            first_seen = item.get('first_seen_at')
            elapsed = round(max(0, (datetime.fromisoformat(snapshot['checked_at']) -
                datetime.fromisoformat(first_seen)).total_seconds()) / 3600, 1) if first_seen else None
            cells = [item.get('severity', 'REGULARIZAR'), buyers, item['preliminary'], item['ip'],
                     item.get('ov', ''), purchase_history.request_dates_text(item),
                     f'{elapsed} h' if elapsed is not None else 'Sin medicion']
            rows.append('<tr>' + ''.join('<td style="border:1px solid #ccc;padding:8px;overflow-wrap:anywhere">'
                + '<span class="cell-label" style="display:none">' + heading + '</span>'
                + html.escape(str(cell)) + '</td>' for heading, cell in zip(headings, cells)) + '</tr>')
        return ('<table class="purchase-report" style="border-collapse:collapse;width:100%;table-layout:fixed"><thead><tr>' +
                ''.join('<th scope="col" style="border:1px solid #ccc;padding:8px;text-align:left;overflow-wrap:anywhere">' + heading + '</th>'
                        for heading in headings) +
                '</tr></thead><tbody>' + ''.join(rows) + '</tbody></table>')

    body = ('<!doctype html><html lang="es"><head><meta name="viewport" content="width=device-width,initial-scale=1">'
            '<style>body{font:14px Arial,sans-serif;color:#29343b;overflow-wrap:anywhere}h2{font-size:18px}h3{font-size:16px}'
            '@media(max-width:600px){.purchase-report,.purchase-report tbody,.purchase-report tr,.purchase-report td'
            '{display:block;width:100%;box-sizing:border-box}.purchase-report thead{display:none}'
            '.purchase-report tr{margin:12px 0}.purchase-report .cell-label{display:block!important;font-weight:bold;margin-bottom:4px}}</style>'
            '</head><body><p>' + message + '</p><p>' + html.escape(intro) + '</p>')
    if any(item.get('issue_type') == 'COT' for item in cases):
        body += '<p><strong>GRAVE: cotizacion COT con OC o IP registrada. Regularizar con prioridad; una cotizacion no habilita una compra.</strong></p>'
    cases = sorted(cases, key=lambda item: item.get('severity') != 'GRAVE')
    for grave, title in [(True, 'Graves: cotizacion con OC o pedido a fabrica'),
                         (False, 'Por regularizar: preliminar / OC con pedido a fabrica')]:
        group = [item for item in cases if (item.get('severity') == 'GRAVE') == grave]
        if not group:
            continue
        body += '<h2>' + title + f' ({len(group)})</h2>'
        if kind == 'incremental':
            new = [item for item in group if item['key'] in new_keys]
            old = [item for item in group if item['key'] not in new_keys]
            if new:
                body += '<h3>Casos adicionados</h3>' + table(new)
            if old:
                body += '<h3>Demas pendientes</h3>' + table(old)
        else:
            body += table(group)
    body += '<p>Saludos,<br>' + html.escape(config['sender_name']) + '</p>'
    body += '<p>Revision: ' + html.escape(tokens['fecha']) + ' (Lima).</p></body></html>'
    return {'subject': subject.format(**tokens), 'html': body, 'case_keys': [item['key'] for item in cases],
            'to': config['to'], 'cc': config['cc'], 'sender': config['sender']}


def status():
    with db() as connection:
        config = {**DEFAULTS, **_get(connection, 'settings', {})}
        state = _get(connection, 'state', {})
        history = [dict(row) for row in connection.execute('SELECT id,kind,created_at,sent_at,status,detail '
                   'FROM outbox ORDER BY created_at DESC LIMIT 30')]
    auth = cloud_auth.status()
    return {'settings': config, 'state': state, 'history': history,
            'worker_enabled': os.getenv('WMS_PURCHASE_ALERT_WORKER', '0') == '1',
            'mail_authorized': bool(config['sender'] and auth.get('mail_username', '').casefold() == config['sender'].casefold()),
            'microsoft_authorized': auth['authorized']}


def case_history(**options):
    with db() as connection:
        return purchase_history.query(connection, **options)


def history_csv(**options):
    with db() as connection:
        return purchase_history.csv_bytes(purchase_history.query(connection, export=True, **options))


def _daily_due(config, state, now):
    local = datetime.fromtimestamp(now, timezone.utc).astimezone(ZoneInfo(config['timezone']))
    return (local.weekday() in config['weekdays'] and local.strftime('%H:%M') >= config['daily_time']
            and state.get('last_daily') != local.date().isoformat())


def _post_mail(token, payload):
    def recipients(field):
        return [{'emailAddress': item} for item in _recipients(payload[field])]
    message = {'subject': payload['subject'], 'body': {'contentType': 'HTML', 'content': payload['html']},
               'toRecipients': recipients('to'), 'ccRecipients': recipients('cc')}
    return requests.post('https://graph.microsoft.com/v1.0/me/sendMail',
        headers={'Authorization': 'Bearer ' + token}, json={'message': message, 'saveToSentItems': True},
        timeout=30, allow_redirects=False)


def run(force=False, send_now=False, now=None):
    now = time.time() if now is None else now
    with job_lock() as acquired:
        if not acquired:
            raise ValueError('Ya hay una revision o envio de compras en curso.')
        config = settings()
        with db() as connection:
            # With the exclusive process lock, a remaining SENDING row belongs to a crashed sender.
            connection.execute("UPDATE outbox SET status='REVISAR ENVIO',detail='Comprueba el buzon de enviados antes de repetir.' WHERE status='ENVIANDO'")
            state = _get(connection, 'state', {})
        daily = _daily_due(config, state, now)
        if not force and (not config['enabled'] or (now < state.get('next_check', 0) and not daily)):
            return {'skipped': True}
        try:
            snapshot = fetch_snapshot(config)
            snapshot['checked_at'] = _stamp(now)
            missing = {_normalize_header(s) for s in state.get('snapshot', {}).get('sheets', [])} - {
                _normalize_header(s) for s in snapshot['sheets'] + config['excluded_sheets']}
            if missing:
                raise ValueError('Faltan hojas antes reconocidas. Revisa el Excel; no se cerraron pendientes.')
            keys = {item['key'] for item in snapshot['cases']}
            with db() as connection:
                connection.execute('BEGIN IMMEDIATE')
                purchase_history.observe(connection, snapshot, _stamp(now))
                seen = {row[0] for row in connection.execute('SELECT case_key FROM reported')}
                blocked = set()
                for row in connection.execute("SELECT payload FROM outbox WHERE status IN ('PENDIENTE','ENVIANDO','ERROR','REVISAR ENVIO')"):
                    blocked.update(json.loads(row[0])['case_keys'])
                new_keys = keys - seen - blocked
                connection.executemany('DELETE FROM reported WHERE case_key=?', ((key,) for key in seen - keys))
                state.update(snapshot=snapshot, checked_at=_stamp(now), next_check=now + config['interval_minutes'] * 60, error='')
                _set(connection, 'state', state)
            preview = render(config, snapshot, new_keys, 'incremental' if state.get('started') and new_keys and not daily else 'daily')
            if force and not send_now:
                return {'preview': preview, 'snapshot': snapshot, 'new_count': len(new_keys)}
            if not config['enabled']:
                raise ValueError('Activa la automatizacion antes de enviar un reporte.')
            daily_date = datetime.fromtimestamp(now, timezone.utc).astimezone(ZoneInfo(config['timezone'])).date().isoformat()
            if not keys or (state.get('started') and not daily and not new_keys and not send_now):
                if daily:
                    state['last_daily'] = daily_date
                state['started'] = True
                with db() as connection:
                    _set(connection, 'state', state)
                return {'pending': len(keys), 'sent': False}
            # Authorize before creating any outbox row; missing consent must not consume the first report.
            token = cloud_auth.access_token(scopes=cloud_auth.MAIL_SCOPES, expected_username=config['sender'])
            if send_now:
                kind, event = 'manual', 'manual:' + uuid.uuid4().hex
            elif daily:
                kind, event = 'daily', 'daily:' + daily_date
            elif not state.get('started'):
                kind, event = 'initial', 'initial:' + hashlib.sha256('|'.join(sorted(keys)).encode()).hexdigest()
            else:
                kind, event = 'incremental', 'new:' + hashlib.sha256(('|'.join(sorted(new_keys)) + '|' + str(now)).encode()).hexdigest()
            payload = render(config, snapshot, new_keys, kind)
            with db() as connection:
                current = {**DEFAULTS, **_get(connection, 'settings', {})}
                if current != config or not current['enabled']:
                    raise ValueError('La configuracion cambio durante la revision. Vuelve a revisar.')
                if connection.execute("SELECT 1 FROM outbox WHERE status='REVISAR ENVIO'").fetchone():
                    raise ValueError('Hay un envio incierto. Verifica su resultado antes de enviar otro reporte.')
                identifier = uuid.uuid4().hex
                cursor = connection.execute('INSERT OR IGNORE INTO outbox '
                    '(id,event_key,kind,created_at,status,payload) VALUES(?,?,?,?,?,?)',
                    (identifier, event, kind, _stamp(now), 'ENVIANDO', json.dumps(payload)))
                if not cursor.rowcount:
                    return {'pending': len(keys), 'sent': False}
                state['started'] = True
                if daily:
                    state['last_daily'] = daily_date
                _set(connection, 'state', state)
            try:
                response = _post_mail(token, payload)
                result = 'ENVIADO' if response.status_code == 202 else 'ERROR' if 400 <= response.status_code < 500 and response.status_code != 408 else 'REVISAR ENVIO'
                detail = 'Aceptado por Microsoft 365; no confirma lectura.' if result == 'ENVIADO' else f'Microsoft respondio HTTP {response.status_code}; revisa permisos o el buzon de enviados.'
            except requests.RequestException:
                result, detail = 'REVISAR ENVIO', 'Resultado incierto. Comprueba el buzon de enviados antes de repetir.'
            with db() as connection:
                connection.execute('UPDATE outbox SET status=?,detail=?,sent_at=? WHERE id=?',
                    (result, detail, _stamp(now) if result == 'ENVIADO' else None, identifier))
                if result == 'ENVIADO':
                    connection.executemany('INSERT OR IGNORE INTO reported VALUES(?)', ((key,) for key in keys))
            return {'pending': len(keys), 'new_count': len(new_keys), 'sent': result == 'ENVIADO', 'mail_status': result}
        except Exception as error:
            detail = str(error) if isinstance(error, (cloud_auth.CloudAuthError, cloud_excel.CloudExcelError)) else 'No se completo la revision. Comprueba formato, hojas y configuracion; los pendientes anteriores se conservaron.'
            with db() as connection:
                state['error'] = detail
                state['next_check'] = now + 60
                _set(connection, 'state', state)
            if isinstance(error, (ValueError, cloud_auth.CloudAuthError, cloud_excel.CloudExcelError)):
                raise ValueError(detail) from error
            raise ValueError('No se completo la revision de compras.') from error


def review_delivery(identifier, delivered, actor):
    with job_lock() as acquired:
        if not acquired:
            raise ValueError('Hay un envio en curso.')
        with db() as connection:
            row = connection.execute("SELECT * FROM outbox WHERE id=? AND status='REVISAR ENVIO'", (identifier,)).fetchone()
            if not row:
                raise ValueError('El envio no necesita revision.')
            payload = json.loads(row['payload'])
            connection.execute('UPDATE outbox SET status=?,detail=? WHERE id=?',
                ('ENVIADO' if delivered else 'DESCARTADO', 'Resultado verificado por administrador.', identifier))
            if delivered:
                connection.executemany('INSERT OR IGNORE INTO reported VALUES(?)', ((key,) for key in payload['case_keys']))
            state = _get(connection, 'state', {})
            state['next_check'] = 0
            _set(connection, 'state', state)
            connection.execute('INSERT INTO audit VALUES(?,?,?)', (_stamp(), actor, 'VERIFICAR ENVIO ' + identifier))
    return status()


def start_worker():
    global _worker_started
    with _guard:
        if _worker_started:
            return
        _worker_started = True

    def loop():
        while True:
            try:
                if settings()['enabled']:
                    run()
            except Exception:
                pass  # Status retains a sanitized error; no OAuth/Excel data in process logs.
            threading.Event().wait(30)
    threading.Thread(target=loop, daemon=True, name='wms-purchase-alerts').start()


def handle(request, user):
    if user['role'] != 'ADMINISTRADOR':
        return {'error': 'Solo el administrador puede gestionar los reportes de compras.'}, 403
    path = request.path.split('?', 1)[0]
    allowed = {'status', 'settings', 'preview', 'send', 'authorize', 'delivery-review', 'cases', 'export'}
    if path not in {'/api/purchase-alerts/' + action for action in allowed}:
        return {'error': 'Ruta no encontrada.'}, 404
    if request.command in {'GET', 'HEAD'} and path == '/api/purchase-alerts/status':
        return status(), 200
    if request.command in {'GET', 'HEAD'} and path in {'/api/purchase-alerts/cases', '/api/purchase-alerts/export'}:
        params = parse_qs(urlsplit(request.path).query)
        options = {name: params.get(name, [default])[0] for name, default in
                   [('state', 'pending'), ('severity', 'ALL')]}
        if path.endswith('/export'):
            return history_csv(**options), 200
        return case_history(**options, limit=int(params.get('limit', ['50'])[0]),
                            offset=int(params.get('offset', ['0'])[0])), 200
    if path in {'/api/purchase-alerts/cases', '/api/purchase-alerts/export'}:
        return {'error': 'Usa GET para consultar el historial.'}, 405
    if request.command != 'POST':
        return {'error': 'Usa POST para esta solicitud.'}, 405
    if request.headers.get('X-WMS-Request') != '1' or request.headers.get('Sec-Fetch-Site') == 'cross-site':
        return {'error': 'Actualiza la pagina antes de enviar cambios.'}, 403
    if int(request.headers.get('Content-Length', '0')) > 32768:
        raise ValueError('Solicitud demasiado grande.')
    data = request.body()
    if not isinstance(data, dict):
        raise ValueError('Solicitud no valida.')
    if path.endswith('/settings'):
        return save(data, user['username']), 200
    if path.endswith('/preview'):
        return run(force=True), 200
    if path.endswith('/send'):
        if data.get('confirm') is not True:
            raise ValueError('Confirma el envio a los destinatarios configurados.')
        return run(force=True, send_now=True), 200
    if path.endswith('/authorize'):
        sender = settings()['sender']
        if not sender:
            raise ValueError('Guarda primero el remitente del reporte.')
        return cloud_auth.begin(cloud_auth.owner_key(request.headers), mail_sender=sender), 200
    if path.endswith('/delivery-review'):
        if data.get('confirm') is not True or type(data.get('delivered')) is not bool:
            raise ValueError('Confirma el resultado despues de revisar el buzon de enviados.')
        return review_delivery(str(data.get('id', '')), data['delivered'], user['username']), 200
    return {'error': 'Ruta no encontrada.'}, 404


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Instalar configuracion privada de compras, sin habilitar envios.')
    parser.add_argument('--configure', type=Path, required=True)
    args = parser.parse_args()
    config = json.loads(args.configure.read_text(encoding='utf-8'))
    if config.get('enabled') is not False:
        parser.error('La instalacion inicial requiere enabled=false. Activa luego desde el panel administrativo.')
    save(config, 'server-admin')
    print('Configuracion privada instalada. Automatizacion pausada; no se enviaron correos.')
