"""Bound XLSX archives before openpyxl parses them; never extract to disk."""
import os
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook as _load_workbook
from openpyxl.xml import DEFUSEDXML

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_MEMBER_BYTES = 128 * 1024 * 1024
MAX_ENTRIES = 10000


def _validate(stream):
    stream.seek(0, 2)
    if stream.tell() > MAX_FILE_BYTES:
        raise ValueError('El Excel supera el limite de 50 MB.')
    try:
        with ZipFile(stream) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ENTRIES:
                raise ValueError('El Excel contiene demasiadas entradas.')
            if sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ValueError('El Excel es demasiado grande al descomprimir.')
            if any(entry.file_size > MAX_MEMBER_BYTES for entry in entries):
                raise ValueError('Una hoja o entrada del Excel es demasiado grande.')
            if any(entry.flag_bits & 1 for entry in entries):
                raise ValueError('Usa un Excel sin cifrado ni contrasena.')
    except BadZipFile:
        raise ValueError('El archivo no es un Excel .xlsx valido.') from None


def load_workbook(filename, *args, **kwargs):
    """Keep openpyxl options/ownership; reject unsafe inputs before DB changes.

    Existing read-only worksheet iteration also uses openpyxl's defusedxml
    parser. Do not disable that parser via OPENPYXL_DEFUSEDXML in deployment.
    """
    if not DEFUSEDXML:
        raise ValueError('Falta el lector XML protegido; revisa las dependencias del servidor.')
    if isinstance(filename, (str, bytes, os.PathLike)):
        with open(filename, 'rb') as stream:
            _validate(stream)
    else:
        try:
            position = filename.tell()
        except (AttributeError, OSError, ValueError):
            raise ValueError('El lector de Excel requiere un archivo valido y reposicionable.') from None
        try:
            _validate(filename)
        finally:
            filename.seek(position)
    return _load_workbook(filename, *args, **kwargs)
