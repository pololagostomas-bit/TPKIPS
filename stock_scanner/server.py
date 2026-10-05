"""Standalone read-only NP scanner backed by the warehouse stock workbook."""

from __future__ import annotations

import json
import os
import re
import argparse
import ssl
import base64
import hmac
import getpass
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = Path(
    r"C:\Users\practicante04\OneDrive - Triton Trading\Escritorio\ANGELA SOLICITUDES\REPORTE JOSE 2109\STOCK2109.xlsx"
)
WORKBOOK = Path(os.environ.get("TRITON_STOCK_FILE", str(DEFAULT_WORKBOOK)))
HOST = os.environ.get("TRITON_SCANNER_HOST", "127.0.0.1")
PORT = int(os.environ.get("TRITON_SCANNER_PORT", "8765"))
def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def key(value) -> str:
    return re.sub(r"\s+", "", clean(value)).upper()


def quantity(value) -> float:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0


def load_inventory(path: Path) -> tuple[dict[str, dict], int]:
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo de stock: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = next((s for s in workbook.worksheets if s.title.strip().casefold() == "hoja1"), workbook.worksheets[0])
    inventory: dict[str, dict] = {}
    rows = 0
    for row in sheet.iter_rows(min_row=3, values_only=True):
        code = clean(row[0] if len(row) > 0 else None)
        description = clean(row[1] if len(row) > 1 else None)
        if not code or not description or key(code) in {"ALMACEN", "NUMERODEARTICULO"}:
            continue
        rows += 1
        item_key = key(code)
        entry = inventory.setdefault(item_key, {
            "np": code, "description": description,
            "unit": clean(row[2] if len(row) > 2 else None),
            "stock": 0, "committed": 0, "requested": 0, "available": 0,
            "locations": set(),
        })
        on_hand = quantity(row[4] if len(row) > 4 else None)
        entry["stock"] += on_hand
        entry["committed"] += quantity(row[5] if len(row) > 5 else None)
        entry["requested"] += quantity(row[6] if len(row) > 6 else None)
        entry["available"] += quantity(row[7] if len(row) > 7 else None)
        actual_location = clean(row[3] if len(row) > 3 else None)
        default_location = clean(row[8] if len(row) > 8 else None)
        # Algunas exportaciones SAP traen la tilde de "UBICACIÓN" dañada;
        # detectar la ubicación virtual sin depender de la codificación.
        system_location = "UBICACI" in actual_location.upper() and "SISTEMA" in actual_location.upper()
        location = default_location if system_location else actual_location
        if on_hand > 0 and location:
            entry["locations"].add(location)
    workbook.close()
    for entry in inventory.values():
        entry["locations"] = sorted(entry["locations"])
    return inventory, rows


class ScannerHandler(BaseHTTPRequestHandler):
    inventory: dict[str, dict] = {}
    source_rows = 0

    def is_authorized(self):
        username = os.environ.get("TRITON_SCANNER_USER", "")
        password = os.environ.get("TRITON_SCANNER_PASSWORD", "")
        if not username or not password:
            self.send_error(503, "Configura TRITON_SCANNER_USER y TRITON_SCANNER_PASSWORD antes de iniciar.")
            return False
        supplied = self.headers.get("Authorization", "")
        scheme = ""
        try:
            scheme, token = supplied.split(" ", 1)
            decoded = base64.b64decode(token, validate=True).decode("utf-8")
            given_user, given_password = decoded.split(":", 1)
        except (ValueError, UnicodeDecodeError):
            given_user, given_password = "", ""
        if scheme.lower() == "basic" and hmac.compare_digest(given_user, username) and hmac.compare_digest(given_password, password):
            return True
        self.send_response(401)
        self.send_header("WWW-Authenticate", 'Basic realm="Consulta de stock Triton", charset="UTF-8"')
        self.send_header("Content-Length", "0")
        self.end_headers()
        return False

    def send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        if not self.is_authorized():
            return
        path = urlparse(self.path).path
        if path == "/api/health":
            return self.send_json(200, {"ok": True, "products": len(self.inventory), "source_rows": self.source_rows, "source": WORKBOOK.name})
        if path.startswith("/api/product/"):
            scanned = unquote(path.removeprefix("/api/product/"))
            item = self.inventory.get(key(scanned))
            if not item:
                return self.send_json(404, {"error": "No encontré ese NP en el corte de stock."})
            return self.send_json(200, {**item, "locations": item["locations"]})
        if path in {"/", "/index.html"}:
            return self.serve_file("index.html", "text/html; charset=utf-8")
        if path in {"/app.js", "/styles.css"}:
            mime = "text/javascript; charset=utf-8" if path.endswith(".js") else "text/css; charset=utf-8"
            return self.serve_file(path.lstrip("/"), mime)
        self.send_error(404)

    def serve_file(self, name, mime):
        content = (ROOT / name).read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def log_message(self, fmt, *args):
        print(f"[{self.log_date_time_string()}] {fmt % args}")


def main():
    parser = argparse.ArgumentParser(description="Consulta de stock por escaneo de NP")
    parser.add_argument("--host", default=HOST, help="Dirección de escucha; usa 0.0.0.0 para la red local")
    parser.add_argument("--port", type=int, default=PORT)
    parser.add_argument("--cert", help="Certificado HTTPS PEM emitido para el nombre/IP del servidor")
    parser.add_argument("--key", help="Clave privada PEM del certificado HTTPS")
    args = parser.parse_args()
    if not os.environ.get("TRITON_SCANNER_USER"):
        os.environ["TRITON_SCANNER_USER"] = input("Usuario para proteger el lector: ").strip()
    if not os.environ.get("TRITON_SCANNER_PASSWORD"):
        os.environ["TRITON_SCANNER_PASSWORD"] = getpass.getpass("Contraseña para proteger el lector: ")
    if not os.environ["TRITON_SCANNER_USER"] or not os.environ["TRITON_SCANNER_PASSWORD"]:
        raise SystemExit("Debes indicar un usuario y una contraseña.")
    try:
        ScannerHandler.inventory, ScannerHandler.source_rows = load_inventory(WORKBOOK)
    except Exception as exc:
        raise SystemExit(f"No se pudo cargar el stock: {exc}") from exc
    if bool(args.cert) != bool(args.key):
        raise SystemExit("Para HTTPS indica ambos archivos: --cert y --key.")
    server = ThreadingHTTPServer((args.host, args.port), ScannerHandler)
    scheme = "http"
    if args.cert and args.key:
        context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        context.load_cert_chain(certfile=args.cert, keyfile=args.key)
        server.socket = context.wrap_socket(server.socket, server_side=True)
        scheme = "https"
    print(f"Consulta de NP lista: {scheme}://{args.host}:{args.port}")
    print(f"Fuente: {WORKBOOK} | {len(ScannerHandler.inventory):,} productos")
    print("Solo lectura. Detén el servidor con Ctrl+C.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServidor detenido.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
